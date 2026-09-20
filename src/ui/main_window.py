"""Главное окно приложения: связывает main_window.ui с расчётными сервисами."""

from PyQt6.QtWidgets import (QApplication, QMainWindow, QPlainTextEdit, QComboBox,
                             QPushButton, QSpinBox, QDateEdit, QTableWidgetItem, QTableWidget,
                             QMessageBox, QFileDialog, QGroupBox,
                             QTreeWidgetItem, QInputDialog, QMenu, QListWidgetItem,
                             QDialog, QLineEdit, QVBoxLayout, QHBoxLayout, QDialogButtonBox,
                             QLabel, QWidget, QToolButton, QWidgetAction, QFormLayout,
                             QListWidget, QListView, QFrame, QAbstractItemView, QSizePolicy,
                             QScrollArea)
from PyQt6.QtCore import QLocale, Qt, QDate, QPointF, QTimer, QSize, QMimeData, QSignalBlocker
from PyQt6.QtGui import QPixmap, QIcon, QPainter, QColor, QPen, QGuiApplication, QDrag
from PyQt6.uic import loadUi
from typing import Dict, Union
from pathlib import Path
from uuid import uuid4
from docx.shared import Mm
from docx.oxml.ns import qn
from docxtpl import DocxTemplate, InlineImage, RichText

import copy
import os
import re
import shutil
import subprocess
import html
import math
import functools
from types import SimpleNamespace

from ..equipment_types import EquipmentType, REGISTRY
from ..services.calculations import (
    calculate_strength,
    calculate_residual_life,
    generate_thickness_measurements,
    generate_ovalness_measurements,
    calculate_hardness_range,
    generate_hardness_measurements,
    find_min_thickness,
    format_year_range,
)
from ..services.formatting import (
    format_ru, format_ru_fixed, parse_ru, format_thickness_block, format_fio_initials,
    number_to_words_ru,
)
from ..services.calculations_pipeline import (
    calculate_pipeline_strength,
    calculate_pipeline_residual_life,
    generate_pipeline_thickness_measurements,
    get_allowable_stress,
    SegmentSpec,
)
from ..services.employees_store import (
    store_kleishe_image, load_employees, resolve_kleishe_path, find_employee_id_by_name,
)
from ..services.docx_layout import float_drawings_behind_text
from ..services import workspace
from ..services import custom_sections_store
from . import icons
from .open_with import open_with_prompt
from .template_location import choose_template_save_path
from .growable_placeholder_field import GrowablePlaceholderField
from .formula_editor_dialog import FormulaEditorDialog
from .table_editor_dialog import TableEditorDialog
from ..models.project import Project
from ..config import NK_SCHEME_DIR, PNEVMO_GRAPH_DIR
from .widget_names_pipeline import SEGMENT_TYPES, PROGRAM_DEFAULT_ITEMS, AE_CLASS_TYPES


# Маркер табличного поля (field_tables, см. table_editor_dialog.py) --
# get_form_data() подставляет ЭТОТ текст вместо значения поля (см.
# MainWindow._render_slot_fields(): сам виджет остаётся readOnly и
# ПУСТЫМ, показывая подсказку через setPlaceholderText(), а не этот
# маркер -- маркер существует только для Jinja-рендера). После
# tpl.render() маркер -- уже обычный текст где-то в готовом .docx;
# _splice_table_placeholders() ищет его по паттерну и меняет на
# настоящую .docx-таблицу (см. её докстринг). Символы U+E000 -- Private
# Use Area Юникода, гарантированно не встретятся в тексте, который
# реально печатает оператор -- обычный текст не может содержать их
# "случайно", в отличие, например, от "---" или "{{".
_TABLE_MARKER_RE = re.compile("table:([^]+)")


def _table_field_marker(field_id: str) -> str:
    return f"table:{field_id}"


# Тёмная тема окна конструктора документов -- палитра 1:1 из
# docs/design/constructor_mockup.html (#1c1c1e/#2c2c2e/#0a84ff/#8e8e93/
# #38383a). Применяется ТОЛЬКО к окну конструктора (self.setStyleSheet() в
# __init__ при equipment_type.id == "constructor") -- баллоны и трубопровод
# по-прежнему рендерятся нативным стилем Qt, эта тема их не затрагивает.
# В проекте нет иконочного шрифта/.qrc (см. CLAUDE.md) -- там, где в мокапе
# иконка, здесь текстовые символы (▾/▸, ×) на обычных QPushButton/QLabel.
CONSTRUCTOR_QSS = """
QMainWindow, #constructorCentral { background: #1c1c1e; }
#constructorSidebar { background: #242426; border-right: 0.5px solid #38383a; }
#constructorSidebar QLabel { color: #8e8e93; font-size: 11px; }
QPushButton[role="groupToggle"] {
    background: transparent; border: none; color: #e5e5e7; font-size: 12px;
    font-weight: 500; text-align: left; padding: 6px; border-radius: 5px;
}
QPushButton[role="groupToggle"]:hover { background: #2c2c2e; }
QPushButton[role="addVariantBtn"] {
    background: transparent; border: 0.5px dashed #48484a; border-radius: 7px;
    color: #8e8e93; font-size: 11.5px; text-align: left; padding: 7px 9px;
    margin: 2px 0;
}
QPushButton[role="addVariantBtn"]:hover {
    border-color: #0a84ff; color: #e5e5e7;
}
/* «+ Добавить раздел» (см. __init__/_open_add_section_dialog()) -- сплошная
   (не пунктирная) синяя рамка, чтобы визуально отличаться от «Добавить»
   ВНУТРИ раздела (тот пунктирный, role="addVariantBtn" выше) -- тот же
   приём, что и в мокапе (docs/design/вводная_часть.html). */
#addSectionBtn {
    background: transparent; border: 0.5px solid rgba(10, 132, 255, 110); border-radius: 7px;
    color: #0a84ff; font-size: 11.5px; text-align: left; padding: 7px 9px;
    margin: 2px 0 10px;
}
#addSectionBtn:hover { background: rgba(10, 132, 255, 24); }
/* [role="..."] вместо перечисления #titleFoo, #introFoo, #appendix1Foo --
   раздел конструктора теперь не фиксированная тройка (см. MainWindow.
   _create_section()/_build_custom_section_widgets()), а произвольный набор
   слотов, заранее не известный на момент написания этого QSS -- имя
   объекта каждого нового раздела уникально ("section1FieldsPanel" и т.п.),
   а вот role -- одна из фиксированных ролей ("fieldsPanel" и т.п.),
   выставляется программно на КАЖДЫЙ слот одинаково в _init_constructor_slot()
   независимо от того, встроенный он или добавлен пользователем. */
QListWidget[role="availableBlocksList"] {
    background: transparent; border: none; outline: none; font-size: 11.5px;
}
QListWidget[role="availableBlocksList"]::item {
    background: #2c2c2e; border: 0.5px solid #38383a; border-radius: 7px;
    padding: 5px 8px; margin: 2px 0; color: #e5e5e7;
}
QListWidget[role="availableBlocksList"]::item:hover { border-color: #0a84ff; }
QListWidget[role="availableBlocksList"]::item:selected { background: #2c2c2e; }
/* Карточка варианта, включённого в документ сейчас (см.
   _highlight_available_block_items()) -- тот же стиль, что и у чипа
   плейсхолдера в реквизитах (QLabel[titleChip="true"] ниже), просто
   применённый к itemWidget всей карточки, а не к одной подписи поля. */
QWidget#availableBlockChip {
    background: rgba(10, 132, 255, 40); border: 1px solid rgba(10, 132, 255, 110); border-radius: 6px;
}
QWidget#availableBlockChip QLabel { background: transparent; color: #5ab4ff; font-size: 11.5px; }
QLabel#crumbLabel {
    background: #202022; color: #8e8e93; font-size: 11px; padding: 8px 18px;
    border-bottom: 0.5px solid #2c2c2e;
}
QLabel#documentSectionLabel { color: #8e8e93; font-size: 11px; }
QListWidget[role="includedBlockList"] {
    background: transparent; outline: none; border: none;
}
QListWidget[role="includedBlockList"][filled="false"] {
    border: 1.5px dashed #38383a; border-radius: 8px;
}
QListWidget[role="includedBlockList"][filled="true"]::item {
    background: #2c2c2e; border-radius: 8px; padding: 0; margin: 0;
}
QListWidget[role="includedBlockList"][dragOver="true"] {
    border: 1.5px dashed #0a84ff; border-radius: 8px; background: rgba(10, 132, 255, 24);
}
QGroupBox[role="fieldsPanel"] {
    border: none; margin-top: 14px; padding-top: 0;
}
QFrame[role="slotSeparator"] {
    background: #545456; margin-top: 20px; margin-bottom: 12px; border-radius: 1px;
}
QGroupBox[role="fieldsPanel"] QLabel {
    color: #c7c7cc; font-size: 12px;
}
QGroupBox[role="fieldsPanel"] QLabel#titleFieldsSectionLabel {
    color: #8e8e93; font-size: 11px; margin-top: 4px;
}
QGroupBox[role="fieldsPanel"] QPlainTextEdit {
    background: #2c2c2e; border: 0.5px solid #38383a; border-radius: 5px;
    color: #e5e5e7; font-size: 12px; padding: 6px 8px;
}
QGroupBox[role="fieldsPanel"] QPlainTextEdit:focus { border-color: #0a84ff; }
QGroupBox[role="fieldsPanel"] QToolButton {
    background: transparent; border: none; color: #c7c7cc; font-size: 11.5px;
    padding: 4px 8px; border-radius: 5px;
}
QGroupBox[role="fieldsPanel"] QToolButton:hover { background: #3a3a3c; color: #e5e5e7; }
QGroupBox[role="fieldsPanel"] QToolButton::menu-indicator {
    width: 8px; height: 8px; subcontrol-position: right center;
    subcontrol-origin: padding; right: 4px;
}
QGroupBox[role="fieldsPanel"] QLabel[titleChip="true"] {
    background: rgba(10, 132, 255, 40); color: #5ab4ff;
    border: 1px solid rgba(10, 132, 255, 110); border-radius: 6px;
    padding: 4px 9px; font-size: 11.5px;
}
QGroupBox[role="fieldsPanel"] QLabel[titleChipCopied="true"] {
    background: rgba(48, 209, 88, 40); color: #30d158;
    border: 1px solid rgba(48, 209, 88, 140);
}
QGroupBox[role="fieldsPanel"] QLabel[chipDragOver="true"] {
    border: 1.5px dashed #0a84ff; background: rgba(10, 132, 255, 70);
}
QGroupBox[role="fieldsPanel"] QLabel[titleChipFormula="true"] {
    background: rgba(191, 90, 242, 40); color: #d29dfa;
    border: 1px solid rgba(191, 90, 242, 140);
}
QGroupBox[role="fieldsPanel"] QPlainTextEdit[computed="true"] {
    background: rgba(191, 90, 242, 24); border-color: rgba(191, 90, 242, 140); color: #d29dfa;
}
/* Чип поля, представленного таблицей (см. table_editor_dialog.py) --
   отдельный акцент (голубой), чтобы отличать и от обычных плейсхолдеров
   (синие), и от вычисляемых по формуле (фиолетовые). */
QGroupBox[role="fieldsPanel"] QLabel[titleChipTable="true"] {
    background: rgba(100, 210, 255, 40); color: #64d2ff;
    border: 1px solid rgba(100, 210, 255, 140);
}
QGroupBox[role="fieldsPanel"] QPlainTextEdit[computedTable="true"] {
    background: rgba(100, 210, 255, 24); border-color: rgba(100, 210, 255, 140);
}
QWidget#titleTemplateDropHint {
    border: 1.5px dashed #38383a; border-radius: 8px;
}
QWidget#titleTemplateDropZone[dragActive="true"] QWidget#titleTemplateDropHint {
    border-color: #0a84ff; background: rgba(10, 132, 255, 24);
}
QWidget#titleTemplateLoadedCard {
    background: #2c2c2e; border: 0.5px solid #38383a; border-radius: 7px;
}
QWidget#titleTemplateLoadedCard QLabel { color: #e5e5e7; font-size: 11.5px; }
QWidget#titleTemplateLoadedCard QPushButton {
    background: transparent; border: none; border-radius: 5px;
}
QWidget#titleTemplateLoadedCard QPushButton:hover { background: #3a3a3c; }
QWidget#titlePreviewField {
    background: #2c2c2e; border: 0.5px solid #38383a; border-radius: 7px;
}
QWidget#titlePreviewField QLabel { color: #8e8e93; font-size: 11.5px; }
QWidget#titlePreviewField QLabel[titlePreviewFieldActive="true"] { color: #e5e5e7; }
#constructorBottomBar { background: #1c1c1e; border-top: 0.5px solid #38383a; }
#constructorBottomBar QPushButton {
    background: transparent; border: none; color: #c7c7cc; font-size: 12px;
    padding: 9px; border-left: 0.5px solid #38383a;
}
#constructorBottomBar QPushButton#pushButt_generateWord { border-left: none; }
#constructorBottomBar QPushButton:disabled { color: #5a5a5c; }

/* «Объекты» -- дерево папок/документов (objectsTree, тот же виджет, что и
   у трубопровода, см. _refresh_objects_tree() в main_window.py), здесь
   просто одетый в тёмную тему, а не нативный стиль Qt. */
#revealStrip { background: #242426; }
#revealStrip QPushButton { background: transparent; border: none; color: #8e8e93; }
#revealStrip QPushButton:hover { color: #e5e5e7; }
QWidget#sidebar { background: #242426; border-right: 0.5px solid #38383a; }
QWidget#sidebar QLabel#sidebarHeader_objects { color: #8e8e93; font-size: 11px; font-weight: 500; }
QWidget#sidebar QPushButton#sidebarBtn_collapse {
    background: transparent; border: none; color: #8e8e93; font-size: 11px;
}
QWidget#sidebar QPushButton#sidebarBtn_collapse:hover { color: #e5e5e7; }
QLineEdit#sidebarSearchBox {
    background: #1c1c1e; border: 0.5px solid #38383a; border-radius: 6px;
    color: #e5e5e7; font-size: 12px; padding: 6px 8px;
}
QLineEdit#sidebarSearchBox:focus { border-color: #0a84ff; }
QWidget#sidebar QPushButton#sidebarBtn_createObject, QWidget#sidebar QPushButton#sidebarBtn_createDocument {
    background: transparent; border: 0.5px dashed #48484a; border-radius: 7px;
    color: #8e8e93; font-size: 11.5px; text-align: left; padding: 7px 9px; margin: 2px 0;
}
QWidget#sidebar QPushButton#sidebarBtn_createObject:hover, QWidget#sidebar QPushButton#sidebarBtn_createDocument:hover {
    border-color: #0a84ff; color: #e5e5e7;
}
QTreeWidget#objectsTree {
    background: transparent; border: none; outline: none; color: #e5e5e7; font-size: 11.5px;
}
QTreeWidget#objectsTree::item { padding: 4px 2px; border-radius: 5px; }
QTreeWidget#objectsTree::item:hover { background: #2c2c2e; }
QTreeWidget#objectsTree::item:selected { background: rgba(10, 132, 255, 40); color: #e5e5e7; }
QTreeWidget#objectsTree QHeaderView::section {
    background: #242426; color: #5a5a5c; font-size: 10px; border: none; padding: 4px 2px;
}
QSplitter#mainSplitter::handle { background: #38383a; }

/* Активити-бар (docs/design/вводная_часть.html, #activityBar) -- узкая
   колонка иконок слева от viewStack, см. _switch_activity_view(). */
QWidget#activityBar { background: #1c1c1e; border-right: 0.5px solid #38383a; }
QWidget#activityBar QToolButton { border: none; border-radius: 0; color: #8e8e93; }
QWidget#activityBar QToolButton:hover { color: #e5e5e7; }
QWidget#activityBar QToolButton:checked { color: #e5e5e7; border-left: 2px solid #0a84ff; }
QWidget#page_database QLabel#databaseBreadcrumbLabel {
    color: #8e8e93; font-size: 11px; text-transform: uppercase; margin-bottom: 4px;
}
QWidget#page_database QLabel#databaseHintLabel { color: #5a5a5c; font-size: 11.5px; font-style: italic; }
QWidget#page_database QLabel#databasePreviewTitle { color: #e5e5e7; font-size: 12.5px; font-weight: 600; }
QWidget#page_database QLabel#databasePreviewFieldLabel { color: #8e8e93; font-size: 11.5px; }
QWidget#page_database QLabel#databasePreviewFieldValue { color: #c7c7cc; font-size: 12px; }
QWidget#page_stub QLabel#activityStubTitle { color: #8e8e93; font-size: 13px; }
QWidget#page_stub QLabel#activityStubHint { color: #5a5a5c; font-size: 12px; }
"""


class MainWindow(QMainWindow):
    def __init__(self, equipment_type: EquipmentType = REGISTRY["balloon"]):
        """Инициализация конструктора класса. Пишем все атрибуты,
        что пригодятся нам по коду.

        equipment_type определяет, какой .ui загружать, какие виджеты
        резолвить и в каком порядке требовать шаги расчёта — см.
        src/equipment_types.py. По умолчанию баллоны, чтобы существующие
        вызовы MainWindow() не меняли поведение.
        """
        super().__init__()
        self.data = {}
        self.text = []
        self.s_min_lst = []
        self.file_handler = None
        self._completed_steps = set()
        self._current_document_path = None

        self.equipment_type = equipment_type
        # list(...) -- копия, не ссылка: конструктор документов дописывает
        # сюда имена динамически созданных полей реквизитов в рантайме (см.
        # _render_slot_fields()); без копии .append() мутировал бы сам
        # модуль widget_names_constructor.py между запусками окна.
        self.PLAIN_TEXT_EDIT_NAMES = list(equipment_type.widget_names.PLAIN_TEXT_EDIT_NAMES)
        self.COMBO_BOX_NAMES = equipment_type.widget_names.COMBO_BOX_NAMES
        self.DATE_EDIT_NAMES = equipment_type.widget_names.DATE_EDIT_NAMES
        self.BUTTON_NAMES = equipment_type.widget_names.BUTTON_NAMES
        self.SPIN_BOX_NAMES = equipment_type.widget_names.SPIN_BOX_NAMES
        self.TABLE_WIDGET = equipment_type.widget_names.TABLE_WIDGET
        self.STEP_ORDER = equipment_type.step_order
        self.STEP_LABELS = equipment_type.step_labels

        loadUi(str(equipment_type.ui_path), self)

        # Автоматическая инициализация виджетов
        self.init_widgets()

        # Инициализация FileHandler
        self.init_file_handler()

        # Подключение сигналов — общие для всех типов
        self.pushButt_generateWord.clicked.connect(self.calculate)
        self.pushButton_saveProject.clicked.connect(self.save_project)
        self.pushButton_openProject.clicked.connect(self.open_project)

        if equipment_type.id == "balloon":
            self.pushButton_exportCSV.clicked.connect(self.export_csv)
            self.pushButt_amount.clicked.connect(self.fill_table)
            self.pushButt_sMinMin.clicked.connect(self.s_min_min_calc)
            self.pushButton_creatThickness.clicked.connect(self.calc_thick)
            self.pushButton_creatRasschProchn.clicked.connect(self.prochnost)
            self.pushButton_creatOstRes.clicked.connect(self.ost_res)
            self.pushButt_ovalnost.clicked.connect(self.ovalnost_calc)
            self.pushButt_tverdost.clicked.connect(self.tverdost)
            self.pushButton_importCSV.clicked.connect(self.import_csv)
        elif equipment_type.id == "pipeline":
            # Комбобоксы выбора специалиста (Приложения 1-6, 8, 9) --
            # хранят не текст, а индекс строки table_specialists (см.
            # _refresh_specialist_combo()), поэтому не входят ни в один
            # список widget_names_pipeline.py. Из-за этого выбор нигде
            # не попадал в JSON проекта и при открытии всегда откатывался
            # на первую строку -- см. get_form_data()/_fill_ui_from_project()
            # в file_handler.py и _refresh_program_specialist_combo() ниже.
            self.SPECIALIST_COMBO_NAMES = [
                "program_specialist", "act2_specialist", "vik_specialist",
                "thick_specialist", "uzk_specialist", "calc_specialist",
                "pnevmo_specialist", "ae_zakl_specialist",
            ]
            # employee_id сотрудника из справочника «Сотрудники» для каждой
            # строки table_specialists, по порядку строк -- см.
            # _add_specialist_row()/_remove_specialist_row(). Не входит ни в
            # TABLE_WIDGET, ни в один список widget_names_pipeline.py --
            # сохраняется/восстанавливается отдельно, см. file_handler.py.
            self._specialist_employee_ids = []
            # Префиксы плейсхолдеров подписи, для которых в шаблон
            # подставляется клише специалиста (InlineImage), см. calculate()
            # и _specialist_kleishe_image(). lead_specialist -- всегда первая
            # строка table_specialists, остальные 8 -- SPECIALIST_COMBO_NAMES
            # (program_specialist -> programm_specialist_* и т.п., названия
            # плейсхолдеров исторически с двумя "м", см. calculate()).
            self.KLEISHE_ROLE_PREFIXES = [
                "lead_specialist", "programm_specialist", "act2_specialist",
                "vik_specialist", "thick_specialist", "uzk_specialist",
                "calc_specialist", "pnevmo_specialist", "ae_zakl_specialist",
            ]
            self.pushButt_segments.clicked.connect(self.fill_segments_table)
            self.pushButton_genThickness.clicked.connect(self.calc_pipeline_thickness)
            self.pushButton_calcStrength.clicked.connect(self.calc_pipeline_strength_ui)
            self.pushButton_calcResidualLife.clicked.connect(self.calc_pipeline_residual_life_ui)
            self.pushButt_addSpecialist.clicked.connect(self._add_specialist_row)
            self.pushButt_removeSpecialist.clicked.connect(self._remove_specialist_row)
            self.pushButt_addReviewedDoc.clicked.connect(lambda: self._add_table_row(self.table_reviewed_docs))
            self.pushButt_removeReviewedDoc.clicked.connect(lambda: self._remove_table_row(self.table_reviewed_docs))
            self.pushButt_addDocSection5.clicked.connect(lambda: self._add_table_row(self.table_docs_section5))
            self.pushButt_removeDocSection5.clicked.connect(lambda: self._remove_table_row(self.table_docs_section5))
            self.pushButt_addVikVisualRow.clicked.connect(lambda: self._add_table_row(self.table_vik_visual))
            self.pushButt_removeVikVisualRow.clicked.connect(lambda: self._remove_table_row(self.table_vik_visual))
            self.pushButt_addVikMeasureRow.clicked.connect(lambda: self._add_table_row(self.table_vik_measure))
            self.pushButt_removeVikMeasureRow.clicked.connect(lambda: self._remove_table_row(self.table_vik_measure))
            self.pushButt_removeThickDevice.clicked.connect(lambda: self._remove_combo_current_item(self.thick_device))
            self.pushButt_removeUzkDevice.clicked.connect(lambda: self._remove_combo_current_item(self.uzk_device))
            self.pushButt_addUzkRow.clicked.connect(self._add_uzk_row)
            self.pushButt_removeUzkRow.clicked.connect(self._remove_uzk_row)
            self.pushButt_removePnevmoDevice.clicked.connect(lambda: self._remove_combo_current_item(self.pnevmo_device))
            self.pushButt_removeReportTitle.clicked.connect(lambda: self._remove_combo_current_item(self.report_title))
            self.pushButt_addPipeMaterial.clicked.connect(self._add_pipe_material_row)
            self.pushButt_removePipeMaterial.clicked.connect(lambda: self._remove_table_row(self.table_pipe_materials))
            self.p_rab_mpa.textChanged.connect(self._update_p_rab_kgs)
            self.p_rab_mpa.textChanged.connect(self._update_calc_p_rab_display)
            self.p_rab_kgs.textChanged.connect(self._update_pnevmo_pressure)
            self.p_rab_kgs.textChanged.connect(self._fill_pnevmo_stages_table)
            self.work_temp.textChanged.connect(self._update_calc_temp_display)
            self.years_of_operation.textChanged.connect(self._update_calc_years_operation_display)
            self.obj_naznach.textChanged.connect(self._update_pnevmo_obj_naznach_display)
            self.pnevmo_obj_naznach.textChanged.connect(self._update_ae_zakl_obj_control_display)
            self.pnevmo_date.dateChanged.connect(self._update_ae_zakl_date_display)
            self.report_date.dateChanged.connect(self._update_ae_zakl_report_date_display)
            self.pnevmo_pressure.textChanged.connect(self._update_pnevmo_pressure_hint)
            self.pnevmo_pressure.textChanged.connect(self._update_result_76_pressure)
            self.final_years_allowed.textChanged.connect(self._update_final_deadline_date)
            self.report_year.textChanged.connect(self._update_final_deadline_date)
            self.report_year.textChanged.connect(self._update_years_of_operation_display)
            self.year_start.textChanged.connect(self._update_years_of_operation_display)
            self.pushButt_addProgramItem.clicked.connect(self._add_program_item_row)
            self.pushButt_addProgramSubitem.clicked.connect(self._add_program_subitem_row)
            self.pushButt_removeProgramRow.clicked.connect(self._remove_program_row)
            self._seed_program_table_defaults()
            self.pushButt_chooseNkScheme.clicked.connect(self._choose_nk_scheme)
            self.pushButt_clearNkScheme.clicked.connect(self._clear_nk_scheme)
            self.pushButt_fillPnevmoAeTable.clicked.connect(self._fill_pnevmo_ae_table)
            self.pushButt_addPnevmoAeRow.clicked.connect(self._add_pnevmo_ae_row)
            self.pushButt_removePnevmoAeRow.clicked.connect(lambda: self._remove_table_row(self.table_pnevmo_ae))
            self.pushButt_addPnevmoStageRow.clicked.connect(lambda: self._add_table_row(self.table_pnevmo_stages))
            self.pushButt_removePnevmoStageRow.clicked.connect(lambda: self._remove_table_row(self.table_pnevmo_stages))
            self.pushButt_choosePnevmoGraph.clicked.connect(self._choose_pnevmo_graph)
            self.pushButt_clearPnevmoGraph.clicked.connect(self._clear_pnevmo_graph)

            from .employees_tab import EmployeesTabController
            self.employees_tab = EmployeesTabController(self)

            from .instruments_tab import InstrumentsTabController
            self.instruments_tab = InstrumentsTabController(self)

            from .orgdocs_tab import OrgDocsTabController
            self.orgdocs_tab = OrgDocsTabController(self)

            # Сайдбар: переключатели «Сотрудники»/«Приборы»/«Документы» —
            # общие справочники компании, не часть текущего отчёта: кнопки
            # генерации Word и работы с проектом там неуместны, см.
            # _update_report_buttons_visibility(). Дерево объектов/
            # документов (objectsTree) — навигация внутри "document".
            self.sidebarBtn_employees.clicked.connect(lambda: self._switch_view("employees"))
            self.sidebarBtn_instruments.clicked.connect(lambda: self._switch_view("instruments"))
            self.sidebarBtn_orgdocs.clicked.connect(lambda: self._switch_view("orgdocs"))
            self._current_view = "document"

            # Дерево "объект (папка) -> документ (.json внутри неё)" --
            # см. src/services/workspace.py. Клик по документу открывает
            # его (см. _open_document()); клик по объекту — стандартное
            # разворачивание/сворачивание QTreeWidget, ничего сверху не
            # навешено.
            self.sidebarBtn_createObject.clicked.connect(self._create_object_dialog)
            self.sidebarBtn_createDocument.clicked.connect(self._create_document_dialog)
            self.sidebarBtn_templates.clicked.connect(self._show_templates_menu)
            self.objectsTree.itemClicked.connect(self._on_objects_tree_item_clicked)
            self.objectsTree.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
            self.objectsTree.customContextMenuRequested.connect(self._show_objects_tree_context_menu)
            self._refresh_objects_tree()
            self.sidebarSearchBox.textChanged.connect(self._filter_objects_tree)
            self._sidebar_collapsed = False
            # mainSplitter -- три ребёнка (revealStrip, sidebar, viewStack,
            # Фаза 7), sizes должен покрывать все три -- список короче
            # count() даёт непредсказуемое распределение (найдено
            # реальным багом: sidebar защёлкивался на maximumWidth=400
            # вместо заданных 240, потому что setSizes([240, 760]) на
            # трёх виджетах интерпретировался не так, как рассчитывали).
            self._sidebar_expanded_sizes = [0, 240, 760]
            self.sidebarBtn_collapse.clicked.connect(self._toggle_sidebar)
            self.sidebarBtn_expand.clicked.connect(self._toggle_sidebar)
            # sidebarBtn_pin -- чисто визуальный тумблер, как и в
            # референсе (togglePin() там тоже только перекрашивает
            # иконку, без функционального эффекта): checkable=true +
            # QPushButton:checked в styleSheet (.ui) уже даёт нужный вид
            # сам, без единой строки кода здесь.

            # Оглавление документа (Фаза 6): вложено в objectsTree как
            # дочерние строки открытого документа -- список groupbox'ов
            # статичен (не зависит от того, какой документ открыт,
            # форма всегда одна и та же), берётся из самой ленты
            # tab_document, а не хардкодится -- если разделы переставят
            # в Designer'е, список подстроится сам. Клики по строкам TOC
            # идут через тот же objectsTree.itemClicked, что и по
            # объектам/документам -- см. _on_objects_tree_item_clicked().
            self._suppress_toc_spy = False
            self._toc_groupboxes = self._collect_toc_groupboxes()
            self._current_toc_items = []
            self._current_toc_active_item = None
            self._toc_check_icon = self._make_check_icon()
            self._toc_circle_icon = self._make_circle_icon()
            self._toc_lock_icon = self._make_lock_icon()
            self.tab_document_scroll.verticalScrollBar().valueChanged.connect(self._on_document_scrolled)

            # Индикатор несохранённых изменений (Фаза 5.4) -- точка на
            # строке текущего документа в objectsTree.
            self._document_dirty = False
            self._dirty_icon = self._make_dot_icon("#e0983c")
            self._connect_dirty_tracking()

            # Ширина сайдбара задаётся кодом: .ui-формат не умеет
            # сериализовать QSplitter.sizes (нет XML-типа под QList<int>).
            # Диапазон перетаскивания ограничен minimumSize/maximumSize
            # самого sidebar (180..400 px в .ui). Индексы -- по составу
            # mainSplitter после Фазы 7: 0=revealStrip, 1=sidebar,
            # 2=viewStack; растягивается при изменении размера окна
            # только форма (индекс 2), сайдбар держит выставленную
            # ширину.
            self.mainSplitter.setSizes(self._sidebar_expanded_sizes)
            self.mainSplitter.setStretchFactor(0, 0)
            self.mainSplitter.setStretchFactor(1, 0)
            self.mainSplitter.setStretchFactor(2, 1)
            # Ручка перетаскивания шире дефолтной (~3-4px) -- за неё
            # неудобно было попасть мышью, отсюда и ощущение, что
            # ширину нельзя менять руками.
            self.mainSplitter.setHandleWidth(6)

            # Нижняя панель -- три равные колонки, как в референсе.
            # <property name="stretch"> в .ui не сработал бы: uic.loadUi()
            # (в отличие от кодогенератора pyuic6) не умеет разбирать
            # строку "1,1,1" для QHBoxLayout.setStretch(), падает на
            # старте (проверено headless-тестом) -- выставляется кодом.
            for i in range(self.bottom_buttons.count()):
                self.bottom_buttons.setStretch(i, 1)

            # Устанавливает начальный вид -- при запуске ни один документ
            # ещё не открыт (_current_document_path is None), TOC-строк
            # в дереве нет (появляются в _open_document()).
            self._switch_view("document")
            self.breadcrumbLabel.setOpenExternalLinks(False)
            self.breadcrumbLabel.linkActivated.connect(self._on_breadcrumb_link_activated)
            self._update_breadcrumb()
        elif equipment_type.id == "constructor":
            # Перетаскивание между списками настроено декларативно в самом
            # .ui (dragEnabled/acceptDrops/dragDropMode) -- Qt по умолчанию
            # кодирует все роли item'а (включая UserRole) в MIME при
            # перетаскивании между двумя QListWidget, так что специального
            # кода на сам drag не нужно. pushButt_generateWord уже
            # подключена к self.calculate выше (общая кнопка для всех
            # типов) -- она же и «Собрать документ» для конструктора.
            #
            # Три встроенных раздела -- «Титульные листы» (title), «Вводная
            # часть» (intro) и «Приложение 1» (appendix1), см.
            # _CONSTRUCTOR_SLOTS -- заполняются (или нет) независимо друг от
            # друга, у каждого свой набор виджетов сайдбара/области документа
            # (см. .ui: available*BlocksList/add*VariantBtn/*GroupToggle/
            # included*BlockList/*fieldsPanel). Инициализация вынесена в
            # generic _init_constructor_slot(), вызывается один раз на
            # каждый слот -- и на эти три, и (ниже) на любой раздел, который
            # оператор добавит поверх них (_create_section()).
            #
            # self._CONSTRUCTOR_SLOTS -- собственная (per-instance) копия
            # словаря класса, В ГЛУБИНУ НА ОДИН УРОВЕНЬ: _create_section()/
            # _delete_section() ниже мутируют сам словарь (добавляют/убирают
            # ключи-слоты), а _rename_section() мутирует ЗНАЧЕНИЕ отдельного
            # слота (display_label/empty_hint/section_label) -- голого
            # dict(self._CONSTRUCTOR_SLOTS) для второго недостаточно: он
            # копирует только внешний словарь, вложенный dict каждого
            # встроенного слота остался бы ТЕМ ЖЕ объектом, что и в
            # классе -- переименование "title" в одном окне тихо
            # испортило бы значение по умолчанию для всех остальных
            # (включая уже открытые/будущие окна и тесты в том же процессе).
            self._CONSTRUCTOR_SLOTS = {slot: dict(config) for slot, config in self._CONSTRUCTOR_SLOTS.items()}
            # Та же причина -- _rename_section() убирает пункт слота из
            # этих двух словарей (см. её докстринг), они тоже class-level.
            self._GENERATE_DIALOG_SLOT_LABELS = dict(self._GENERATE_DIALOG_SLOT_LABELS)
            self._ADD_VARIANT_DIALOG_TITLE = dict(self._ADD_VARIANT_DIALOG_TITLE)
            self.setStyleSheet(CONSTRUCTOR_QSS)
            self._dynamic_field_names = {slot: [] for slot in self._CONSTRUCTOR_SLOTS}
            # (slot, variant_id) -> путь последнего сгенерированного .docx
            # предпросмотра (лежит рядом с самим файлом шаблона, см.
            # _open_variant_preview()) -- сам файл переживает перезапуск
            # приложения (перезаписывается на месте, не в temp-каталоге),
            # а вот эта карта -- нет: она только про то, открывали ли
            # предпросмотр этого варианта в ТЕКУЩЕМ запуске (визуальное
            # состояние поля, см. _build_preview_field()). Значение (не
            # только сам факт наличия ключа) нужно кнопкам «Сохранить
            # как»/«Показать в Finder» на поле предпросмотра
            # (_reveal_preview()/_save_preview_as()), чтобы знать, какой
            # именно файл открывать.
            self._preview_generated = {}

            # Путь документа, чьё содержимое сейчас показано в правой
            # части «Базы документов» (см. _show_document_preview()) --
            # None, пока ничего не выбрано (тогда видна только
            # databaseHintLabel). Нужен, чтобы _delete_document() мог
            # сбросить превью, если удалили именно просматриваемый сейчас
            # документ (иначе справа осталось бы превью уже
            # несуществующего файла).
            self._database_preview_path = None

            for slot in self._CONSTRUCTOR_SLOTS:
                self._init_constructor_slot(slot)

            # «+ Добавить раздел» -- не «Добавить» ВНУТРИ раздела (тот
            # добавляет вариант в существующий, см. add*VariantBtn выше) --
            # создание НОВОГО раздела целиком, см. _open_add_section_dialog()/
            # _create_section(). Вставляется последней в sidebarLayout,
            # перед хвостовым spacer'ом -- тем же приёмом (insertWidget по
            # индексу count()-1), что и группы разделов, добавленных позже
            # (см. _build_custom_section_widgets()), чтобы кнопка всегда
            # оставалась под уже существующими группами, а не над ними.
            add_section_btn = QPushButton("  Добавить раздел")
            add_section_btn.setObjectName("addSectionBtn")
            add_section_btn.setFlat(True)
            add_section_btn.setCursor(Qt.CursorShape.PointingHandCursor)
            add_section_btn.setIcon(icons.icon("plus", "#0a84ff", 13))
            add_section_btn.setIconSize(QSize(13, 13))
            add_section_btn.clicked.connect(self._open_add_section_dialog)
            self.addSectionBtn = add_section_btn
            self.sidebarLayout.insertWidget(self.sidebarLayout.count() - 1, add_section_btn)

            # Разделы, добавленные оператором в прошлых запусках приложения
            # (см. custom_sections_store.py) -- тем же путём, что и
            # «+ Добавить раздел» выше, только без повторной записи в файл
            # (id уже там, persist=False).
            for section in custom_sections_store.load_sections():
                self._create_section(section['id'], section['label'], persist=False)

            self.pushButt_generateWord.setEnabled(False)

            # Активити-бар (docs/design/вводная_часть.html, #activityBar) --
            # узкая колонка иконок слева от viewStack, переключает ЦЕЛИКОМ
            # страницу viewStack (см. _switch_activity_view()): «Редактор
            # документов» (палитра блоков + канва, страница tab_document,
            # уже существовала до этой правки) и «База документов» (дерево
            # объектов/документов, страница page_database, см. ниже) --
            # взаимоисключающие разделы, а не два одновременно видимых
            # сайдбара, как было в предыдущей (ошибочной) версии этой
            # правки. «Поиск»/«Сотрудники» -- заглушка (page_stub), как и в
            # самом мокапе (activityViewTitles там же).
            _activity_icons = {"search": "search", "database": "database", "editor": "edit", "employees": "users"}
            for name, icon_name in _activity_icons.items():
                btn = getattr(self, f"activityBtn_{name}")
                btn.setIcon(icons.icon(icon_name, "#8e8e93", 19))
                btn.setIconSize(QSize(19, 19))
            self.activityBtn_search.clicked.connect(lambda: self._switch_activity_view("search"))
            self.activityBtn_database.clicked.connect(lambda: self._switch_activity_view("database"))
            self.activityBtn_editor.clicked.connect(lambda: self._switch_activity_view("editor"))
            self.activityBtn_employees.clicked.connect(lambda: self._switch_activity_view("employees"))

            # Дерево "объект (папка) -> документ (.json внутри неё)" --
            # тот же workspace.py и тот же objectsTree/mainSplitter/
            # sidebar, что и у трубопровода (equipment_type.id ==
            # "pipeline" выше), просто без TOC/специалистов/индикатора
            # несохранённых правок -- их у конструктора нет, см.
            # _open_constructor_document()/_reset_constructor_form(). В
            # отличие от трубопровода, здесь этот блок живёт на отдельной
            # странице viewStack (page_database), а не постоянно рядом с
            # формой -- см. constructor_window.ui. constructor_window.ui
            # объявляет те же имена виджетов (sidebar/objectsTree/
            # sidebarBtn_*/mainSplitter/revealStrip), поэтому общие методы
            # (_refresh_objects_tree(), _filter_objects_tree(),
            # _toggle_sidebar(), _show_objects_tree_context_menu() и т.п.)
            # переиспользуются без изменений.
            self.sidebarBtn_createObject.clicked.connect(self._create_object_dialog)
            self.sidebarBtn_createDocument.clicked.connect(self._create_document_dialog)
            self.objectsTree.itemClicked.connect(self._on_objects_tree_item_clicked)
            self.objectsTree.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
            self.objectsTree.customContextMenuRequested.connect(self._show_objects_tree_context_menu)
            self._refresh_objects_tree()
            self.sidebarSearchBox.textChanged.connect(self._filter_objects_tree)
            self._sidebar_collapsed = False
            # mainSplitter -- три ребёнка (revealStrip, sidebar,
            # databaseContent), тот же приём, что и у трубопровода (см. его
            # инициализацию выше) -- sizes должен покрывать все три сразу.
            self._sidebar_expanded_sizes = [0, 220, 680]
            self.sidebarBtn_collapse.clicked.connect(self._toggle_sidebar)
            self.sidebarBtn_expand.clicked.connect(self._toggle_sidebar)
            self.mainSplitter.setSizes(self._sidebar_expanded_sizes)
            self.mainSplitter.setStretchFactor(0, 0)
            self.mainSplitter.setStretchFactor(1, 0)
            self.mainSplitter.setStretchFactor(2, 1)
            self.mainSplitter.setHandleWidth(6)
            self._switch_activity_view("editor")

    # Имена виджетов на каждый слот конструктора -- title (титульные листы,
    # оригинальный Phase 1), intro (вводная часть, второй независимый слот)
    # и appendix1 (приложение 1, третий независимый слот). Ключи --
    # логические роли, используемые везде ниже через
    # _slot_widget()/_slot_config(); сами .ui-объекты см.
    # src/ui/designer/constructor_window.ui. Порядок словаря = порядок
    # инициализации в __init__ и порядок в собранном документе (см.
    # _calculate_constructor()) -- title первым, как и раньше.
    #
    # text_label_attr/chevron_attr -- имена Python-атрибутов (НЕ .ui-виджетов
    # -- те динамические, создаются в _process_block_drop(), setattr(self, ...))
    # под подпись и шеврон уже вставленного в документ блока этого слота.
    # title исторически без "Title" в имени (оставлено как есть, слот
    # существовал ДО того, как появилась сама концепция "слот" -- менять
    # имя атрибута задним числом незачем), остальные слоты -- по образцу intro.
    #
    # separator_attr -- имя .ui-виджета визуального разделителя ПЕРЕД этим
    # слотом на канве (нет у title -- перед первым разделителю нечего
    # отделять); display_label -- сырое название раздела для диалогов
    # «Собрать документ»/«Новый вариант» и подтверждения удаления
    # (_GENERATE_DIALOG_SLOT_LABELS/_ADD_VARIANT_DIALOG_TITLE ниже для
    # встроенных трёх всё ещё хранят вручную выверенные формулировки под
    # падеж -- display_label только запасной вариант для случаев, где такой
    # словарь не заведён, и единственный источник для разделов, которые
    # завёл сам оператор, см. _build_custom_section_widgets()).
    #
    # Этот словарь -- ТОЛЬКО стартовый набор (title/intro/appendix1),
    # прописанный в Qt Designer. self._CONSTRUCTOR_SLOTS в __init__
    # становится собственной (per-instance) копией этого словаря -- именно
    # в неё _create_section()/_delete_section() добавляют/убирают
    # пользовательские разделы; общий для класса словарь ниже остаётся
    # нетронутым (иначе один раздел, добавленный в одном окне, утёк бы в
    # любое другое активное окно -- редкий, но реальный сценарий, класс
    # используется не только для конструктора).
    _CONSTRUCTOR_SLOTS = {
        "title": {
            "available_list": "availableBlocksList",
            "add_btn": "addTitleVariantBtn",
            "group_toggle": "titleGroupToggle",
            "group_content": "titleGroupContent",
            "included_list": "includedBlockList",
            "fields_panel": "fieldsPanel",
            "fields_layout": "titleFieldsLayout",
            "empty_hint": "Перетащите титульный лист сюда",
            "section_label": "Реквизиты титульного листа",
            "text_label_attr": "includedBlockTextLabel",
            "chevron_attr": "includedBlockChevron",
            "display_label": "Титульный лист",
        },
        "intro": {
            "available_list": "availableIntroBlocksList",
            "add_btn": "addIntroVariantBtn",
            "group_toggle": "introGroupToggle",
            "group_content": "introGroupContent",
            "included_list": "includedIntroBlockList",
            "fields_panel": "introFieldsPanel",
            "fields_layout": "introFieldsLayout",
            "empty_hint": "Перетащите вводную часть сюда",
            "section_label": "Реквизиты вводной части",
            "text_label_attr": "includedIntroBlockTextLabel",
            "chevron_attr": "includedIntroBlockChevron",
            "separator_attr": "slotSeparator",
            "display_label": "Вводная часть",
        },
        "appendix1": {
            "available_list": "availableAppendix1BlocksList",
            "add_btn": "addAppendix1VariantBtn",
            "group_toggle": "appendix1GroupToggle",
            "group_content": "appendix1GroupContent",
            "included_list": "includedAppendix1BlockList",
            "fields_panel": "appendix1FieldsPanel",
            "fields_layout": "appendix1FieldsLayout",
            "empty_hint": "Перетащите приложение 1 сюда",
            "section_label": "Реквизиты приложения 1",
            "text_label_attr": "includedAppendix1BlockTextLabel",
            "chevron_attr": "includedAppendix1BlockChevron",
            "separator_attr": "slotSeparator2",
            "display_label": "Приложение 1",
        },
    }

    def _slot_widget(self, slot: str, key: str):
        return getattr(self, self._CONSTRUCTOR_SLOTS[slot][key])

    def _slot_config(self, slot: str, key: str):
        return self._CONSTRUCTOR_SLOTS[slot][key]

    def _slot_placeholder_name(self, slot: str, field_id: str) -> str:
        """Имя плейсхолдера конкретного поля в конкретном слоте -- одновременно
        имя Python-атрибута динамического виджета, ключ self.data (см.
        get_form_data()) и буквальный текст "{{ ... }}", вписанный в
        .docx-фрагмент варианта (add_title(), см. generate_title_fragment()/
        generate_intro_fragment()).

        Для "title" -- голое field_id, БЕЗ префикса: так уже устроены
        реальные, вручную доработанные в Word файлы title_*.docx --
        сменить схему значило бы молча сломать плейсхолдеры в уже
        существующих пользовательских файлах.

        Для остальных слотов (intro, appendix1, ...) -- с префиксом
        "{slot}_". Каталог полей (get_all_field_labels()) общий для всех
        слотов -- один и тот же field_id можно вставить в любой из них. Без
        префикса виджеты реквизитов нескольких слотов претендовали бы на
        один и тот же self.<field_id> (setattr в _render_slot_fields()) --
        виджет, созданный позже, молча перезаписывал бы атрибут созданного
        раньше, и значение, введённое в один из них, никогда не попадало бы
        в форму (get_form_data() читает только то, на что сейчас указывает
        атрибут)."""
        return field_id if slot == "title" else f"{slot}_{field_id}"

    def _cross_slot_placeholder_value(self, field_id: str) -> str:
        """Значение того же field_id, уже введённое в ДРУГОМ отрисованном
        слоте -- один и тот же плейсхолдер (например, «заводской №»)
        нередко вставлен сразу в несколько слотов (общий каталог полей, см.
        _slot_placeholder_name()), и вводить его в каждом отдельно неудобно.
        Разовое предзаполнение в момент создания виджета, НЕ живая привязка
        -- дальнейшая правка в одном слоте на другие не влияет (виджеты
        остаются независимыми QPlainTextEdit, ничего не связывает их после
        этого вызова).

        Каждый слот хранит виджет под своим именем атрибута
        (_slot_placeholder_name()) -- проверяем все варианты, не зная
        заранее, в каком слоте плейсхолдер уже заполнен.

        Проверка "attr in self.PLAIN_TEXT_EDIT_NAMES" -- не просто
        getattr(self, attr, None): удаление плейсхолдера из варианта
        (_remove_variant_placeholder()) вычёркивает имя из
        PLAIN_TEXT_EDIT_NAMES, но НЕ трогает сам Python-атрибут self.<attr>
        (setattr() никогда не выставляется заново для убранного поля) --
        без этой проверки getattr() возвращал бы висячую ссылку на уже
        удалённый Qt-виджет (QFormLayout.removeRow() в _render_slot_fields()
        удаляет и сам C++-объект), а вызов .toPlainText() на нём падает
        RuntimeError'ом (см. traceback при добавлении того же поля в
        другой слот после того, как его убрали откуда-то ещё)."""
        for slot in self._CONSTRUCTOR_SLOTS:
            attr = self._slot_placeholder_name(slot, field_id)
            if attr not in self.PLAIN_TEXT_EDIT_NAMES:
                continue
            widget = getattr(self, attr, None)
            if widget is None:
                continue
            text = widget.toPlainText()
            if text:
                return text
        return ""

    def _placeholder_numeric_value(self, field_id: str):
        """Сырое значение field_id, введённое ВРУЧНУЮ в реквизитах (тот же
        поиск по всем слотам, что и в _cross_slot_placeholder_value(),
        включая ту же защиту от висячих ссылок через "attr in
        self.PLAIN_TEXT_EDIT_NAMES"). Не учитывает формулы других полей --
        их обходит formula_engine.evaluate_formula() сам, рекурсивно,
        передавая этот метод как resolve_placeholder только для полей БЕЗ
        формулы (см. _open_formula_editor()/_refresh_computed_fields()).

        None, если поле не найдено, пустое или не является числом (formula_engine
        показывает "—", а не падает) -- в отличие от _cross_slot_placeholder_value(),
        которая возвращает "" (эта нужна для текстового предзаполнения, а
        не для арифметики)."""
        from ..services.formula_engine import parse_formula_number

        for slot in self._CONSTRUCTOR_SLOTS:
            attr = self._slot_placeholder_name(slot, field_id)
            if attr not in self.PLAIN_TEXT_EDIT_NAMES:
                continue
            widget = getattr(self, attr, None)
            if widget is None:
                continue
            text = widget.toPlainText().strip()
            if text:
                value = parse_formula_number(text)
                if value is not None:
                    return value
        return None

    def _refresh_computed_fields(self):
        """Живой пересчёт вычисляемых полей -- подключён к textChanged
        КАЖДОГО обычного (не вычисляемого) поля реквизитов во всех слотах
        (см. _render_slot_fields()), а не только тех, что реально нужны
        какой-то формуле -- заранее не известно, от какого именно поля она
        зависит. Трогает только уже отрисованные вычисляемые виджеты
        (QPlainTextEdit[computed="true"]) -- добавление/удаление самого
        поля или его формулы идёт через _render_slot_fields()/
        _open_formula_editor()."""
        from ..services.formula_engine import evaluate_formula, format_formula_result
        from ..services.title_variants_store import load_field_formulas

        formulas = load_field_formulas()
        for slot in self._CONSTRUCTOR_SLOTS:
            prefix = "" if slot == "title" else f"{slot}_"
            for widget_name in self._dynamic_field_names.get(slot, []):
                field_id = widget_name[len(prefix):] if prefix else widget_name
                formula = formulas.get(field_id)
                if formula is None:
                    continue
                widget = getattr(self, widget_name, None)
                if widget is None:
                    continue
                value = evaluate_formula(formula.get("tokens", []), self._placeholder_numeric_value, formulas)
                # QSignalBlocker -- обязателен: widget тут не всегда уже
                # пересобранный "вычисляемый" (readOnly, без textChanged на
                # этот же метод, см. _render_slot_fields()) -- на выходе из
                # _open_formula_editor() формула сохраняется, ДО того как
                # пересобран слот, которому она реально принадлежит (порядок
                # цикла по self._CONSTRUCTOR_SLOTS): пока не дошла очередь
                # до его пересборки, здесь ещё ЖИВОЙ старый обычный виджет
                # с textChanged.connect(self._refresh_computed_fields) --
                # без блокировки setPlainText() ниже сам вызывал бы этот же
                # метод повторно через textChanged, и так до
                # RecursionError (воспроизведено).
                with QSignalBlocker(widget):
                    widget.setPlainText(format_formula_result(value, formula.get("decimals", 2)))

    def _sync_cross_slot_placeholders(self):
        """Досылает предзаполнение одинаковых плейсхолдеров между слотами
        (_cross_slot_placeholder_value()) для УЖЕ существующих виджетов --
        вызывается после «Открыть проект» (см.
        FileHandler._fill_ui_from_project()), когда все слоты
        восстановлены и заполнены своими сохранёнными значениями, но
        документ мог быть сохранён ДО того, как заработало предзаполнение
        при создании виджета, или заполнен только в части слотов.
        Трогает только пустые поля -- уже заполненное (в т.ч. намеренно
        оставленное пустым другим текстом при сохранении) значение не
        перезаписывает."""
        for slot in self._CONSTRUCTOR_SLOTS:
            prefix = "" if slot == "title" else f"{slot}_"
            for widget_name in self._dynamic_field_names.get(slot, []):
                widget = getattr(self, widget_name, None)
                if widget is None or widget.toPlainText():
                    continue
                field_id = widget_name[len(prefix):] if prefix else widget_name
                prefill = self._cross_slot_placeholder_value(field_id)
                if prefill:
                    widget.setPlainText(prefill)

    def _slots_using_field(self, slot: str, field_id: str) -> list:
        """Список display_label слотов (кроме slot), в реквизиты
        включённого блока которых уже добавлен field_id (см.
        _build_placeholder_menu()) -- статус "уже вставлен" (✓, серым,
        некликабельно) в меню «Вставить плейсхолдер» должен относиться
        ТОЛЬКО к текущему варианту текущего слота (в отличие от общего
        каталога полей -- он один на все слоты), иначе пользователь не
        смог бы вставить в этот слот плейсхолдер, который уже стоит в
        другом. Использование в другом слоте -- отдельная, не блокирующая
        пометка рядом со строкой (см. вызывающую сторону), с указанием
        КОНКРЕТНОГО раздела (или нескольких), а не просто "в другом
        шаблоне"."""
        labels = []
        for other_slot in self._CONSTRUCTOR_SLOTS:
            if other_slot == slot:
                continue
            other_variant_id = self._filled_slot_variant(other_slot)
            if other_variant_id is None:
                continue
            other_variant = self._slot_store(other_slot).get_all_variants().get(other_variant_id)
            if other_variant is not None and field_id in other_variant.subtitle_fields:
                labels.append(self._CONSTRUCTOR_SLOTS[other_slot]["display_label"])
        return labels

    def _init_constructor_slot(self, slot: str):
        """Инициализация одного слота конструктора -- вызывается по разу
        для "title" и "intro" из __init__. Дословно то, что раньше было
        одноразовым кодом прямо в __init__ (см. историю), только вместо
        self.availableBlocksList/self.titleGroupToggle/... -- self._slot_widget(slot, ...),
        и коннекты сигналов принимают slot первым аргументом через
        functools.partial (сигналы сами передают только свои штатные
        аргументы -- pos/expanded/parent,first,last -- slot нужно
        зафиксировать заранее)."""
        available_list = self._slot_widget(slot, "available_list")
        add_btn = self._slot_widget(slot, "add_btn")
        group_toggle = self._slot_widget(slot, "group_toggle")
        included_list = self._slot_widget(slot, "included_list")
        fields_panel = self._slot_widget(slot, "fields_panel")

        # "role" -- динамическое свойство под CONSTRUCTOR_QSS (см. её
        # комментарий): раньше эти виджеты стилизовались перечислением
        # #titleFoo, #introFoo, #appendix1Foo -- набор из ровно трёх имён,
        # захардкоженный в самом QSS. Раздел конструктора больше не
        # фиксированная тройка (см. _create_section()) -- имя объекта
        # каждого нового раздела уникально и заранее не известно QSS,
        # поэтому стилизация переехала на роль (одну из пяти фиксированных
        # значений), выставляемую здесь одинаково для встроенного и
        # добавленного раздела. unpolish()/polish() -- тот же приём, что и
        # у остальных динамических свойств в этом файле (dragOver и т.п.),
        # заставляет QSS перечитать свойство сразу, а не только при
        # следующей естественной перерисовке.
        for widget, role in (
            (group_toggle, "groupToggle"), (add_btn, "addVariantBtn"),
            (available_list, "availableBlocksList"), (included_list, "includedBlockList"),
            (fields_panel, "fieldsPanel"),
        ):
            widget.setProperty("role", role)
            widget.style().unpolish(widget)
            widget.style().polish(widget)
        separator_attr = self._CONSTRUCTOR_SLOTS[slot].get("separator_attr")
        if separator_attr:
            separator = getattr(self, separator_attr)
            separator.setProperty("role", "slotSeparator")
            separator.style().unpolish(separator)
            separator.style().polish(separator)

        available_list.setIconSize(QSize(15, 15))
        self._refresh_available_blocks_list(slot)
        available_list.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        available_list.customContextMenuRequested.connect(
            functools.partial(self._show_available_block_context_menu, slot)
        )
        add_btn.setIcon(icons.icon("plus", "#8e8e93", 13))
        add_btn.setIconSize(QSize(13, 13))
        add_btn.clicked.connect(functools.partial(self._open_add_variant_dialog, slot))

        group_toggle.setIconSize(QSize(30, 13))
        group_toggle.toggled.connect(functools.partial(self._toggle_constructor_group, slot))
        self._toggle_constructor_group(slot, group_toggle.isChecked())

        # ПКМ по заголовку раздела -- «Удалить раздел» целиком (в отличие от
        # ПКМ по карточке варианта в available_list -- та удаляет один
        # вариант, см. _show_available_block_context_menu()). Разрешено для
        # любого раздела, включая встроенные title/intro/appendix1 -- по
        # решению пользователя (полная симметрия с мокапом,
        # docs/design/вводная_часть.html), без особой защиты.
        group_toggle.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        group_toggle.customContextMenuRequested.connect(
            functools.partial(self._show_section_context_menu, slot)
        )

        # В списке ровно один блок на слот (Phase 1, см. _process_block_drop()) --
        # родная линия-подсказка Qt "вставить выше/ниже" вводит в
        # заблуждение при замене (выглядит так, будто можно вставить
        # второй блок рядом), хотя реально всегда остаётся один. Сама
        # вставка above/below при этом всё равно у Qt происходит --
        # обрабатывается в _process_block_drop() независимо от того,
        # видна эта линия или нет.
        included_list.setDropIndicatorShown(False)
        included_list.model().rowsInserted.connect(functools.partial(self._on_block_dropped, slot))

        # Подсветка при перетаскивании варианта из сайдбара сюда -- своя,
        # поверх уже готового нативного DnD у QListWidget (сама вставка
        # блока уже работает и без этого, см. rowsInserted выше). Родные
        # dragEnterEvent/dragLeaveEvent/dropEvent сохраняются и по-прежнему
        # вызываются первыми (orig(event)) -- без этого перестала бы
        # работать сама вставка; подсветка (QSS-свойство dragOver, тот же
        # unpolish/polish-приём, что и у titleTemplateDropZone/
        # titleTemplateDropHint) добавляется поверх, только когда orig()
        # реально принял перетаскивание (event.isAccepted()) -- чужой файл
        # из Finder, например, не подсвечивает область. dragMoveEvent не
        # трогается -- у QAbstractItemView он уже принимает события сам, в
        # отличие от titleTemplateDropZone (голого QWidget без встроенного
        # DnD).
        orig_drag_enter = included_list.dragEnterEvent
        orig_drag_leave = included_list.dragLeaveEvent
        orig_drop_event = included_list.dropEvent

        def _block_list_drag_enter(event, orig=orig_drag_enter, lw=included_list):
            orig(event)
            if event.isAccepted():
                lw.setProperty("dragOver", True)
                lw.style().unpolish(lw)
                lw.style().polish(lw)
                lw.update()

        def _block_list_drag_leave(event, orig=orig_drag_leave, lw=included_list):
            orig(event)
            lw.setProperty("dragOver", False)
            lw.style().unpolish(lw)
            lw.style().polish(lw)
            lw.update()

        def _block_list_drop(event, orig=orig_drop_event, lw=included_list):
            orig(event)
            lw.setProperty("dragOver", False)
            lw.style().unpolish(lw)
            lw.style().polish(lw)
            lw.update()

        included_list.dragEnterEvent = _block_list_drag_enter
        included_list.dragLeaveEvent = _block_list_drag_leave
        included_list.dropEvent = _block_list_drop

        # sizeHint единственного item'а (карточка блока или
        # placeholder-подсказка, см. _process_block_drop()/
        # _show_included_block_placeholder()) фиксирует ширину строки в
        # пикселях под viewport().width() НА МОМЕНТ создания item'а --
        # если список потом станет уже (окно сузили, или появился
        # вертикальный скролл у обёртки mainAreaScrollArea), старая
        # ширина не пересчитывается сама, и правый край строки (крестик
        # «×») обрезается видом списка. Держим sizeHint в ногу с
        # фактической шириной на каждый resize.
        orig_resize_event = included_list.resizeEvent

        def _block_list_resize(event, orig=orig_resize_event, lw=included_list):
            orig(event)
            if lw.count() == 0:
                return
            item = lw.item(0)
            width = lw.viewport().width()
            if width and item.sizeHint().width() != width:
                item.setSizeHint(QSize(width, item.sizeHint().height()))

        included_list.resizeEvent = _block_list_resize

        self._show_included_block_placeholder(slot)

    def _show_section_context_menu(self, slot: str, pos):
        """ПКМ по заголовку раздела -- «Переименовать»/«Удалить раздел» (см.
        _init_constructor_slot()). Отдельное меню от #blockMenu (карточка
        варианта, _show_available_block_context_menu()) -- те же два
        пункта, но другой смысл действия (раздел целиком, не один вариант в
        нём)."""
        menu = QMenu(self)
        rename_action = menu.addAction(icons.icon("edit", "#c7c7cc", 14), "Переименовать раздел")
        delete_action = menu.addAction(icons.icon("trash", "#ff453a", 14), "Удалить раздел")
        group_toggle = self._slot_widget(slot, "group_toggle")
        chosen = menu.exec(group_toggle.mapToGlobal(pos))
        if chosen == rename_action:
            self._rename_section(slot)
        elif chosen == delete_action:
            self._delete_section(slot)

    def _rename_section(self, slot: str):
        """«Переименовать раздел» -- меняет отображаемое название везде,
        где оно показывается (заголовок группы в сайдбаре, строка в
        «Собрать документ», заголовок «Новый вариант: ...», подсказка
        пустой drop-зоны, заголовок «Реквизиты ...»). Разрешено для любого
        раздела, включая встроенные -- та же симметрия, что и у удаления
        (см. _delete_section()).

        empty_hint/section_label и словари _GENERATE_DIALOG_SLOT_LABELS/
        _ADD_VARIANT_DIALOG_TITLE у встроенных трёх переходят на
        нейтральные формулировки через двоеточие (тот же приём, что и у
        раздела, заведённого оператором, см.
        _build_custom_section_widgets()) -- их вручную выверенный под
        падеж текст относился к СТАРОМУ названию и для нового может больше
        не подходить грамматически; откатить обратно к исходным
        формулировкам после переименования уже нельзя (то же самое,
        безвозвратное, что и у остальных переименований в этом файле --
        _rename_variant() тоже не хранит историю)."""
        slot_config = self._CONSTRUCTOR_SLOTS[slot]
        current_label = slot_config["display_label"]
        new_label, ok = QInputDialog.getText(
            self, "Переименовать раздел", "Новое название:", text=current_label,
        )
        new_label = new_label.strip()
        if not ok or not new_label or new_label == current_label:
            return

        slot_config["display_label"] = new_label
        slot_config["empty_hint"] = f"Перетащите сюда: {new_label}"
        slot_config["section_label"] = f"Реквизиты: {new_label}"
        self._GENERATE_DIALOG_SLOT_LABELS.pop(slot, None)
        self._ADD_VARIANT_DIALOG_TITLE.pop(slot, None)

        self._slot_widget(slot, "group_toggle").setText(f"  {new_label}")

        variant_id = self._filled_slot_variant(slot)
        if variant_id is not None:
            self._render_slot_fields(slot, variant_id)
        else:
            self._show_included_block_placeholder(slot)

        if slot not in ("title", "intro", "appendix1"):
            sections = custom_sections_store.load_sections()
            for section in sections:
                if section['id'] == slot:
                    section['label'] = new_label
            custom_sections_store.save_sections(sections)

    def _build_custom_section_widgets(self, slot: str, label: str):
        """Строит виджеты раздела slot программно -- та же иерархия, что
        constructor_window.ui объявляет для intro/appendix1 (QPushButton-
        переключатель + QWidget-содержимое в сайдбаре; QListWidget+
        QGroupBox(QFormLayout)+разделитель на канве), но через код, а не Qt
        Designer -- единственный способ завести раздел, которого не было на
        момент сборки .ui (см. _create_section()/обсуждение задачи).

        Виджеты регистрируются через setattr() под теми же именами, что уже
        использует _CONSTRUCTOR_SLOTS для intro/appendix1
        (available{Cap}BlocksList и т.п.) -- после этого
        _slot_widget()/getattr(self, name) не отличает такой раздел от
        встроенного, и весь остальной код (сборка, генерация, ПКМ-меню,
        drag-and-drop, _init_constructor_slot()) работает без единой правки.

        Формулировки (empty_hint/section_label/...) для падежа непроизвольно
        введённого названия не подобрать -- те же нейтральные формулировки
        через двоеточие, что и в мокапе (docs/design/вводная_часть.html,
        sectionLabels()). Виджеты вставляются последними в
        sidebarLayout/mainAreaLayout -- перед хвостовым spacer'ом каждого
        (спейсер не именован, единственный надёжный способ остаться перед
        ним -- вставка по индексу count()-1, а не по имени)."""
        cap = slot[0].upper() + slot[1:]

        group_toggle = QPushButton(f"  {label}")
        group_toggle.setObjectName(f"{slot}GroupToggle")
        group_toggle.setCheckable(True)
        group_toggle.setChecked(True)
        group_toggle.setFlat(True)
        setattr(self, f"{slot}GroupToggle", group_toggle)

        group_content = QWidget()
        group_content.setObjectName(f"{slot}GroupContent")
        group_content_layout = QVBoxLayout(group_content)
        group_content_layout.setSpacing(2)
        group_content_layout.setContentsMargins(0, 0, 0, 0)
        setattr(self, f"{slot}GroupContent", group_content)

        available_list = QListWidget()
        available_list.setObjectName(f"available{cap}BlocksList")
        available_list.setDragEnabled(True)
        available_list.setFrameShape(QFrame.Shape.NoFrame)
        available_list.setWordWrap(True)
        available_list.setTextElideMode(Qt.TextElideMode.ElideNone)
        available_list.setResizeMode(QListView.ResizeMode.Adjust)
        available_list.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        available_list.setSizePolicy(QSizePolicy.Policy.Preferred, QSizePolicy.Policy.Maximum)
        setattr(self, f"available{cap}BlocksList", available_list)
        group_content_layout.addWidget(available_list)

        add_btn = QPushButton("  Добавить")
        add_btn.setObjectName(f"add{cap}VariantBtn")
        add_btn.setFlat(True)
        add_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        setattr(self, f"add{cap}VariantBtn", add_btn)
        group_content_layout.addWidget(add_btn)

        # count()-1 (перед хвостовым spacer'ом) здесь НЕ годится -- между
        # последней группой и spacer'ом уже лежит addSectionBtn («+
        # Добавить раздел», см. __init__), и count()-1 воткнул бы новую
        # группу МЕЖДУ этой кнопкой и spacer'ом -- визуально новый раздел
        # оказывался бы ПОД кнопкой «Добавить раздел», а не над ней (баг,
        # найденный на реальном запуске: см. обсуждение задачи). indexOf()
        # находит саму кнопку -- новая группа встаёт непосредственно перед
        # ней, независимо от того, сколько разделов уже добавлено раньше.
        insert_at = self.sidebarLayout.indexOf(self.addSectionBtn)
        self.sidebarLayout.insertWidget(insert_at, group_toggle)
        self.sidebarLayout.insertWidget(insert_at + 1, group_content)

        # Разделитель -- только если это не первый раздел вообще (см.
        # _refresh_section_separators(), вызывается сразу после этого метода
        # из _create_section() и пересчитывает видимость у всех разделов
        # заново, а не только у нового).
        separator = QFrame()
        separator.setObjectName(f"{slot}Separator")
        separator.setFrameShape(QFrame.Shape.NoFrame)
        separator.setMinimumSize(QSize(0, 2))
        separator.setMaximumSize(QSize(16777215, 2))
        setattr(self, f"{slot}Separator", separator)

        included_list = QListWidget()
        included_list.setObjectName(f"included{cap}BlockList")
        included_list.setAcceptDrops(True)
        included_list.setDragDropMode(QAbstractItemView.DragDropMode.DragDrop)
        included_list.setDefaultDropAction(Qt.DropAction.MoveAction)
        included_list.setFrameShape(QFrame.Shape.NoFrame)
        included_list.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        included_list.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        included_list.setMaximumHeight(56)
        setattr(self, f"included{cap}BlockList", included_list)

        fields_panel = QGroupBox()
        fields_panel.setObjectName(f"{slot}FieldsPanel")
        fields_panel.setTitle("")
        fields_panel.setVisible(False)
        fields_layout = QFormLayout(fields_panel)
        setattr(self, f"{slot}FieldsPanel", fields_panel)
        setattr(self, f"{slot}FieldsLayout", fields_layout)

        insert_at = self.mainAreaLayout.count() - 1
        self.mainAreaLayout.insertWidget(insert_at, separator)
        self.mainAreaLayout.insertWidget(insert_at + 1, included_list)
        self.mainAreaLayout.insertWidget(insert_at + 2, fields_panel)

        self._CONSTRUCTOR_SLOTS[slot] = {
            "available_list": f"available{cap}BlocksList",
            "add_btn": f"add{cap}VariantBtn",
            "group_toggle": f"{slot}GroupToggle",
            "group_content": f"{slot}GroupContent",
            "included_list": f"included{cap}BlockList",
            "fields_panel": f"{slot}FieldsPanel",
            "fields_layout": f"{slot}FieldsLayout",
            "empty_hint": f"Перетащите сюда: {label}",
            "section_label": f"Реквизиты: {label}",
            "text_label_attr": f"included{cap}BlockTextLabel",
            "chevron_attr": f"included{cap}BlockChevron",
            "separator_attr": f"{slot}Separator",
            "display_label": label,
        }
        self._dynamic_field_names[slot] = []

    def _refresh_section_separators(self):
        """Разделитель раздела виден, только если это НЕ первый по порядку
        раздел в _CONSTRUCTOR_SLOTS -- добавление/удаление раздела может
        выдвинуть в первые тот, что раньше первым не был (у title
        отдельного разделителя нет вовсе -- getattr(..., None) тогда
        просто пропускает раздел)."""
        for i, slot in enumerate(self._CONSTRUCTOR_SLOTS):
            separator_attr = self._CONSTRUCTOR_SLOTS[slot].get("separator_attr")
            if not separator_attr:
                continue
            separator = getattr(self, separator_attr, None)
            if separator is not None:
                separator.setVisible(i > 0)

    def _create_section(self, slot: str, label: str, persist: bool = True):
        """Заводит новый раздел конструктора -- строит виджеты
        (_build_custom_section_widgets()) и инициализирует его тем же общим
        путём, что и три встроенных при старте (_init_constructor_slot()).

        persist=False -- при восстановлении уже сохранённого раздела на
        старте приложения (см. __init__, custom_sections_store.load_sections())
        -- id уже есть в файле, повторно дописывать не нужно."""
        self._build_custom_section_widgets(slot, label)
        self._init_constructor_slot(slot)
        self._refresh_section_separators()
        if persist:
            sections = custom_sections_store.load_sections()
            sections.append({'id': slot, 'label': label})
            custom_sections_store.save_sections(sections)

    def _open_add_section_dialog(self):
        """«+ Добавить раздел» -- заводит раздел целиком (не вариант внутри
        существующего, см. _open_add_variant_dialog()). Название -- любое,
        id -- отдельный счётчик (custom_sections_store.next_section_id()),
        не производная от названия (см. её докстринг)."""
        label, ok = QInputDialog.getText(self, "Новый раздел", "Название раздела:")
        label = label.strip()
        if not ok or not label:
            return
        slot = custom_sections_store.next_section_id()
        self._create_section(slot, label)

    def _delete_section(self, slot: str):
        """Удаляет раздел целиком -- разрешено для любого, включая
        встроенные title/intro/appendix1 (см. _show_section_context_menu()
        и обсуждение задачи). Варианты раздела стираются из JSON безвозвратно
        (store.save_variants([])) -- их .docx-фрагменты на диске НЕ
        удаляются, та же осторожность, что и в _delete_variant() (не она их
        удаляет, только запись в каталоге).

        Удаление встроенного раздела действует до перезапуска приложения --
        сам раздел объявлен в constructor_window.ui и на следующем старте
        появится снова (пустым, варианты уже стёрты); custom_sections_store
        правится только для пользовательских разделов, у встроенных трёх
        такой записи и не было."""
        label = self._CONSTRUCTOR_SLOTS[slot].get("display_label", slot)
        answer = QMessageBox.question(
            self, "Удалить раздел",
            f"Удалить раздел «{label}»? Все его варианты и введённые в них данные будут удалены безвозвратно.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return

        self._slot_store(slot).save_variants([])

        for widget_name in self._dynamic_field_names.get(slot, []):
            if widget_name in self.PLAIN_TEXT_EDIT_NAMES:
                self.PLAIN_TEXT_EDIT_NAMES.remove(widget_name)
        self._dynamic_field_names.pop(slot, None)

        group_toggle = self._slot_widget(slot, "group_toggle")
        group_content = self._slot_widget(slot, "group_content")
        self.sidebarLayout.removeWidget(group_toggle)
        self.sidebarLayout.removeWidget(group_content)
        group_toggle.deleteLater()
        group_content.deleteLater()

        included_list = self._slot_widget(slot, "included_list")
        fields_panel = self._slot_widget(slot, "fields_panel")
        self.mainAreaLayout.removeWidget(included_list)
        self.mainAreaLayout.removeWidget(fields_panel)
        included_list.deleteLater()
        fields_panel.deleteLater()

        separator_attr = self._CONSTRUCTOR_SLOTS[slot].get("separator_attr")
        if separator_attr:
            separator = getattr(self, separator_attr, None)
            if separator is not None:
                self.mainAreaLayout.removeWidget(separator)
                separator.deleteLater()

        del self._CONSTRUCTOR_SLOTS[slot]
        self._refresh_section_separators()
        self._update_crumb()
        self.pushButt_generateWord.setEnabled(
            any(self._filled_slot_variant(s) is not None for s in self._CONSTRUCTOR_SLOTS)
        )

        if slot not in ("title", "intro", "appendix1"):
            sections = [s for s in custom_sections_store.load_sections() if s['id'] != slot]
            custom_sections_store.save_sections(sections)

    def init_file_handler(self):
        """Инициализация FileHandler для импорта/экспорта."""
        from .file_handler import FileHandler
        self.file_handler = FileHandler(self)

    def init_widgets(self):
        """Автоматически инициализирует все виджеты из UI"""
        # Инициализация QPlainTextEdit
        for name in self.PLAIN_TEXT_EDIT_NAMES:
            widget = self.findChild(QPlainTextEdit, name)
            if widget is None:
                raise ValueError(f"Не найден QPlainTextEdit с именем {name}")
            setattr(self, name, widget)

        # Инициализация QComboBox
        for name in self.COMBO_BOX_NAMES:
            widget = self.findChild(QComboBox, name)
            if widget is None:
                raise ValueError(f"Не найден QComboBox с именем {name}")
            setattr(self, name, widget)

        # Инициализация QPushButton
        for name in self.BUTTON_NAMES:
            widget = self.findChild(QPushButton, name)
            if widget is None:
                raise ValueError(f"Не найден QPushButton с именем {name}")
            setattr(self, name, widget)

        # Инициализация QSpinBox
        for name in self.SPIN_BOX_NAMES:
            widget = self.findChild(QSpinBox, name)
            if widget is None:
                raise ValueError(f"Не найден QSpinBox с именем {name}")
            setattr(self, name, widget)

        # Инициализация QDateEdit
        for name in self.DATE_EDIT_NAMES:
            widget = self.findChild(QDateEdit, name)
            if widget is None:
                raise ValueError(f"Не найден QDateEdit с именем {name}")
            setattr(self, name, widget)

        # Инициализация QTableWidget.
        for name in self.TABLE_WIDGET:
            widget = self.findChild(QTableWidget, name)
            if widget is None:
                raise ValueError(f"Не найден QTableWidget с именем {name}")
            setattr(self, name, widget)

    def get_form_data(self) -> Dict[str, Union[str, float]]:
        """Возвращает все данные формы в виде словаря"""

        # Получаем текст из всех QPlainTextEdit
        for name in self.PLAIN_TEXT_EDIT_NAMES:
            widget = getattr(self, name)
            self.data[name] = widget.toPlainText()

        # Получаем текущий текст из QComboBox
        for name in self.COMBO_BOX_NAMES:
            widget = getattr(self, name)
            self.data[name] = widget.currentText()

        # Комбобоксы выбора специалиста (только трубопровод) -- хранят
        # индекс строки table_specialists (currentData()), а не текст,
        # поэтому не в COMBO_BOX_NAMES. Сохраняем отдельно, иначе выбор
        # не переживает "Сохранить проект" -> "Открыть проект" и после
        # открытия откатывается на первую строку, см.
        # FileHandler._fill_ui_from_project()/_refresh_program_specialist_combo().
        for name in getattr(self, "SPECIALIST_COMBO_NAMES", []):
            widget = getattr(self, name)
            self.data[name] = widget.currentData()

        # Получаем текст из QSpinBox.
        for name in self.SPIN_BOX_NAMES:
            widget = getattr(self, name)
            self.data[name] = widget.value()

        # Получаем даты из QDateEdit.
        for name in self.DATE_EDIT_NAMES:
            widget = getattr(self, name)
            date = widget.date()
            if date.isValid():
                locale = QLocale('ru_RU')
                # Суффикс "г." -- только для трубопровода: в его шаблоне
                # везде, кроме final_deadline_date, дата стоит "голой"
                # (Дата проведения контроля {{ ... }}, титул {{ report_date
                # }}) -- эталонный документ TD_720291 везде добавляет "г.".
                # final_deadline_date -- исключение: за ним в шаблоне уже
                # стоит своё литеральное "года" (8.2, 8.4.1), суффикс
                # задвоил бы "г. года". Баллонный шаблон не проверялся на
                # такое же дублирование, поэтому его не трогаем.
                is_pipeline = self.equipment_type.id == "pipeline"
                suffix = " г." if (is_pipeline and name != "final_deadline_date") else ""
                if name == "report_date":
                    # Дата отчёта на титульном листе -- день в кавычках-ёлочках
                    # по канцелярской традиции: «15» ноября 2024. Текст в
                    # одинарных кавычках в формате QLocale выводится буквально.
                    self.data[name] = locale.toString(date, "'«'dd'»' MMMM yyyy") + suffix
                else:
                    # Формат: dd MMMM yyyy (пробелы, месяц в родительном падеже на русском)
                    self.data[name] = locale.toString(date, 'dd MMMM yyyy') + suffix
            else:
                self.data[name] = ""

        # Получаем текст из QTableWidget. Ячейка может быть обычным
        # QTableWidgetItem или виджетом (например, QComboBox -- см.
        # table_segments), поэтому при отсутствии item проверяем cellWidget.
        for name in self.TABLE_WIDGET:
            widget = getattr(self, name)
            table_data = []
            for row in range(widget.rowCount()):
                row_data = []
                for col in range(widget.columnCount()):
                    item = widget.item(row, col)
                    if item is not None:
                        row_data.append(item.text())
                    else:
                        cell_widget = widget.cellWidget(row, col)
                        if isinstance(cell_widget, QComboBox):
                            row_data.append(cell_widget.currentText())
                        else:
                            row_data.append("")
                table_data.append(row_data)
            self.data[name] = table_data

        # Табличные поля конструктора (field_tables, см.
        # table_editor_dialog.py) -- значение виджета всегда "" (см.
        # _render_slot_fields(), ветка `table is not None`), сюда вместо
        # него подставляется маркер (_table_field_marker()), который
        # _splice_table_placeholders() ПОСЛЕ tpl.render() меняет на
        # настоящую .docx-таблицу. Поиск по всем слотам -- тот же приём,
        # что и в _cross_slot_placeholder_value()/_placeholder_numeric_value()
        # (widget_name разный в разных слотах, кроме title). Только для
        # equipment_type == "constructor" -- у баллонов/трубопровода
        # field_tables не существует вовсе, отдельный импорт стора тут не
        # нужен без реальной надобности.
        if self.equipment_type.id == "constructor":
            from ..services.title_variants_store import load_field_tables

            tables = load_field_tables()
            for field_id in tables:
                for slot in self._CONSTRUCTOR_SLOTS:
                    widget_name = self._slot_placeholder_name(slot, field_id)
                    if widget_name in self.PLAIN_TEXT_EDIT_NAMES:
                        self.data[widget_name] = _table_field_marker(field_id)

        return self.data

    def calculate(self):
        """Обработчик нажатия кнопки генерации Word с улучшенной обработкой ошибок"""
        if self.equipment_type.id == "constructor":
            # Конструктор собирает документ из перетащенных блоков, а не из
            # одного большого статического шаблона -- рендер устроен
            # совсем иначе (см. _calculate_constructor()), чем у
            # balloon/pipeline ниже, поэтому отдельная ветка, а не третий
            # elif в теле этого метода. Перед самой сборкой -- диалог
            # выбора шаблонов (_open_generate_dialog()): пользователь
            # решает, из чего собирать, а не только из того, что уже
            # лежит в included_list слотов.
            self._open_generate_dialog()
            return

        missing = [step for step in self.STEP_ORDER if step not in self._completed_steps]
        if missing:
            missing_labels = ", ".join(self.STEP_LABELS[step] for step in missing)
            self.show_message(
                "Не выполнены обязательные шаги",
                f"Перед генерацией документа выполните: {missing_labels}.",
                QMessageBox.Icon.Warning,
            )
            return

        try:
            # 1. Проверка заполненности полей
            if not all(
                getattr(self, name).toPlainText().strip()
                for name in self.equipment_type.required_fields
            ):
                raise ValueError("Не все обязательные поля заполнены")

            # 2. Получаем данные формы
            if self.equipment_type.id == "pipeline":
                # specialists -- список словарей для {% for %} в шаблоне;
                # table_specialists (список списков "как есть") тоже
                # попадёт в form_data ниже через get_form_data() -- нужен
                # для сохранения/восстановления проекта.
                # Колонки таблицы: 0 -- Должность, 1 -- ФИО, 2 -- Удостоверение
                # (полное, только для Таблицы 2), 3 -- Удостоверение (кратко,
                # для подписей после каждого из 9 приложений, см. ниже).
                self.data["specialists"] = self._table_to_dicts(
                    self.table_specialists,
                    ["position", "name", "cert_number", "cert_number_short"]
                )
                # name_initials -- "И. О. Фамилия" для таблицы подписи после
                # раздела 8 (Таблица 2 продолжает использовать полное "name").
                # employee_id -- сотрудник справочника «Сотрудники», из
                # которого взята эта строка (см. _add_specialist_row()),
                # нужен ниже для подстановки клише (InlineImage) в каждое из
                # 9 мест подписи, см. _specialist_kleishe_image(). Если связи
                # нет (строка вписана вручную или сохранена в project.json
                # до появления привязки к справочнику) -- пробуем найти
                # сотрудника по точному совпадению ФИО, иначе документы,
                # набранные раньше этой доработки, никогда не получили бы
                # клише, хотя вписанное ФИО совпадает с сотрудником в
                # справочнике буква в букву.
                roster = load_employees()
                for row, specialist in enumerate(self.data["specialists"]):
                    specialist["name_initials"] = format_fio_initials(specialist["name"])
                    employee_id = (
                        self._specialist_employee_ids[row]
                        if row < len(self._specialist_employee_ids) else None
                    )
                    if not employee_id:
                        employee_id = find_employee_id_by_name(roster, specialist["name"])
                    specialist["employee_id"] = employee_id

                # Таблица подписи после раздела 8 -- подписывает только один
                # (первый по списку) специалист, не цикл по всем; см.
                # template_schema.py. Явные плоские поля вместо {{
                # specialists[0]... }} в шаблоне -- Jinja упал бы на пустом
                # списке, а так ошибка ловится здесь с понятным текстом.
                if not self.data["specialists"]:
                    raise ValueError(
                        "Добавьте хотя бы одного специалиста в Таблице 2 "
                        "(1.3 Сведения о специалистах) -- он подписывает Отчёт"
                    )
                lead_specialist = self.data["specialists"][0]
                self.data["lead_specialist_position"] = lead_specialist["position"]
                self.data["lead_specialist_name_initials"] = lead_specialist["name_initials"]
                self.data["lead_specialist_cert_number"] = lead_specialist["cert_number_short"]
                self.data["lead_specialist_employee_id"] = lead_specialist["employee_id"]

                # "Программу составил" (Приложение 1) -- специалиста выбирает
                # оператор через program_specialist (индекс строки Таблицы 2),
                # а не жёстко первая строка, как у lead_specialist_* выше.
                programm_specialist_idx = self.program_specialist.currentData()
                if (
                    programm_specialist_idx is None
                    or programm_specialist_idx >= len(self.data["specialists"])
                ):
                    raise ValueError(
                        "Выберите специалиста в поле «Программу составил» "
                        "(Приложение 1) -- источник вариантов: Таблица 2 (1.3)"
                    )
                programm_specialist = self.data["specialists"][programm_specialist_idx]
                self.data["programm_specialist_position"] = programm_specialist["position"]
                self.data["programm_specialist_name_initials"] = programm_specialist["name_initials"]
                self.data["programm_specialist_cert_number"] = programm_specialist["cert_number_short"]
                self.data["programm_specialist_employee_id"] = programm_specialist["employee_id"]

                # "Анализ документации провёл" (Приложение 2) -- тот же
                # паттерн, что и "Программу составил" выше.
                act2_specialist_idx = self.act2_specialist.currentData()
                if (
                    act2_specialist_idx is None
                    or act2_specialist_idx >= len(self.data["specialists"])
                ):
                    raise ValueError(
                        "Выберите специалиста в поле «Анализ документации провёл» "
                        "(Приложение 2) -- источник вариантов: Таблица 2 (1.3)"
                    )
                act2_specialist = self.data["specialists"][act2_specialist_idx]
                self.data["act2_specialist_position"] = act2_specialist["position"]
                self.data["act2_specialist_name_initials"] = act2_specialist["name_initials"]
                self.data["act2_specialist_cert_number"] = act2_specialist["cert_number_short"]
                self.data["act2_specialist_employee_id"] = act2_specialist["employee_id"]

                # "Контроль провёл" (Приложение 3) -- тот же паттерн.
                vik_specialist_idx = self.vik_specialist.currentData()
                if (
                    vik_specialist_idx is None
                    or vik_specialist_idx >= len(self.data["specialists"])
                ):
                    raise ValueError(
                        "Выберите специалиста в поле «Контроль провёл» "
                        "(Приложение 3) -- источник вариантов: Таблица 2 (1.3)"
                    )
                vik_specialist = self.data["specialists"][vik_specialist_idx]
                self.data["vik_specialist_position"] = vik_specialist["position"]
                self.data["vik_specialist_name_initials"] = vik_specialist["name_initials"]
                self.data["vik_specialist_cert_number"] = vik_specialist["cert_number_short"]
                self.data["vik_specialist_employee_id"] = vik_specialist["employee_id"]

                # "Измерение провёл" (Приложение 4) -- тот же паттерн.
                thick_specialist_idx = self.thick_specialist.currentData()
                if (
                    thick_specialist_idx is None
                    or thick_specialist_idx >= len(self.data["specialists"])
                ):
                    raise ValueError(
                        "Выберите специалиста в поле «Измерение провёл» "
                        "(Приложение 4) -- источник вариантов: Таблица 2 (1.3)"
                    )
                thick_specialist = self.data["specialists"][thick_specialist_idx]
                self.data["thick_specialist_position"] = thick_specialist["position"]
                self.data["thick_specialist_name_initials"] = thick_specialist["name_initials"]
                self.data["thick_specialist_cert_number"] = thick_specialist["cert_number_short"]
                self.data["thick_specialist_employee_id"] = thick_specialist["employee_id"]

                # "Измерение провёл" (Приложение 5) -- тот же паттерн.
                uzk_specialist_idx = self.uzk_specialist.currentData()
                if (
                    uzk_specialist_idx is None
                    or uzk_specialist_idx >= len(self.data["specialists"])
                ):
                    raise ValueError(
                        "Выберите специалиста в поле «Измерение провёл» "
                        "(Приложение 5) -- источник вариантов: Таблица 2 (1.3)"
                    )
                uzk_specialist = self.data["specialists"][uzk_specialist_idx]
                self.data["uzk_specialist_position"] = uzk_specialist["position"]
                self.data["uzk_specialist_name_initials"] = uzk_specialist["name_initials"]
                self.data["uzk_specialist_cert_number"] = uzk_specialist["cert_number_short"]
                self.data["uzk_specialist_employee_id"] = uzk_specialist["employee_id"]

                # "Расчёт выполнил" (Приложение 6) -- тот же паттерн.
                calc_specialist_idx = self.calc_specialist.currentData()
                if (
                    calc_specialist_idx is None
                    or calc_specialist_idx >= len(self.data["specialists"])
                ):
                    raise ValueError(
                        "Выберите специалиста в поле «Расчёт выполнил» "
                        "(Приложение 6) -- источник вариантов: Таблица 2 (1.3)"
                    )
                calc_specialist = self.data["specialists"][calc_specialist_idx]
                self.data["calc_specialist_position"] = calc_specialist["position"]
                self.data["calc_specialist_name_initials"] = calc_specialist["name_initials"]
                self.data["calc_specialist_cert_number"] = calc_specialist["cert_number_short"]
                self.data["calc_specialist_employee_id"] = calc_specialist["employee_id"]

                # "Контроль выполнил" (Приложение 8) -- тот же паттерн.
                pnevmo_specialist_idx = self.pnevmo_specialist.currentData()
                if (
                    pnevmo_specialist_idx is None
                    or pnevmo_specialist_idx >= len(self.data["specialists"])
                ):
                    raise ValueError(
                        "Выберите специалиста в поле «Контроль выполнил» "
                        "(Приложение 8) -- источник вариантов: Таблица 2 (1.3)"
                    )
                pnevmo_specialist = self.data["specialists"][pnevmo_specialist_idx]
                self.data["pnevmo_specialist_position"] = pnevmo_specialist["position"]
                self.data["pnevmo_specialist_name_initials"] = pnevmo_specialist["name_initials"]
                self.data["pnevmo_specialist_cert_number"] = pnevmo_specialist["cert_number_short"]
                self.data["pnevmo_specialist_employee_id"] = pnevmo_specialist["employee_id"]

                # "Заключение составил" (Приложение 9) -- тот же паттерн.
                ae_zakl_specialist_idx = self.ae_zakl_specialist.currentData()
                if (
                    ae_zakl_specialist_idx is None
                    or ae_zakl_specialist_idx >= len(self.data["specialists"])
                ):
                    raise ValueError(
                        "Выберите специалиста в поле «Заключение составил» "
                        "(Приложение 9) -- источник вариантов: Таблица 2 (1.3)"
                    )
                ae_zakl_specialist = self.data["specialists"][ae_zakl_specialist_idx]
                self.data["ae_zakl_specialist_position"] = ae_zakl_specialist["position"]
                self.data["ae_zakl_specialist_name_initials"] = ae_zakl_specialist["name_initials"]
                self.data["ae_zakl_specialist_cert_number"] = ae_zakl_specialist["cert_number_short"]
                self.data["ae_zakl_specialist_employee_id"] = ae_zakl_specialist["employee_id"]

                # Таблица 1 (Приложение 8, п.11) -- список словарей под
                # {%tr for %} в шаблоне; колонки: 0 -- ПАЭ №, 1 -- Нагрузка,
                # 2 -- Класс источника (см. table_pnevmo_ae, AE_CLASS_TYPES).
                self.data["pnevmo_ae_results"] = self._table_to_dicts(
                    self.table_pnevmo_ae, ["paje_num", "nagruzka", "klass"]
                )

                # 8.2 -- "5 (пять) лет": число прописью в скобках рядом с цифрой.
                years_allowed_text = self.final_years_allowed.toPlainText().strip()
                try:
                    years_allowed_int = int(years_allowed_text)
                except ValueError:
                    raise ValueError(
                        f"«Разрешённый срок дальнейшей эксплуатации» должен быть "
                        f"целым числом лет: {years_allowed_text!r}"
                    )
                self.data["final_years_allowed_words"] = number_to_words_ru(years_allowed_int)

                # Приложение 8, п.1 -- дата контроля числом (15.08.2024), а не
                # текстом (pnevmo_date остаётся текстовым везде, где уже
                # используется, включая Приложение 9 -- см. отчёт о сравнении).
                self.data["pnevmo_date_numeric"] = self.pnevmo_date.date().toString("dd.MM.yyyy")
            # Копия, а не self.data напрямую: ниже в form_data пишутся
            # InlineImage (держат ссылку на DocxTemplate/дерево .docx) и
            # RichText (_apply_line_breaks) -- одноразовые объекты только
            # для рендера. self.get_form_data() возвращает self.data по
            # ссылке; без copy() эти объекты оседали бы в self.data
            # навсегда и следующее "Сохранить проект" падало бы с
            # RecursionError на dataclasses.asdict() -> copy.deepcopy()
            # лежащего в них lxml-дерева документа.
            form_data = dict(self.get_form_data())
            if self.equipment_type.id == "pipeline":
                self._apply_line_breaks(form_data)
            print("Данные для Word:", form_data)

            # 3. Проверяем наличие шаблона (используем config.py)
            from ..config import find_template, OUTPUT_DIR
            template_path = find_template(self.equipment_type.id)

            # 4. Загружаем и заполняем шаблон
            doc = DocxTemplate(template_path)

            if self.equipment_type.id == "pipeline":
                # Приложение 7 -- Схема НК: картинка не обязательна, при
                # отсутствии на месте {{ nk_scheme_image }} остаётся пусто.
                nk_scheme_filename = self.data.get("nk_scheme_filename")
                if nk_scheme_filename:
                    form_data["nk_scheme_image"] = InlineImage(
                        doc, str(NK_SCHEME_DIR / nk_scheme_filename), width=Mm(150)
                    )
                else:
                    form_data["nk_scheme_image"] = ""

                # Приложение 8 -- Рисунок 1 (график нагружения): тот же
                # приём, картинка не обязательна.
                pnevmo_graph_filename = self.data.get("pnevmo_graph_filename")
                if pnevmo_graph_filename:
                    form_data["pnevmo_graph_image"] = InlineImage(
                        doc, str(PNEVMO_GRAPH_DIR / pnevmo_graph_filename), width=Mm(150)
                    )
                else:
                    form_data["pnevmo_graph_image"] = ""

                # Клише специалиста (картинка из справочника «Сотрудники»,
                # см. EmployeesTabController) в каждое из 9 мест подписи --
                # тот же приём, не обязательно (сотрудник мог не загрузить
                # клише, а специалист может быть не найден в справочнике
                # даже через find_employee_id_by_name() выше). roster --
                # тот же справочник, что уже читали выше для employee_id.
                # kleishe_paths_used -- для _float_kleishe_drawings_behind_
                # text() ниже, после рендера.
                kleishe_paths_used = set()
                for prefix in self.KLEISHE_ROLE_PREFIXES:
                    form_data[f"{prefix}_kleishe"] = self._specialist_kleishe_image(
                        doc, roster, self.data.get(f"{prefix}_employee_id"), kleishe_paths_used
                    )

            doc.render(form_data)

            if self.equipment_type.id == "pipeline":
                # Клише должно быть "за текстом" (Word: Обтекание текстом ->
                # За текстом), а не обычной инлайн-картинкой, растягивающей
                # ячейку -- docxtpl не умеет вставлять плавающие картинки
                # напрямую, поэтому уже отрисованный XML патчится постфактум,
                # см. src/services/docx_layout.py. Размер (<wp:extent>) при
                # этом не меняется -- переносится тот же узел, а не пересоздаётся.
                self._float_kleishe_drawings_behind_text(doc, kleishe_paths_used)

            # 5. Генерируем имя файла
            if self.equipment_type.id == "balloon":
                output_filename = (
                    f"закл_{self.zakl_number.toPlainText().strip()}_"
                    f"рег-{self.reg_number.toPlainText().strip()}_"
                    f"р-{self.p_rab.toPlainText().strip()}_"
                    f"{self.rab_sreda.currentText()}_"
                    f"кбХиммаш_{self.amount.value()}шт.docx"
                )
            else:
                # Реальные номера отчётов часто содержат "/" (напр.
                # "681/24") -- недопустимо в имени файла, заменяем на "-".
                safe_report_number = self.report_number.toPlainText().strip().replace("/", "-")
                output_filename = (
                    f"отчёт_{safe_report_number}_"
                    f"рег-{self.reg_number.toPlainText().strip()}_"
                    f"р-{self.p_rab_mpa.toPlainText().strip()}_"
                    f"{self.work_medium.currentText()}.docx"
                )

            # 6. Место сохранения и имя файла -- оператор может изменить
            # и то, и другое; output_filename выше только подсказка по
            # умолчанию.
            output_dir = str(OUTPUT_DIR)
            os.makedirs(output_dir, exist_ok=True)
            default_path = os.path.join(output_dir, output_filename)

            output_path, _ = QFileDialog.getSaveFileName(
                self, "Сохранить документ", default_path, "Документы Word (*.docx)"
            )
            if not output_path:
                return  # отменено пользователем -- не ошибка

            doc.save(output_path)

            # 7. Уведомление об успехе
            self.show_message(
                "Готово!",
                f"Документ успешно сохранён:\n{output_path}",
                QMessageBox.Icon.Information  # PyQt6 использует QMessageBox.Icon
            )
            print(f"Документ успешно сохранён: {output_path}")

        except ValueError as ve:
            self.show_message("Ошибка ввода", str(ve), QMessageBox.Icon.Warning)
        except FileNotFoundError as fe:
            self.show_message("Файл не найден", str(fe), QMessageBox.Icon.Critical)
        except PermissionError:
            self.show_message(
                "Ошибка доступа",
                "Нет прав для записи в указанную папку",
                QMessageBox.Icon.Critical
            )
        except Exception as e:
            self.show_message(
                "Ошибка генерации",
                f"Неизвестная ошибка: {str(e)}",
                QMessageBox.Icon.Critical
            )

    def _filled_slot_variant(self, slot: str):
        """id варианта, лежащего в слоте slot, или None, если слот пуст
        (includedBlockList/includedIntroBlockList хранят либо ровно один
        реальный item, либо только плейсхолдер INCLUDED_BLOCK_PLACEHOLDER,
        см. _show_included_block_placeholder())."""
        included_list = self._slot_widget(slot, "included_list")
        if included_list.count() == 0:
            return None
        variant_id = included_list.item(0).data(Qt.ItemDataRole.UserRole)
        return None if variant_id == self.INCLUDED_BLOCK_PLACEHOLDER else variant_id

    def _restore_included_variant(self, slot: str, variant_id: str):
        """Программно включает variant_id в слот -- кладёт в included_list
        слота item с тем же UserRole=variant_id, что при обычном
        перетаскивании из available_list, и прогоняет его через тот же
        _process_block_drop(), который иначе запускается только сигналом
        rowsInserted настоящего drag-and-drop. Два вызывающих: «Открыть
        проект» (см. FileHandler._create_project()/_fill_ui_from_project())
        и выбор шаблона в диалоге «Собрать документ» (_open_generate_dialog()).
        Если вариант с этим id с тех пор удалили из каталога (JSON правили
        руками, или «Удалить» после сохранения проекта) -- молча ничего не
        делает, слот остаётся в прежнем состоянии, как и было бы после
        обычного drop чужого/несуществующего id (см. _process_block_drop())."""
        get_all_variants = self._slot_store(slot).get_all_variants
        variants = get_all_variants()
        if variant_id not in variants:
            return

        block_list = self._slot_widget(slot, "included_list")
        block_list.clear()
        item = QListWidgetItem(variants[variant_id].document_title)
        item.setIcon(icons.icon("file-text", "#0a84ff", 15))
        item.setData(Qt.ItemDataRole.UserRole, variant_id)
        block_list.addItem(item)
        self._process_block_drop(slot)

    _GENERATE_DIALOG_SLOT_LABELS = {
        "title": "Титульный лист", "intro": "Вводная часть", "appendix1": "Приложение 1",
    }

    def _open_generate_dialog(self):
        """«Собрать документ» -- перед самим рендером (_calculate_constructor())
        даёт явно выбрать, какой вариант каждого слота пойдёт в документ, а
        не молча берёт то, что сейчас лежит в included_list (пользователь
        мог захотеть собрать другую комбинацию без лишнего drag-and-drop
        по канвасу). По умолчанию в каждом выпадающем списке выбран
        вариант, уже включённый в этот слот (или «— не включать —», если
        слот пуст) -- открыть диалог и просто нажать OK равносильно
        нынешнему поведению без диалога.

        Выбор в диалоге ПРИМЕНЯЕТСЯ к included_list слотов (через
        _restore_included_variant()/_clear_included_block() -- те же
        функции, что и обычный drag-and-drop/восстановление проекта),
        поэтому после закрытия диалога канвас (карточки, реквизиты)
        отражает именно то, что попадёт в документ, а не расходится с
        ним -- следующее открытие диалога снова покажет актуальный
        выбор."""
        dialog = QDialog(self)
        dialog.setWindowTitle("Собрать документ")
        layout = QVBoxLayout(dialog)
        layout.addWidget(QLabel("Выберите, из каких шаблонов собрать документ:"))

        combos = {}
        for slot in self._CONSTRUCTOR_SLOTS:
            get_all_variants = self._slot_store(slot).get_all_variants
            variants = get_all_variants()
            current_variant_id = self._filled_slot_variant(slot)

            display_label = self._GENERATE_DIALOG_SLOT_LABELS.get(slot) or self._CONSTRUCTOR_SLOTS[slot]["display_label"]
            layout.addWidget(QLabel(display_label))
            combo = QComboBox()
            combo.addItem("— не включать —", None)
            selected_index = 0
            for i, (variant_id, config) in enumerate(variants.items(), start=1):
                combo.addItem(config.document_title, variant_id)
                if variant_id == current_variant_id:
                    selected_index = i
            combo.setCurrentIndex(selected_index)
            layout.addWidget(combo)
            combos[slot] = combo

        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)

        if dialog.exec() != QDialog.DialogCode.Accepted:
            return

        for slot, combo in combos.items():
            selected_variant_id = combo.currentData()
            if selected_variant_id == self._filled_slot_variant(slot):
                continue
            if selected_variant_id is None:
                self._clear_included_block(slot)
            else:
                self._restore_included_variant(slot, selected_variant_id)

        self._calculate_constructor()

    @staticmethod
    def _append_docx_body(target_doc, source_doc):
        """Переносит содержимое source_doc в конец target_doc -- ПЕРЕД его
        финальным sectPr (секция -- настройки страницы, поля, рамка --
        остаётся от target_doc, у source_doc своя секция отбрасывается).
        Используется для склейки заголовок+вводная часть в один файл (см.
        _calculate_constructor()) без сторонних зависимостей (docxcompose и
        т.п. в проекте нет) -- прямая работа с телом документа через
        python-docx OXML, тот же уровень API, что уже использует
        template_generator.py."""
        target_body = target_doc.element.body
        sect_pr = target_body.find(qn('w:sectPr'))
        for child in list(source_doc.element.body):
            if child.tag == qn('w:sectPr'):
                continue
            if sect_pr is not None:
                sect_pr.addprevious(child)
            else:
                target_body.append(child)

    def _build_table_element(self, table: Dict):
        """Строит НАСТОЯЩУЮ .docx-таблицу из содержимого редактора таблиц
        (field_tables[field_id] -- {"has_header":.., "rows": [[токены
        ячейки, ...], ...]}, см. table_editor_dialog.py) и возвращает её
        как независимый OXML-элемент <w:tbl>, готовый быть вставленным в
        ЛЮБОЙ документ (_splice_table_placeholders() ниже).

        Таблица строится в отдельном, ни с чем не связанном python-docx
        Document() -- ЕГО собственный table.style = "Table Grid" работает
        так же надёжно, как и в template_generator.py (тот же built-in
        стиль Word, доступный в любом .docx без явного объявления в
        styles.xml). copy.deepcopy() -- обязателен: без него element
        остаётся частью дерева scratch-документа, а вставка ЧУЖОГО lxml-
        элемента в другое дерево (addprevious() в _splice_table_placeholders())
        молча вырезает его из исходного дерева, не копируя, что тут не
        имеет значения (scratch выбрасывается), но именно эта копия и
        нужна, чтобы вставка была безопасна.

        Плейсхолдер-токен ячейки резолвится в ТЕКУЩИЙ текст того же поля
        через _cross_slot_placeholder_value() -- ту же подстановку, что
        уже применяется для предзаполнения полей между слотами; для
        вычисляемого поля это уже готовый посчитанный текст (см.
        _render_slot_fields()). Ссылка на ДРУГОЕ табличное поле внутри
        ячейки не разворачивается рекурсивно (вернёт "") -- вложенные
        таблицы вне охвата, тот же принцип, что и запрет ссылки таблицы на
        саму себя в TableEditorDialog."""
        from docx import Document as _ScratchDocument

        rows = table.get("rows", [])
        n_rows = len(rows)
        n_cols = len(rows[0]) if rows else 0
        scratch = _ScratchDocument()
        if not n_rows or not n_cols:
            # Пустая таблица (сохранена без единой заполненной ячейки, но
            # хотя бы с одной строкой/столбцом -- TableEditorDialog не
            # позволяет меньше) -- всё равно рисуем видимую сетку размера
            # rows×cols, а не молча ничего не вставляем.
            n_rows, n_cols = max(n_rows, 1), max(n_cols, 1)
        doc_table = scratch.add_table(rows=n_rows, cols=n_cols)
        doc_table.style = "Table Grid"
        has_header = bool(table.get("has_header"))
        for r, row in enumerate(rows):
            for c, cell_tokens in enumerate(row):
                text = "".join(
                    token.get("value", "") if token.get("type") == "text"
                    else self._cross_slot_placeholder_value(token.get("id", ""))
                    for token in cell_tokens
                )
                cell = doc_table.cell(r, c)
                cell.text = text
                if has_header and r == 0:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
        return copy.deepcopy(doc_table._tbl)

    def _splice_table_placeholders(self, doc):
        """Меняет маркеры табличных полей (_table_field_marker(), см.
        get_form_data()) в уже отрендеренном doc на настоящие .docx-
        таблицы (_build_table_element()) -- вызывается ПОСЛЕ tpl.render(),
        маркер к этому моменту -- обычный текст где-то в теле документа,
        раньше искать/менять нечего.

        Ищет только среди doc.paragraphs (прямые абзацы тела документа,
        как и у _append_docx_body()) -- если поле с таблицей вставлено НЕ
        отдельным абзацем (например, руками помещено внутрь ячейки другой
        таблицы шаблона), маркер не будет найден и останется видимым
        текстом в документе. В этом приложении так исторически не делают
        -- каждый плейсхолдер реквизитов рендерится в свой отдельный
        абзац (см. _render_slot_fields()), другое расположение вне
        охвата.

        Абзац с маркером убирается ЦЕЛИКОМ, а таблица вставляется на его
        место -- предполагается, что маркер был единственным содержимым
        абзаца (обычный случай для реквизитов конструктора); если в
        абзаце реально был ещё текст вперемешку с маркером, он тоже
        пропадёт вместе с абзацем -- известное упрощение."""
        from ..services.title_variants_store import load_field_tables

        tables = load_field_tables()
        if not tables:
            return

        for paragraph in list(doc.paragraphs):
            match = _TABLE_MARKER_RE.search(paragraph.text)
            if not match:
                continue
            field_id = match.group(1)
            table = tables.get(field_id)
            p_element = paragraph._p
            parent = p_element.getparent()
            if table is not None:
                p_element.addprevious(self._build_table_element(table))
            parent.remove(p_element)

    def _calculate_constructor(self):
        """Сборка документа конструктора (equipment_type == "constructor").

        Независимые разделы -- «Титульные листы» (title), «Вводная часть»
        (intro), «Приложение 1» (appendix1) и любые добавленные оператором
        (см. _CONSTRUCTOR_SLOTS/_create_section()), заполненные независимо
        друг от друга, склеиваются в ОДИН .docx: каждый раздел рендерится
        docxtpl отдельно (свой фрагмент, свои плейсхолдеры), затем
        содержимое всех разделов, кроме первого, дописывается в тело
        документа первого (_append_docx_body()) в ФИКСИРОВАННОМ порядке
        _CONSTRUCTOR_SLOTS -- порядке добавления раздела, title первым, если
        он ещё существует -- порядок в документе не зависит от того, какой
        раздел заполнили раньше. Если заполнен только один раздел --
        результат тот же, что и в Phase 1 (просто сохранённый рендер одного
        фрагмента, склеивать нечего)."""
        filled = [
            (slot, variant_id)
            for slot in self._CONSTRUCTOR_SLOTS
            for variant_id in [self._filled_slot_variant(slot)]
            if variant_id is not None
        ]
        if not filled:
            self.show_message(
                "Нечего собирать",
                "Перетащите титульный лист, вводную часть и/или приложение 1 из списка слева в "
                "основную область.",
                QMessageBox.Icon.Warning,
            )
            return

        try:
            from ..config import OUTPUT_DIR

            form_data = self.get_form_data()

            rendered_docs = []
            for slot, variant_id in filled:
                tpl = DocxTemplate(self._find_slot_template(slot, variant_id))
                tpl.render(form_data)
                # НЕ tpl.get_docx() -- он вызывает init_docx(reload=True),
                # который при is_rendered=True (выставляется в render())
                # заново перечитывает ИСХОДНЫЙ файл с диска, откатывая
                # рендер (проверено эмпирически на реальной версии
                # docxtpl, установленной в проекте -- поведение нигде явно
                # не задокументировано). tpl.docx -- тот же объект
                # python-docx Document, но БЕЗ этого отката: render()
                # подменяет его корневой body напрямую (map_tree()) ещё до
                # возврата из render().
                #
                # _splice_table_placeholders() -- ОБЯЗАТЕЛЬНО на КАЖДЫЙ
                # tpl.docx отдельно, до _append_docx_body() ниже: маркер
                # табличного поля (см. get_form_data()) попадает буквально
                # в текст ОДНОГО конкретного фрагмента (title/intro/
                # appendix1/...), искать его стоит, пока фрагменты ещё не
                # склеены -- после склейки результат тот же (абзацы никуда
                # не деваются), но по одному фрагменту дешевле и проще для
                # чтения, чем один проход по уже смешанному телу.
                self._splice_table_placeholders(tpl.docx)
                rendered_docs.append(tpl.docx)

            combined = rendered_docs[0]
            for extra_doc in rendered_docs[1:]:
                self._append_docx_body(combined, extra_doc)

            doc_number_widget = getattr(self, "doc_number", None)
            safe_doc_number = (
                doc_number_widget.toPlainText().strip().replace("/", "-") if doc_number_widget else ""
            )
            output_filename = f"констр_{safe_doc_number or 'документ'}.docx"

            output_dir = str(OUTPUT_DIR)
            os.makedirs(output_dir, exist_ok=True)
            default_path = os.path.join(output_dir, output_filename)

            output_path, _ = QFileDialog.getSaveFileName(
                self, "Сохранить документ", default_path, "Документы Word (*.docx)"
            )
            if not output_path:
                return  # отменено пользователем -- не ошибка

            combined.save(output_path)

            self.show_message(
                "Готово!", f"Документ успешно сохранён:\n{output_path}", QMessageBox.Icon.Information,
            )
            print(f"Документ успешно сохранён: {output_path}")

        except FileNotFoundError as fe:
            self.show_message("Файл не найден", str(fe), QMessageBox.Icon.Critical)
        except PermissionError:
            self.show_message(
                "Ошибка доступа", "Нет прав для записи в указанную папку", QMessageBox.Icon.Critical,
            )
        except Exception as e:
            self.show_message("Ошибка генерации", f"Неизвестная ошибка: {str(e)}", QMessageBox.Icon.Critical)

    # -- Конструктор документов: сайдбар «Титульные листы» -------------------

    INCLUDED_BLOCK_PLACEHOLDER = "__empty__"

    # -- Пары (store, generator) на каждый слот -- маленькая диспетчерская
    # таблица вместо if slot == "title" в каждом методе ниже. Собирается
    # ЛЕНИВО внутри метода (не на уровне класса) -- как и раньше в этом
    # файле, импорт services.* на верхнем уровне модуля дал бы циклический
    # импорт (см. комментарий у find_title_template() в src/config.py).
    # Модуль-стор каждого слота (title/intro/appendix_variants_store.py, а
    # для раздела, заведённого оператором, -- custom_section_variants_store.py)
    # экспортирует одинаковые имена load_variants()/save_variants()/
    # get_all_variants() (см. их докстринг в title_variants_store.py) --
    # вызывающему коду не нужно знать, какой это слот, достаточно один раз
    # выбрать правильный модуль здесь.
    @staticmethod
    def _slot_store(slot: str):
        """Раньше конечная ветка была голым else -- безопасно, пока
        реальных слотов было ровно три (title/intro/appendix1). Раздел,
        заведённый оператором (_create_section()), сюда тоже попадал бы --
        молча читал/писал бы appendix_variants.json вместо своего --
        поэтому теперь явное elif "appendix1", а всё остальное собирается
        через custom_section_variants_store (один файл на все
        пользовательские разделы, см. src/config.py,
        CUSTOM_SECTION_VARIANTS_FILE) -- functools.partial() фиксирует slot
        первым позиционным аргументом, чтобы вызывающий код
        (self._slot_store(slot).load_variants()) не отличал этот случай от
        обычного модуля с теми же тремя именами функций."""
        if slot == "title":
            from ..services import title_variants_store as store
            return store
        if slot == "intro":
            from ..services import intro_variants_store as store
            return store
        if slot == "appendix1":
            from ..services import appendix_variants_store as store
            return store
        from ..services import custom_section_variants_store
        return SimpleNamespace(
            load_variants=functools.partial(custom_section_variants_store.load_variants, slot),
            save_variants=functools.partial(custom_section_variants_store.save_variants, slot),
            get_all_variants=functools.partial(custom_section_variants_store.get_all_variants, slot),
        )

    @staticmethod
    def _slot_generate_fragment(slot: str):
        from ..services.template_generator import (
            generate_appendix_fragment, generate_intro_fragment, generate_section_fragment, generate_title_fragment,
        )
        if slot == "title":
            return generate_title_fragment
        if slot == "intro":
            return generate_intro_fragment
        if slot == "appendix1":
            return generate_appendix_fragment
        # generate_section_fragment(variant, output_path, slot) берёт третий
        # позиционный аргумент -- вызывающий код (_open_add_variant_dialog(),
        # _remove_variant_template()) зовёт результат как
        # generate_fn(variant, output_path), поэтому slot фиксируется здесь.
        return functools.partial(generate_section_fragment, slot=slot)

    @staticmethod
    def _variant_fragment_path(slot: str, variant) -> Path:
        """Куда фактически сохранён/будет сохранён .docx-файл варианта --
        выбранный оператором путь (variant.template_path), если он есть, иначе
        фиксированный FRAGMENTS_DIR/{slot}_{id}.docx, как и раньше (см.
        find_title_template()/find_intro_template()/find_appendix_template()/
        find_section_template() в src/config.py -- та же логика приоритета,
        этот хелпер её только дублирует для UI-кода, которому нужен путь ДО
        того, как файл реально создан на диске)."""
        from ..config import FRAGMENTS_DIR

        if variant.template_path:
            return Path(variant.template_path)
        return FRAGMENTS_DIR / f"{slot}_{variant.id}.docx"

    @staticmethod
    def _slot_builtin_variants(slot: str):
        from ..services.template_schema import APPENDIX_VARIANTS, INTRO_VARIANTS, TITLE_VARIANTS
        if slot == "title":
            return TITLE_VARIANTS
        if slot == "intro":
            return INTRO_VARIANTS
        if slot == "appendix1":
            return APPENDIX_VARIANTS
        return {}  # у пользовательского раздела встроенных вариантов не бывает

    @staticmethod
    def _slot_default_fields(slot: str):
        """Стартовый набор плейсхолдеров нового варианта этого слота (см.
        _open_add_variant_dialog()) -- своя константа на каждый слот, та же
        дисциплина лениво импортируемых пар, что и у соседних
        _slot_store()/_slot_generate_fragment() выше. У раздела, заведённого
        оператором, стартового набора нет -- пустой список, оператор
        добавляет поля через «Вставить плейсхолдер» с нуля."""
        from ..services.template_schema import (
            DEFAULT_APPENDIX_SUBTITLE_FIELDS, DEFAULT_INTRO_SUBTITLE_FIELDS, DEFAULT_TITLE_SUBTITLE_FIELDS,
        )
        if slot == "title":
            return DEFAULT_TITLE_SUBTITLE_FIELDS
        if slot == "intro":
            return DEFAULT_INTRO_SUBTITLE_FIELDS
        if slot == "appendix1":
            return DEFAULT_APPENDIX_SUBTITLE_FIELDS
        return []

    def _toggle_constructor_group(self, slot: str, expanded: bool):
        """Сворачивание/разворачивание группы сайдбара («Титульные листы»
        или «Вводная часть») -- список вариантов и add*VariantBtn лежат в
        общем *GroupContent (.ui, свой QVBoxLayout с spacing=2 -- вплотную,
        как строки одного списка, а не 6px общего sidebarLayout, из-за
        которых «Добавить» смотрелся самостоятельной секцией) и
        прячутся/показываются одним setVisible() на контейнере. Шеврон на
        кнопке-заголовке меняет направление. QPushButton поддерживает
        только один icon() -- шеврон и папка (как в мокапе, см.
        docs/design/constructor_mockup.html) собираются в один
        composite-пиксель через icons.combine()."""
        self._slot_widget(slot, "group_content").setVisible(expanded)
        chevron = "chevron-down" if expanded else "chevron-right"
        self._slot_widget(slot, "group_toggle").setIcon(
            icons.combine([(chevron, "#8e8e93"), ("files", "#8e8e93")], size=13, gap=4)
        )

    ITEM_CHARS_PER_LINE = 18  # см. _set_wrapped_item_size_hint()

    def _set_wrapped_item_size_hint(self, slot: str, item: QListWidgetItem):
        """QListWidget не пересчитывает высоту item'а под перенесённый на
        несколько строк текст сам по себе -- реальная отрисованная высота
        строки (QSS font-size, паддинг item'а) оказалась заметно больше,
        чем даёт QFontMetrics/QLabel.sizeHint() при замере до полного
        применения стиля (проверено: даже нарочно завышенная в 2 раза
        оценка не спасала от обрезки многоточием). Вместо хрупкого замера
        шрифтом -- грубая, но надёжная оценка по числу символов: пусть
        лучше немного лишнего пустого места снизу, чем обрезанный текст.

        QSize с отрицательной шириной Qt считает невалидным и тихо
        игнорирует весь setSizeHint() целиком (проверено) -- ширину нельзя
        оставить «как есть» через -1, нужно явное значение.

        Длина считается НЕ от item.text() -- у подсвеченной карточки
        (см. _highlight_available_block_items()) текст самого item'а
        нарочно пуст (виден только через itemWidget), длина от него
        всегда давала бы минимальную 1-строчную оценку независимо от
        реальной длины названия варианта. Название берётся из каталога
        по variant_id (UserRole) -- источник, который не зависит от
        того, подсвечен сейчас item или нет."""
        get_all_variants = self._slot_store(slot).get_all_variants
        variant = get_all_variants().get(item.data(Qt.ItemDataRole.UserRole))
        label = variant.document_title if variant is not None else item.text()
        width = self._slot_widget(slot, "available_list").viewport().width() or 200
        lines = max(1, math.ceil(len(label) / self.ITEM_CHARS_PER_LINE))
        # +14 -- запас под вертикальный padding item'а (QSS: 5px сверху и
        # снизу + пара px буфера), уменьшен вместе с самим padding'ом (был
        # 8px -> +20) -- иначе после уменьшения padding карточки остались
        # бы той же высоты, с пустым местом снизу вместо более компактного
        # вида.
        item.setSizeHint(QSize(width, lines * 22 + 14))

    def _refresh_available_blocks_list(self, slot: str):
        """Перестраивает available_list слота из get_all_*_variants()
        (встроенные варианты + пользовательские из JSON) -- вызывается при
        старте и после любой правки списка вариантов (добавление/удаление).
        Кнопка «Добавить» -- отдельный add*VariantBtn под списком (.ui), не
        item в этом списке (см. мокап -- .add-block-row с пунктирной
        рамкой и hover, недостижимо через QSS ::item на отдельном item'е
        одного списка)."""
        available_list = self._slot_widget(slot, "available_list")
        get_all_variants = self._slot_store(slot).get_all_variants

        available_list.clear()
        for variant_id, title_config in get_all_variants().items():
            item = QListWidgetItem(title_config.document_title)
            item.setIcon(icons.icon("file-text", "#0a84ff", 15))
            item.setData(Qt.ItemDataRole.UserRole, variant_id)
            available_list.addItem(item)

        # Пересчёт высоты -- следующим тиком цикла событий: во время
        # заполнения (в т.ч. при старте, до первого show()) viewport() ещё
        # не имеет окончательной геометрии, ширина под перенос текста и
        # суммарная высота item'ов посчитались бы неверно.
        QTimer.singleShot(0, functools.partial(self._resize_available_blocks_list_to_content, slot))
        # Полная пересборка списка (clear() выше) стирает и цвет
        # иконки/жирность, выставленные _highlight_available_block_items()
        # -- пересобираем подсветку сразу же, иначе после любой правки
        # каталога (добавили/переименовали/удалили вариант) уже
        # включённая в документ карточка перестала бы выделяться.
        self._highlight_available_block_items(slot)

    def _highlight_available_block_items(self, slot: str):
        """Подсвечивает в палитре (available_list слота) карточку
        варианта, который сейчас реально включён в документ (см.
        _filled_slot_variant()) -- тем же стилем, что и чип-плейсхолдер в
        реквизитах (QLabel[titleChip="true"], синий полупрозрачный фон +
        синяя рамка), а не жирным шрифтом/цветом иконки: так подсветка
        выглядит согласованно с уже принятым в приложении визуальным
        языком "это что-то особое" вместо изобретения нового.

        item.setBackground()/setForeground() на подсвечиваемый item НЕ
        работают -- ::item уже задаёт background/color в CONSTRUCTOR_QSS
        (QListWidget#availableBlocksList::item {...}), и активный
        стиль списка побеждает данные конкретного item'а (проверено
        эмпирически: цвет из данных item'а не проявляется, пока список
        стилизован через QSS). Поэтому подсвеченный item получает
        itemWidget (см. _build_available_block_chip()) -- тот же приём,
        что уже используется для includedBlockList, itemWidget рисует
        свой блок целиком поверх содержимого item'а. Остальные item'ы --
        removeItemWidget(), обычный вид (иконка+текст через ::item), как
        было до подсветки.

        item.setText("")/setIcon(QIcon()) на подсвечиваемом item -- иначе
        собственные иконка и текст item'а всё равно проступают рядом с
        itemWidget (тот же приём уже применялся в этом файле для
        перетаскиваемого item'а includedBlockList, см. историю
        "Fixed duplicated icon"). Название варианта в этом случае
        берётся не из item.text() (он его лишится), а заново из
        get_all_variants() по variant_id -- при снятии подсветки текст и
        иконка item'а восстанавливаются оттуда же, а не из чего-то
        закэшированного на самом item'е."""
        included_variant_id = self._filled_slot_variant(slot)
        available_list = self._slot_widget(slot, "available_list")
        get_all_variants = self._slot_store(slot).get_all_variants
        variants = get_all_variants()
        for i in range(available_list.count()):
            item = available_list.item(i)
            variant_id = item.data(Qt.ItemDataRole.UserRole)
            variant = variants.get(variant_id)
            label = variant.document_title if variant is not None else ""
            if variant_id == included_variant_id:
                item.setText("")
                item.setIcon(QIcon())
                available_list.setItemWidget(item, self._build_available_block_chip(label))
                item.setToolTip("Этот вариант сейчас включён в документ")
            else:
                available_list.removeItemWidget(item)
                item.setText(label)
                item.setIcon(icons.icon("file-text", "#0a84ff", 15))
                item.setToolTip("")

    def _build_available_block_chip(self, text: str) -> QWidget:
        """itemWidget подсвеченной карточки в палитре -- визуально та же
        пара иконка+текст, что и у обычной карточки, только в чипе
        (см. _highlight_available_block_items())."""
        row = QWidget()
        row.setObjectName("availableBlockChip")
        layout = QHBoxLayout(row)
        layout.setContentsMargins(9, 4, 9, 4)
        layout.setSpacing(6)
        icon_label = QLabel()
        icon_label.setPixmap(icons.render("file-text", "#5ab4ff", 15))
        layout.addWidget(icon_label)
        text_label = QLabel(text)
        text_label.setWordWrap(True)
        layout.addWidget(text_label, 1)
        return row

    def _resize_available_blocks_list_to_content(self, slot: str):
        """QListWidget с vsizetype=Maximum (.ui) сам по себе не растягивает
        sizeHint() под сумму высот item'ов -- без этого список обрезался бы
        внутренним скроллом даже когда под содержимое хватает места в
        сайдбаре."""
        list_widget = self._slot_widget(slot, "available_list")
        for i in range(list_widget.count()):
            self._set_wrapped_item_size_hint(slot, list_widget.item(i))
        total_height = sum(
            list_widget.item(i).sizeHint().height() for i in range(list_widget.count())
        )
        list_widget.setFixedHeight(total_height + 4)

    _ADD_VARIANT_DIALOG_TITLE = {
        "title": "Новый вариант титульного листа",
        "intro": "Новый вариант вводной части",
        "appendix1": "Новый вариант приложения 1",
    }

    def _open_add_variant_dialog(self, slot: str):
        """Модалка «Новый вариант титульного листа»/«Новый вариант вводной
        части» -- сохраняет метаданные (id + заголовок + стартовый набор
        полей) в JSON, как справочники сотрудников/приборов, И СРАЗУ
        генерирует .docx-заготовку (generate_title_fragment()/
        generate_intro_fragment()) -- раньше вариант оставался "сиротой"
        без файла, пока оператор вручную не запускал CLI (см.
        find_title_template() и historical note в src/config.py); теперь
        find_title_template()/find_intro_template() у нового варианта не
        падает никогда. Заготовка молча ложится в FRAGMENTS_DIR -- без
        диалога выбора места (было и сразу убрано: прерывало создание
        варианта лишним вопросом при каждом «Добавить»). Перенести файл в
        другое место можно потом, отдельным действием -- «Сохранить как»
        на уже загруженной карточке (_save_variant_template_as(),
        src/ui/template_location.py)."""
        dialog = QDialog(self)
        dialog_title = self._ADD_VARIANT_DIALOG_TITLE.get(slot) or f"Новый вариант: {self._CONSTRUCTOR_SLOTS[slot]['display_label']}"
        dialog.setWindowTitle(dialog_title)
        layout = QVBoxLayout(dialog)
        layout.addWidget(QLabel("Заголовок:"))
        line_edit = QLineEdit()
        line_edit.setPlaceholderText("Например, Протокол по результатам контроля")
        layout.addWidget(line_edit)
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)

        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        text = line_edit.text().strip()
        if not text:
            return

        from ..config import FRAGMENTS_DIR
        from ..models.title_variant import TitleVariant

        store = self._slot_store(slot)
        new_variant = TitleVariant(
            id=uuid4().hex[:8], document_title=text, subtitle_fields=list(self._slot_default_fields(slot)),
        )
        variants = store.load_variants()
        variants.append(new_variant)
        store.save_variants(variants)
        self._slot_generate_fragment(slot)(new_variant, FRAGMENTS_DIR / f"{slot}_{new_variant.id}.docx")
        self._refresh_available_blocks_list(slot)

    def _show_available_block_context_menu(self, slot: str, pos):
        """Правый клик по варианту -- «Переименовать»/«Удалить», только для
        пользовательских вариантов (встроенные TITLE_VARIANTS/INTRO_VARIANTS
        через UI не меняются -- сейчас оба словаря пусты, см. project
        memory, но проверка оставлена на случай, если встроенный вариант
        когда-нибудь вернётся)."""
        available_list = self._slot_widget(slot, "available_list")
        item = available_list.itemAt(pos)
        if item is None:
            return
        variant_id = item.data(Qt.ItemDataRole.UserRole)
        if variant_id is None or variant_id in self._slot_builtin_variants(slot):
            return

        menu = QMenu(self)
        rename_action = menu.addAction(icons.icon("edit", "#c7c7cc", 14), "Переименовать")
        delete_action = menu.addAction(icons.icon("trash", "#ff453a", 14), "Удалить")
        chosen = menu.exec(available_list.mapToGlobal(pos))
        if chosen == rename_action:
            self._rename_variant(slot, variant_id, item.text())
        elif chosen == delete_action:
            self._delete_variant(slot, variant_id, item.text())

    def _rename_variant(self, slot: str, variant_id: str, current_label: str):
        """«Переименовать» -- меняет только отображаемое название
        (TitleVariant.document_title в JSON: заголовок карточки в сайдбаре,
        крошка над документом, если это сейчас первый заполненный раздел
        [см. _update_crumb()], подпись перетащенного блока). НЕ трогает уже
        сгенерированный
        .docx-фрагмент варианта -- его текст/вёрстку пользователь правит
        только вручную в Word (см. vsk-21: регенерация задним числом
        сознательно не делается нигде в этом фиче, иначе затирала бы
        ручные правки), так что заголовок на самой странице документа
        переименование не меняет. Если файл фрагмента когда-нибудь
        пропадёт, safety-net _ensure_fragment_exists() перегенерирует его
        уже с новым названием."""
        new_name, ok = QInputDialog.getText(
            self, "Переименовать вариант", "Новое название:", text=current_label,
        )
        new_name = new_name.strip()
        if not ok or not new_name or new_name == current_label:
            return

        store = self._slot_store(slot)
        variants = store.load_variants()
        variant = next((v for v in variants if v.id == variant_id), None)
        if variant is None:
            return
        variant.document_title = new_name
        store.save_variants(variants)
        self._refresh_available_blocks_list(slot)

        included = self._slot_widget(slot, "included_list")
        if included.count() and included.item(0).data(Qt.ItemDataRole.UserRole) == variant_id:
            label_attr = self._slot_config(slot, "text_label_attr")
            if hasattr(self, label_attr):
                getattr(self, label_attr).setText(new_name)
            self._update_crumb()

    def _delete_variant(self, slot: str, variant_id: str, label: str):
        answer = QMessageBox.question(
            self, "Удалить вариант", f"Удалить вариант «{label}»?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return

        store = self._slot_store(slot)
        variants = [v for v in store.load_variants() if v.id != variant_id]
        store.save_variants(variants)
        self._refresh_available_blocks_list(slot)

        included = self._slot_widget(slot, "included_list")
        if included.count() and included.item(0).data(Qt.ItemDataRole.UserRole) == variant_id:
            self._clear_included_block(slot)

    # -- Конструктор документов: область документа ---------------------------

    def _on_block_dropped(self, slot: str, parent, first, last):
        """Реагирует на успешный drop в included_list слота. rowsInserted
        стреляет ДО того, как Qt успевает заполнить данные нового item'а
        (проверено: text() и UserRole внутри этого сигнала ещё пустые,
        заполняются через мгновение уже после dropMimeData()) -- поэтому
        сама обработка отложена на следующий тик цикла событий через
        QTimer.singleShot(0, ...), где item уже полностью готов."""
        QTimer.singleShot(0, functools.partial(self._process_block_drop, slot))

    def _process_block_drop(self, slot: str):
        """Реагирует на успешный drop в included_list слота (Qt сам
        создаёт item со всеми ролями исходного, включая UserRole, при
        перетаскивании между двумя QListWidget -- см. план). В документе
        ровно один блок на слот (Phase 1) -- второй drop в тот же слот
        оставляет только последний добавленный item.

        Все available_list (title/intro/appendix1) настроены с dragEnabled
        -- Qt ничего не мешает перетащить карточку ИЗ ЧУЖОГО списка
        (например, титульный лист -- в included_list вводной части):
        ничего в самом DnD-механизме Qt эту пару список-источник/список-назначение не
        связывает. Поэтому первым делом -- проверка принадлежности
        variant_id этому слоту (get_all_title_variants()/
        get_all_intro_variants()); если карточка не отсюда -- item
        удаляется и обработка на этом заканчивается, как будто drop'а не
        было вовсе.

        rowsInserted стреляет и на программные addItem() (см.
        _show_included_block_placeholder() -- вызывается при старте и при
        очистке, не только на реальный drag-and-drop), поэтому плейсхолдер
        сначала исключается из рассмотрения, а не только последний item --
        Qt может вставить перетащенный item и ПЕРЕД плейсхолдером/старым
        item'ом (см. ниже), не только после."""
        block_list = self._slot_widget(slot, "included_list")
        real_indices = [
            i for i in range(block_list.count())
            if block_list.item(i).data(Qt.ItemDataRole.UserRole) != self.INCLUDED_BLOCK_PLACEHOLDER
        ]
        if not real_indices:
            return  # в списке только сам плейсхолдер -- нечего обрабатывать

        # При замене (два реальных item'а: старый + новый перетащенный, или
        # первый drop -- плейсхолдер + новый) Qt сам решает, вставить ли
        # новый item ДО или ПОСЛЕ уже лежащего, в зависимости от того, в
        # верхнюю или нижнюю половину курсор попал при drop'е
        # (DropIndicatorPosition AboveItem/BelowItem) -- порядок в списке не
        # гарантирован. Раньше код слепо брал последний item -- если Qt
        # вставлял новый ПЕРЕД старым, последним оставался старый item с
        # уже стёртым text() (см. item.setText("") ниже, из предыдущей
        # обработки) -- отсюда пустая строка вместо названия при повторной
        # замене. Новый item опознаётся по отсутствию itemWidget: его
        # выставляет только эта функция, а Qt при перетаскивании копирует
        # роли исходного item'а (текст, UserRole), но не itemWidget.
        #
        # Если ВСЕ реальные item'ы уже с itemWidget (уже были обработаны
        # раньше) -- функция не идемпотентна по конструкции (она стирает
        # text() обрабатываемого item'а), поэтому повторный вызов без
        # нового перетащенного item'а должен быть no-op, а не портить уже
        # готовую карточку.
        keep_index = next(
            (i for i in real_indices if block_list.itemWidget(block_list.item(i)) is None),
            None,
        )
        if keep_index is None:
            return

        new_item = block_list.item(keep_index)
        variant_id = new_item.data(Qt.ItemDataRole.UserRole)
        get_all_variants = self._slot_store(slot).get_all_variants
        if variant_id not in get_all_variants():
            # Перетащили карточку не из своей группы -- откатываем сам
            # drop (убираем только что вставленный item), оставляя слот в
            # прежнем состоянии (пустой плейсхолдер или старый блок).
            block_list.takeItem(keep_index)
            if block_list.count() == 0:
                self._show_included_block_placeholder(slot)
            return

        for i in reversed(range(block_list.count())):
            if i != keep_index:
                block_list.takeItem(i)

        item = block_list.item(0)
        label_text = item.text()
        # Текст и иконку item'а показывает не нативная отрисовка делегата, а
        # сам row (QLabel/QPixmap ниже) -- иначе поверх setItemWidget()
        # проступают оригинальные текст и иконка item'а вторым, наложенным
        # слоем (иконка -- та же самая file-text, что ставит
        # _refresh_available_blocks_list() на карточку в available_list,
        # и Qt копирует DecorationRole вместе с текстом при перетаскивании
        # между списками, отсюда видимое задвоение). sizeHint тоже сбрасываем
        # -- Qt копирует роли (включая SizeHintRole) исходного item'а при
        # перетаскивании, из-за чего сюда попадала многострочная высота
        # карточки в available_list вместо компактной строки.
        item.setText("")
        item.setIcon(QIcon())
        item.setSizeHint(QSize(block_list.viewport().width() or 200, 42))

        row = QWidget()
        row_layout = QHBoxLayout(row)
        row_layout.setContentsMargins(10, 8, 6, 8)
        # WA_TransparentForMouseEvents -- клики по тексту/иконкам должны
        # доходить до row.mousePressEvent (сворачивание/разворачивание
        # реквизитов, см. _toggle_fields_panel()), а не гаситься самими
        # QLabel (мышиные события Qt не всплывают от ребёнка к родителю
        # сами по себе, в отличие от event bubbling в DOM).
        icon_label = QLabel()
        icon_label.setPixmap(icons.render("file-text", "#0a84ff", 15))
        icon_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        row_layout.addWidget(icon_label)
        text_label = QLabel(label_text)
        text_label.setWordWrap(True)
        text_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        row_layout.addWidget(text_label, stretch=1)
        # Сохраняется на self -- _rename_variant() обновляет текст этой
        # подписи вживую, если переименовывают вариант, который сейчас
        # лежит в документе (иначе после переименования тут осталось бы
        # старое название до следующего drop). Имя атрибута -- на слот (см.
        # text_label_attr/chevron_attr в _CONSTRUCTOR_SLOTS), все слоты
        # могут быть заполнены одновременно, общий атрибут потерял бы ссылку
        # на часть из них.
        text_label_attr = self._slot_config(slot, "text_label_attr")
        chevron_attr = self._slot_config(slot, "chevron_attr")
        setattr(self, text_label_attr, text_label)
        chevron_label = QLabel()
        chevron_label.setPixmap(icons.render("chevron-down", "#8e8e93", 13))
        chevron_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        setattr(self, chevron_attr, chevron_label)
        row_layout.addWidget(chevron_label)
        remove_btn = QPushButton()
        remove_btn.setIcon(icons.icon("x", "#8e8e93", 14))
        remove_btn.setFixedSize(22, 22)
        remove_btn.setFlat(True)
        remove_btn.clicked.connect(functools.partial(self._clear_included_block, slot))
        row_layout.addWidget(remove_btn)
        # QPushButton остаётся обычным (не прозрачным для мыши) -- его
        # собственный клик по-прежнему обрабатывается им самим, до row
        # не долетает, поэтому отдельный stopPropagation тут не нужен.
        row.setCursor(Qt.CursorShape.PointingHandCursor)
        row.mousePressEvent = functools.partial(self._toggle_fields_panel, slot)
        # Редактирование каталога плейсхолдеров варианта раньше открывалось
        # отдельным модальным окном по ПКМ на этой строке (см. историю) --
        # теперь оно встроено прямо в fields_panel (_render_slot_fields()),
        # разворачивается тем же кликом, что и реквизиты, отдельного пункта
        # меню/окна больше нет.
        block_list.setItemWidget(item, row)
        self._set_included_block_filled(slot, True)

        # Список без этого остаётся высотой под старый maximumSize из .ui
        # (под пустое состояние с текстом-подсказкой) -- карточка внутри
        # тогда болтается с зазором сверху/снизу вместо того, чтобы
        # заполнять всю площадь блока.
        block_list.setFixedHeight(
            item.sizeHint().height() + 2 * block_list.frameWidth() + 4
        )

        self._render_slot_fields(slot, variant_id)
        self._update_crumb()
        self._slot_widget(slot, "fields_panel").setVisible(True)
        self.pushButt_generateWord.setEnabled(True)
        self._highlight_available_block_items(slot)

    def _set_included_block_filled(self, slot: str, filled: bool):
        """Переключает QSS-состояние included_list слота через динамическое
        свойство ("filled" в CONSTRUCTOR_QSS) -- пунктирная рамка только в
        пустом состоянии, у заполненной карточки своя (сплошной фон,
        никакой рамки у самого списка) -- иначе рамка списка и фон
        item'а никогда не совпадают точь-в-точь (зазоры/наплывы по краям,
        не совпадающие радиусы), что и было исходной проблемой."""
        block_list = self._slot_widget(slot, "included_list")
        block_list.setProperty("filled", filled)
        block_list.style().unpolish(block_list)
        block_list.style().polish(block_list)

    def _show_included_block_placeholder(self, slot: str):
        """Пустое состояние included_list слота -- ненажимаемый
        item-подсказка вместо пустого списка без текста (Qt не даёт
        «placeholder-текст» у QListWidget нативно). Иконка над текстом --
        как в мокапе (docs/design/constructor_mockup.html, #dropEmpty) --
        через itemWidget: сам item лишь резервирует место (NoItemFlags,
        пустой текст), рисует содержимое QWidget с QVBoxLayout. Текст
        подсказки -- свой на каждый слот (_slot_config(slot, "empty_hint"))."""
        block_list = self._slot_widget(slot, "included_list")
        block_list.clear()

        placeholder = QListWidgetItem()
        placeholder.setData(Qt.ItemDataRole.UserRole, self.INCLUDED_BLOCK_PLACEHOLDER)
        placeholder.setFlags(Qt.ItemFlag.NoItemFlags)
        # frameShape=NoFrame (.ui) -- inner-высота виджета совпадает с его
        # maximumSize, поэтому item получает все 56px без поправок на рамку
        # -- иначе (без явного sizeHint) реальная высота item'а была бы
        # только под содержимое itemWidget, а не под весь box.
        placeholder.setSizeHint(QSize(block_list.viewport().width() or 200, 56))
        block_list.addItem(placeholder)

        row = QWidget()
        row_layout = QVBoxLayout(row)
        row_layout.setContentsMargins(0, 0, 0, 0)
        row_layout.setSpacing(6)
        icon_label = QLabel()
        icon_label.setPixmap(icons.render("drag-drop", "#5a5a5c", 22))
        icon_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        row_layout.addWidget(icon_label)
        text_label = QLabel(self._slot_config(slot, "empty_hint"))
        text_label.setStyleSheet("color: #5a5a5c; font-size: 12px;")
        text_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        row_layout.addWidget(text_label)
        block_list.setItemWidget(placeholder, row)

        self._set_included_block_filled(slot, False)
        block_list.setFixedHeight(56)

    def _clear_included_block(self, slot: str):
        self._show_included_block_placeholder(slot)
        self._slot_widget(slot, "fields_panel").setVisible(False)
        self._update_crumb()
        any_filled = any(self._filled_slot_variant(s) is not None for s in self._CONSTRUCTOR_SLOTS)
        self.pushButt_generateWord.setEnabled(any_filled)
        self._highlight_available_block_items(slot)

    def _set_crumb(self, active_section: str):
        """Обновляет crumbLabel -- «Конструктор документов / <активный
        раздел>», где активная часть подсвечена ярче (#e5e5e7) на фоне
        тусклого префикса (#8e8e93 из CONSTRUCTOR_QSS), как crumbBlock в
        docs/design/constructor_mockup.html. html.escape() -- active_section
        приходит из label_text (текст item'а сайдбара), не буквальный
        константный литерал."""
        self.crumbLabel.setText(
            f'Конструктор документов / <span style="color:#e5e5e7;">{html.escape(active_section)}</span>'
        )

    def _update_crumb(self):
        """Пересчитывает крошку -- показывает название варианта ПЕРВОГО ПО
        ПОРЯДКУ _CONSTRUCTOR_SLOTS заполненного раздела (раньше жёстко
        отражала только title -- он и был первым; при удалении раздела
        title может больше не существовать вовсе, см. обсуждение задачи и
        мокап docs/design/вводная_часть.html, updateCrumb()). Вызывается
        после любого drop/очистки/переименования варианта в ЛЮБОМ разделе,
        а не только title -- дешёвая перепроверка, сама подстановка не
        меняется, если первый заполненный раздел не тот, что сейчас
        менялся."""
        for slot in self._CONSTRUCTOR_SLOTS:
            variant_id = self._filled_slot_variant(slot)
            if variant_id is None:
                continue
            variant = self._slot_store(slot).get_all_variants().get(variant_id)
            if variant is not None:
                self._set_crumb(variant.document_title)
                return
        self._set_crumb("без разделов")

    def _toggle_fields_panel(self, slot: str, event):
        """Левый клик по перетащенному блоку в included_list слота
        сворачивает/разворачивает fields_panel (реквизиты) -- см. мокап
        docs/design/constructor_mockup.html, toggleFields(). Клик по
        remove_btn ("×") сюда не долетает (см. _process_block_drop())."""
        if event.button() != Qt.MouseButton.LeftButton:
            return
        fields_panel = self._slot_widget(slot, "fields_panel")
        visible = not fields_panel.isVisible()
        fields_panel.setVisible(visible)
        chevron_attr = self._slot_config(slot, "chevron_attr")
        getattr(self, chevron_attr).setPixmap(
            icons.render("chevron-down" if visible else "chevron-right", "#8e8e93", 13)
        )

    def _render_slot_fields(self, slot: str, variant_id: str):
        """Перестраивает fields_layout слота под набор полей конкретного
        варианта (get_all_title_variants()[variant_id].subtitle_fields или
        get_all_intro_variants() для intro) -- поля разные у разных
        вариантов, поэтому строятся в рантайме, а не заранее в .ui.

        Значения уже заполненных полей, которые остаются в новом наборе,
        сохраняются при пересборке -- вызывается не только при первом дропе
        блока в документ, но и повторно после любого добавления/удаления
        плейсхолдера через каталог (см. _build_placeholder_catalog() ниже),
        где набор плейсхолдеров мог измениться, а уже введённые значения
        оставшихся полей терять не нужно.

        Для пользовательских вариантов первой строкой (до самих полей)
        вставляется каталог плейсхолдеров -- раньше это было отдельное
        модальное окно (TitleContentEditorDialog, удалено), теперь
        встроено прямо сюда: разворачивается/сворачивается тем же кликом,
        что и реквизиты (см. _toggle_fields_panel()). Для встроенных
        вариантов (TITLE_VARIANTS/INTRO_VARIANTS) каталог не показывается
        вообще -- они собраны в Word вручную и через приложение не
        редактируются.

        Подписи полей -- из общего каталога (встроенные + пользовательские,
        см. get_all_field_labels()) -- ОБЩЕГО для обоих слотов, а не
        только встроенного TITLE_FIELD_LABELS -- иначе поле, заведённое
        через «+ Новое поле» в каталоге плейсхолдеров, показывалось бы тут
        просто как голый id.

        Динамически созданные виджеты регистрируются в
        self.PLAIN_TEXT_EDIT_NAMES -- том же списке, что уже читают
        get_form_data()/init_widgets() -- вместо отдельного пути
        сохранения/чтения данных для конструктора. Имя атрибута/ключ --
        _slot_placeholder_name(slot, field_id), НЕ голый field_id (кроме
        title, см. её докстринг) -- иначе одно и то же поле, вставленное
        сразу в оба слота, делило бы один Python-атрибут на два разных
        виджета.

        Сразу под заголовком «Реквизиты...» -- поле «Открыть предпросмотр
        итогового документа» (_build_preview_field(), клик рендерит
        документ с уже введёнными значениями -- _open_variant_preview()).
        Показывается для ЛЮБОГО варианта, включая встроенные -- в отличие
        от каталога плейсхолдеров и «Загрузить шаблон Word»
        (_build_placeholder_catalog()), которые только для
        пользовательских: предпросмотр с подставленными значениями
        одинаково полезен и для готовых встроенных вариантов."""
        from ..services.title_variants_store import get_all_field_labels, load_field_formulas, load_field_tables
        from ..services.formula_engine import evaluate_formula, format_formula_result

        get_all_variants = self._slot_store(slot).get_all_variants
        is_custom = variant_id not in self._slot_builtin_variants(slot)
        field_ids = get_all_variants()[variant_id].subtitle_fields

        dynamic_names = self._dynamic_field_names[slot]
        previous_values = {
            name: getattr(self, name).toPlainText()
            for name in dynamic_names
            if hasattr(self, name)
        }

        for name in dynamic_names:
            if name in self.PLAIN_TEXT_EDIT_NAMES:
                self.PLAIN_TEXT_EDIT_NAMES.remove(name)
        self._dynamic_field_names[slot] = []

        layout = self._slot_widget(slot, "fields_layout")
        while layout.rowCount():
            layout.removeRow(0)

        if is_custom:
            layout.addRow(self._build_placeholder_catalog(slot, variant_id, field_ids))

        # Раньше был заголовком самого fields_panel (QGroupBox.title) --
        # вынесен в обычный QLabel-ряд, чтобы идти ПОСЛЕ каталога
        # плейсхолдеров, а не над ним (QGroupBox всегда рисует свой title
        # над всем содержимым, порядок в layout на это не влияет).
        section_label = QLabel(self._slot_config(slot, "section_label"))
        section_label.setObjectName("titleFieldsSectionLabel")
        layout.addRow(section_label)

        layout.addRow(self._build_preview_field(slot, variant_id))

        if is_custom and not field_ids:
            empty = QLabel("Пока нет ни одного плейсхолдера — добавьте через кнопку «Вставить плейсхолдер» выше.")
            empty.setWordWrap(True)
            empty.setStyleSheet("color: #8e8e93; font-style: italic; font-size: 11px;")
            layout.addRow(empty)

        labels = get_all_field_labels()
        formulas = load_field_formulas()
        tables = load_field_tables()
        for field_id in field_ids:
            widget_name = self._slot_placeholder_name(slot, field_id)
            formula = formulas.get(field_id)
            table = tables.get(field_id)
            # GrowablePlaceholderField, а не голый QPlainTextEdit -- значения
            # бывают длиной в целый абзац (см. содержательные пункты вводной
            # части вроде «1.1. На основании требований п.198 ФНП ТТ ...»).
            # Растёт по высоте вместе с переносом строк вместо
            # горизонтального скролла одной строки; свёрнуто до одной
            # строки, пока не в фокусе или не закреплено шевроном -- см.
            # docs/design/вводная_часть.html и src/ui/growable_placeholder_field.py.
            widget = GrowablePlaceholderField()
            if formula is not None:
                # Вычисляемое поле -- readOnly (не disabled: значение всё
                # ещё можно выделить и скопировать), значение ВСЕГДА
                # пересчитывается заново, а не восстанавливается из
                # previous_values -- формула, а не то, что тут было
                # напечатано руками ДО того, как на поле повесили формулу,
                # теперь единственный источник истины (см. _open_formula_editor()).
                widget.setReadOnly(True)
                widget.setProperty("computed", True)
                value = evaluate_formula(formula.get("tokens", []), self._placeholder_numeric_value, formulas)
                widget.setPlainText(format_formula_result(value, formula.get("decimals", 2)))
            elif table is not None:
                # Таблица -- тоже readOnly, но, в отличие от формулы, НЕ
                # сводится к одному подставляемому значению: get_form_data()
                # уходит не с текстом этого превью, а с отдельным маркером
                # (_table_field_marker()), который потом меняется на
                # настоящую .docx-таблицу (см. _splice_table_placeholders()).
                # Поэтому тут setPlaceholderText() (короткая подсказка про
                # размер и про ПКМ), а НЕ setPlainText() -- у пустого
                # readOnly-поля toPlainText() остаётся "", виджет никак не
                # должен отвечать за то, что реально уйдёт в документ.
                widget.setReadOnly(True)
                widget.setProperty("computedTable", True)
                rows = table.get("rows", [])
                cols = len(rows[0]) if rows else 0
                widget.setPlaceholderText(f"▦ Таблица {len(rows)}×{cols} — редактирование через ПКМ")
            else:
                if widget_name in previous_values:
                    widget.setPlainText(previous_values[widget_name])
                else:
                    prefill = self._cross_slot_placeholder_value(field_id)
                    if prefill:
                        widget.setPlainText(prefill)
                # Живой пересчёт вычисляемых полей -- любое обычное (не
                # вычисляемое) поле реквизитов, в любом слоте, может
                # оказаться операндом чьей-то формулы; заранее не известно,
                # чьей именно, поэтому пересчитываем все видимые
                # вычисляемые поля на любое изменение любого обычного (см.
                # _refresh_computed_fields()), а не только то, что рядом.
                widget.textChanged.connect(self._refresh_computed_fields)
            # Для пользовательских вариантов подпись строки -- тоже сам чип
            # плейсхолдера (QLabel[titleChip="true"]), ПОЛНОСТЬЮ интерактивный
            # (клик копирует и подсвечивает -- _copy_chip(), ПКМ открывает
            # меню -- _show_chip_context_menu()), а не просто стилизованная
            # под чип надпись. Раньше такой же чип рисовался ОТДЕЛЬНЫМ
            # списком в _build_placeholder_catalog() над реквизитами --
            # пользователь поправил: список дублировал поля реквизитов
            # прямо под ним, чипы должны стоять рядом со своими полями, а
            # не отдельно. Видимый текст чипа -- пользовательское название
            # поля (labels.get(field_id)), а не технический тег
            # "{{ widget_name }}" -- оператор не должен видеть
            # Jinja-синтаксис; сам тег остаётся в тултипе и в тексте,
            # который реально копирует _copy_chip() в буфер. Встроенные
            # варианты каталога не имеют -- для них подпись остаётся
            # обычным некликабельным текстом, как раньше.
            if is_custom:
                chip_text = (
                    labels.get(field_id, field_id)
                    + (" ƒ" if formula is not None else "")
                    + (" ▦" if table is not None else "")
                )
                row_label = QLabel(chip_text)
                tooltip = f"{{{{ {widget_name} }}}}"
                if formula is not None:
                    tooltip += "\nВычисляется по формуле — правка недоступна, см. ПКМ."
                if table is not None:
                    tooltip += "\nПредставлено таблицей — правка через ПКМ → «Редактировать таблицу»."
                tooltip += "\nКлик — скопировать для Word. ПКМ — меню. Перетащите на другое поле — переставить порядок."
                row_label.setToolTip(tooltip)
                row_label.setProperty("titleChip", True)
                if formula is not None:
                    row_label.setProperty("titleChipFormula", True)
                if table is not None:
                    row_label.setProperty("titleChipTable", True)
                row_label.setCursor(Qt.CursorShape.PointingHandCursor)
                row_label.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
                row_label.customContextMenuRequested.connect(
                    lambda pos, fid=field_id, w=row_label: self._show_chip_context_menu(
                        slot, variant_id, fid, w.mapToGlobal(pos)
                    )
                )
                self._wire_chip_drag_reorder(row_label, slot, variant_id, field_id, widget_name)
            else:
                row_label = labels.get(field_id, field_id)
            layout.addRow(row_label, widget)
            setattr(self, widget_name, widget)
            self.PLAIN_TEXT_EDIT_NAMES.append(widget_name)
            self._dynamic_field_names[slot].append(widget_name)

        # Программные setPlainText() выше (previous_values/cross-slot
        # prefill) не проходят через textChanged.connect(self._refresh_computed_fields)
        # -- та подписка ставится ПОСЛЕ них, на уже готовое значение, чтобы
        # не тратить пересчёт на каждый промежуточный setPlainText() при
        # самой перестройке. Один пересчёт в конце гарантирует, что
        # вычисляемые поля (в т.ч. в ДРУГОМ слоте, если ссылаются на поле
        # этого) отражают то, что реально сейчас в реквизитах, а не только
        # то, что было на момент их собственного рендера.
        self._refresh_computed_fields()

    _CHIP_DRAG_MIME = "application/x-titlechip-field-id"

    def _wire_chip_drag_reorder(self, row_label: QLabel, slot: str, variant_id: str, field_id: str, widget_name: str):
        """Перетаскивание чипа плейсхолдера -- переставляет порядок полей
        в variant.subtitle_fields (см. _reorder_variant_placeholder()).
        Только для пользовательских вариантов (см. вызывающую сторону) --
        у встроенных чипов нет вовсе, порядок полей там не через UI.

        Клик (копирование, _copy_chip()) и начало перетаскивания различаем
        порогом смещения (QApplication.startDragDistance()) между press и
        move -- иначе обычный клик почти всегда включал бы в себя микро-
        сдвиг мыши и запускал drag вместо копирования. Само копирование
        поэтому перенесено на mouseReleaseEvent (срабатывает, только если
        drag не запускался) вместо mousePressEvent, как было раньше.

        Тот же чип служит и источником, и целью drop'а -- перетащить можно
        на любой другой чип в том же fields_layout; вставка -- до или
        после чипа-цели, в зависимости от того, в верхнюю или нижнюю
        половину его высоты попал курсор при отпускании (тот же принцип,
        что и в _process_block_drop() у самого документа)."""
        row_label._chip_drag_start_pos = None

        def mouse_press(event, w=row_label):
            if event.button() == Qt.MouseButton.LeftButton:
                w._chip_drag_start_pos = event.position().toPoint()

        def mouse_move(event, w=row_label, fid=field_id):
            if w._chip_drag_start_pos is None or not (event.buttons() & Qt.MouseButton.LeftButton):
                return
            moved = (event.position().toPoint() - w._chip_drag_start_pos).manhattanLength()
            if moved < QApplication.startDragDistance():
                return
            w._chip_drag_start_pos = None
            drag = QDrag(w)
            mime = QMimeData()
            mime.setData(self._CHIP_DRAG_MIME, fid.encode("utf-8"))
            drag.setMimeData(mime)
            drag.exec(Qt.DropAction.MoveAction)

        def mouse_release(event, w=row_label, name=widget_name):
            if w._chip_drag_start_pos is not None:
                # Курсор не сдвинулся достаточно для drag -- обычный клик.
                self._copy_chip(name, w)
            w._chip_drag_start_pos = None

        row_label.mousePressEvent = mouse_press
        row_label.mouseMoveEvent = mouse_move
        row_label.mouseReleaseEvent = mouse_release

        row_label.setAcceptDrops(True)

        def drag_enter(event, w=row_label):
            if event.mimeData().hasFormat(self._CHIP_DRAG_MIME):
                event.acceptProposedAction()
                w.setProperty("chipDragOver", True)
                w.style().unpolish(w)
                w.style().polish(w)

        def drag_leave(event, w=row_label):
            w.setProperty("chipDragOver", False)
            w.style().unpolish(w)
            w.style().polish(w)

        def drop_event(event, w=row_label, s=slot, vid=variant_id, target_fid=field_id):
            w.setProperty("chipDragOver", False)
            w.style().unpolish(w)
            w.style().polish(w)
            if not event.mimeData().hasFormat(self._CHIP_DRAG_MIME):
                return
            dragged_fid = bytes(event.mimeData().data(self._CHIP_DRAG_MIME)).decode("utf-8")
            event.acceptProposedAction()
            if dragged_fid == target_fid:
                return
            insert_after = event.position().y() > w.height() / 2
            self._reorder_variant_placeholder(s, vid, dragged_fid, target_fid, insert_after)

        row_label.dragEnterEvent = drag_enter
        row_label.dragLeaveEvent = drag_leave
        row_label.dropEvent = drop_event

    def _reorder_variant_placeholder(
        self, slot: str, variant_id: str, dragged_field_id: str, target_field_id: str, insert_after: bool,
    ):
        """Переставляет dragged_field_id рядом с target_field_id в
        variant.subtitle_fields (до или после, см. insert_after) --
        сохраняет в store слота и перестраивает fields_layout заново
        (_render_slot_fields()), тем же путём, что и добавление/удаление
        плейсхолдера через каталог. Оба id должны реально быть в текущем
        наборе полей варианта -- иначе (устаревший drag, поле успели
        удалить за время перетаскивания) молча ничего не делает."""
        store = self._slot_store(slot)
        variants = store.load_variants()
        variant = next((v for v in variants if v.id == variant_id), None)
        if variant is None:
            return
        fields = variant.subtitle_fields
        if dragged_field_id not in fields or target_field_id not in fields:
            return

        fields.remove(dragged_field_id)
        insert_at = fields.index(target_field_id) + (1 if insert_after else 0)
        fields.insert(insert_at, dragged_field_id)

        store.save_variants(variants)
        self._render_slot_fields(slot, variant_id)

    def _build_placeholder_catalog(self, slot: str, variant_id: str, field_ids) -> QWidget:
        """Верхняя панель пользовательского варианта -- кнопка «Вставить
        плейсхолдер» (общий каталог полей, get_all_field_labels(), в меню
        см. _build_placeholder_menu()), «Загрузить шаблон Word» (выбор
        .docx + запрос приложения для редактирования, см.
        _upload_variant_template()), карточка загруженного файла
        (_build_template_loaded_card()) и подсказка про клик/ПКМ по чипам.

        Отдельной кнопки «Посмотреть шаблон» больше нет -- пользователь
        попросил перенести открытие файла варианта как есть (в том
        состоянии, в котором его оставили в Word, буквально с "{{ field }}"
        там, где не заполнено значением, см. _open_variant_raw_file()) на
        клик по самой карточке загруженного файла
        (_build_template_loaded_card()), а не по отдельной кнопке. Пока
        файл не загружен (drop-рамка вместо карточки), кликнуть для
        просмотра нечего -- там открывать пока нечего.

        Сами чипы плейсхолдеров здесь БОЛЬШЕ НЕ рисуются отдельным
        списком (было раньше -- пользователь поправил: список дублировал
        поля реквизитов, стоящие прямо под ним). Теперь чип -- это подпись
        строки в самих реквизитах, рядом со своим полем ввода (см.
        _render_slot_fields(), где и живёт _copy_chip()/
        _show_chip_context_menu() для этих подписей). «Открыть
        предпросмотр» (рендер с уже введёнными в форму значениями) -- тоже
        НЕ здесь, вынесена в _render_slot_fields() под заголовок
        «Реквизиты...», см. её докстринг.

        Добавление/удаление плейсхолдера (через меню «Вставить
        плейсхолдер» или ПКМ по чипу в реквизитах) сразу сохраняется в
        store слота и перестраивает fields_layout заново
        (_render_slot_fields()) -- отдельного шага "Сохранить шаблон"
        больше нет, изменение каталога и есть сохранение.

        Вся эта панель (не только кнопка «Загрузить шаблон Word») -- ещё и
        drop-зона: перетаскивание .docx-файла из Finder сюда делает то же
        самое, что кнопка + диалог выбора файла (общая установка --
        _install_variant_template()), но без диалога. Кнопка остаётся
        равноправным способом (пользователь явно попросил добавить
        перетаскивание, а не заменить им кнопку) -- полезна, когда Finder
        не открыт рядом или файл лежит в другом окне не на весь экран."""
        store = self._slot_store(slot)
        # get_all_*_variants() отдаёт для пользовательских вариантов
        # урезанный TitleConfig (для единообразия со встроенными в списке
        # сайдбара), без template_filename -- он есть только на самом
        # TitleVariant, поэтому здесь нужен именно load_*_variants().
        variants = store.load_variants()
        variant = next((v for v in variants if v.id == variant_id), None)
        template_filename = variant.template_filename if variant else ""

        container = QWidget()
        container.setObjectName("titleTemplateDropZone")
        outer = QVBoxLayout(container)
        outer.setContentsMargins(6, 6, 6, 10)
        outer.setSpacing(6)

        toolbar = QHBoxLayout()
        insert_btn = QToolButton()
        # objectName -- чтобы _reopen_insert_menu() мог найти именно эту
        # (пере)созданную кнопку и её меню после удаления поля из каталога
        # (см. её докстринг).
        insert_btn.setObjectName(f"insertPlaceholderBtn_{slot}")
        insert_btn.setIcon(icons.icon("tag", "#c7c7cc", 13))
        insert_btn.setText(" Вставить плейсхолдер")
        insert_btn.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextBesideIcon)
        insert_btn.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
        insert_btn.setMenu(self._build_placeholder_menu(slot, variant_id, field_ids))
        toolbar.addWidget(insert_btn)
        toolbar.addStretch(1)

        upload_btn = QToolButton()
        upload_btn.setIcon(icons.icon("upload", "#0a84ff", 13))
        upload_btn.setText(" Загрузить шаблон Word")
        upload_btn.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextBesideIcon)
        upload_btn.setToolTip(
            "Выбрать .docx-файл (например, уже готовый документ) и открыть его, спросив, каким "
            "приложением (Word, Pages, LibreOffice...) редактировать -- именно в него дальше "
            "вставляются плейсхолдеры и реквизиты варианта. Файл можно и просто перетащить на эту "
            "панель вместо диалога выбора"
        )
        upload_btn.clicked.connect(functools.partial(self._upload_variant_template, slot, variant_id))
        toolbar.addWidget(upload_btn)
        outer.addLayout(toolbar)

        if template_filename:
            outer.addWidget(self._build_template_loaded_card(slot, variant_id, template_filename))
        else:
            outer.addWidget(self._build_template_drop_hint())

        hint = QLabel("Клик по плейсхолдеру в реквизитах ниже — скопировать для Word. Правый клик — удалить.")
        hint.setWordWrap(True)
        hint.setStyleSheet("color: #8e8e93; font-size: 11px;")
        outer.addWidget(hint)

        # Перетаскивание .docx должно срабатывать в любой точке панели --
        # включая саму drop-рамку (_build_template_drop_hint()) и все
        # прочие дочерние виджеты (кнопки, подписи). По документации Qt
        # события drag/drop у дочернего виджета без setAcceptDrops(True)
        # должны сами подниматься к первому предку, у которого это
        # свойство включено -- но на практике (проверено пользователем на
        # реальном запуске, не только в headless-тесте с прямой отправкой
        # события контейнеру) это не сработало понадёжнее явно включить
        # приём и на container, и на каждом потомке -- тогда результат не
        # зависит от того, какой именно пиксель ОС посчитает целью drag'а.
        for widget in [container] + container.findChildren(QWidget):
            widget.setAcceptDrops(True)
            widget.dragEnterEvent = (
                lambda event, c=container: self._title_template_drag_enter(event, c)
            )
            widget.dragMoveEvent = self._title_template_drag_move
            widget.dragLeaveEvent = (
                lambda event, c=container: self._title_template_drag_leave(event, c)
            )
            widget.dropEvent = (
                lambda event, s=slot, vid=variant_id, c=container: self._title_template_drop(event, s, vid, c)
            )

        return container

    def _build_preview_field(self, slot: str, variant_id: str) -> QWidget:
        """Поле «Открыть предпросмотр итогового документа» -- заменяет
        прежнюю обычную кнопку «Открыть предпросмотр», тот же визуальный
        язык, что и у карточки загруженного шаблона
        (_build_template_loaded_card()): скруглённая рамка, иконка
        документа слева, клик по всему полю. Показывается для ЛЮБОГО
        варианта (см. докстринг _render_slot_fields()), поэтому живёт
        здесь отдельным методом, а не внутри _build_placeholder_catalog()
        (та -- только для пользовательских вариантов).

        Два состояния, различаются только тем, открывали ли предпросмотр
        этого варианта хотя бы раз в текущем запуске приложения
        (self._preview_generated, см. её комментарий):
        - ещё ни разу -- приглушённый текст "Здесь будет итоговый документ
          для предпросмотра" (аналог пустого состояния карточки шаблона,
          _build_template_drop_hint());
        - хотя бы раз открывали -- обычный активный вид с подписью
          "Открыть предпросмотр итогового документа", плюс те же кнопки
          «Сохранить как»/«Показать в Finder», что и на карточке загруженного
          шаблона (_build_template_loaded_card()) -- файл предпросмотра
          лежит рядом с самим шаблоном варианта (см. _open_variant_preview()),
          эти кнопки дают перенести его в другое место или просто увидеть в
          Finder, не открывая заново в Word (_save_preview_as()/
          _reveal_preview()). Клик по самому полю работает одинаково в
          обоих состояниях (в первый раз он же и генерирует первый
          предпросмотр) -- разница чисто визуальная,
          подсказка "тут появится результат" до первого клика."""
        already_generated = (slot, variant_id) in self._preview_generated

        field = QWidget()
        field.setObjectName("titlePreviewField")
        field.setCursor(Qt.CursorShape.PointingHandCursor)
        field.setToolTip(
            "Собрать документ с уже введёнными значениями во временный файл и открыть его -- "
            "посмотреть, как будет выглядеть готовый документ, не сохраняя его окончательно"
        )
        field.mousePressEvent = lambda event, s=slot, vid=variant_id: self._open_variant_preview(s, vid)
        layout = QHBoxLayout(field)
        layout.setContentsMargins(9, 7, 9, 7)
        layout.setSpacing(8)

        icon_label = QLabel()
        icon_label.setPixmap(icons.render("file-text", "#0a84ff" if already_generated else "#5a5a5c", 15))
        icon_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        layout.addWidget(icon_label)

        text_label = QLabel(
            "Открыть предпросмотр итогового документа" if already_generated
            else "Здесь будет итоговый документ для предпросмотра"
        )
        text_label.setWordWrap(True)
        text_label.setProperty("titlePreviewFieldActive", already_generated)
        text_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        layout.addWidget(text_label, stretch=1)

        if already_generated:
            # Обычные (не прозрачные для мыши) QPushButton -- собственный
            # клик обрабатывается ими самими, до field.mousePressEvent не
            # долетает, тот же приём, что и у reveal_btn/save_as_btn
            # карточки шаблона (_build_template_loaded_card()).
            save_as_btn = QPushButton()
            save_as_btn.setIcon(icons.icon("download", "#8e8e93", 14))
            save_as_btn.setFixedSize(22, 22)
            save_as_btn.setFlat(True)
            save_as_btn.setToolTip("Сохранить как -- сохранить файл предпросмотра в другое место")
            save_as_btn.clicked.connect(functools.partial(self._save_preview_as, slot, variant_id))
            layout.addWidget(save_as_btn)

            reveal_btn = QPushButton()
            reveal_btn.setIcon(icons.icon("files", "#8e8e93", 14))
            reveal_btn.setFixedSize(22, 22)
            reveal_btn.setFlat(True)
            reveal_btn.setToolTip("Показать в Finder")
            reveal_btn.clicked.connect(functools.partial(self._reveal_preview, slot, variant_id))
            layout.addWidget(reveal_btn)

        return field

    def _build_template_drop_hint(self) -> QWidget:
        """Пустое состояние панели шаблона -- пока файл не загружен, вместо
        обычной серой строки текста показывается визуальная drop-зона в том
        же языке, что уже есть у _show_included_block_placeholder()
        (пунктирная рамка, иконка "drag-drop", центрированный текст) --
        пользователь попросил именно этот стиль. Перетаскивание .docx
        работает в любом месте панели независимо от этой рамки (см.
        setAcceptDrops на container в _build_placeholder_catalog()),
        рамка тут только делает зону видимой и даёт понятную подсказку --
        сама по себе она drop-события не обрабатывает."""
        box = QWidget()
        box.setObjectName("titleTemplateDropHint")
        layout = QVBoxLayout(box)
        layout.setContentsMargins(12, 16, 12, 16)
        layout.setSpacing(6)
        icon_label = QLabel()
        icon_label.setPixmap(icons.render("drag-drop", "#5a5a5c", 20))
        icon_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(icon_label)
        text_label = QLabel("Перетащите документ для Word-шаблона сюда для загрузки")
        text_label.setWordWrap(True)
        text_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        text_label.setStyleSheet("color: #5a5a5c; font-size: 12px;")
        layout.addWidget(text_label)
        return box

    def _build_template_loaded_card(self, slot: str, variant_id: str, template_filename: str) -> QWidget:
        """Состояние "файл уже загружен" -- вместо прежней обычной серой
        строки текста ("Файл шаблона: X") видимая карточка в том же языке,
        что и карточки вариантов в сайдбаре (available_list::item --
        тёмный фон, скруглённая рамка, иконка документа слева), плюс
        крестик справа для отвязки файла (_remove_variant_template()).
        Перетаскивание нового .docx поверх карточки по-прежнему работает
        (весь container остаётся drop-зоной, см. цикл setAcceptDrops в
        _build_placeholder_catalog()) и штатно спросит подтверждение на
        замену, как и раньше.

        Клик по самой карточке -- открывает файл варианта как есть
        (_open_variant_raw_file()), тот же результат, что раньше давала
        отдельная кнопка «Посмотреть шаблон» (убрана по просьбе
        пользователя). icon_label/text_label сделаны прозрачными для мыши
        (WA_TransparentForMouseEvents), иначе клик по ним не долетал бы до
        card.mousePressEvent -- тот же приём, что уже используется для
        строки перетащенного блока в _process_block_drop(). remove_btn
        остаётся ОБЫЧНЫМ (не прозрачным) виджетом -- его собственный клик
        по-прежнему обрабатывается им самим, до card не долетает, поэтому
        крестик не запускает просмотр файла."""
        card = QWidget()
        card.setObjectName("titleTemplateLoadedCard")
        card.setCursor(Qt.CursorShape.PointingHandCursor)
        card.setToolTip(
            "Открыть сам файл варианта как есть -- в том состоянии, в котором его оставили в Word "
            "по завершении редактирования"
        )
        card.mousePressEvent = lambda event, s=slot, vid=variant_id: self._open_variant_raw_file(s, vid)
        layout = QHBoxLayout(card)
        layout.setContentsMargins(9, 7, 6, 7)
        layout.setSpacing(8)

        icon_label = QLabel()
        icon_label.setPixmap(icons.render("file-text", "#0a84ff", 15))
        icon_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        layout.addWidget(icon_label)

        text_label = QLabel(f"Файл шаблона: {template_filename}")
        text_label.setWordWrap(True)
        text_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        layout.addWidget(text_label, stretch=1)

        # «Показать в Finder» -- рядом с крестиком, СВОЯ кнопка, а не пункт
        # контекстного меню (в отличие от objectsTree, см.
        # _show_document_context_menu()/_reveal_in_finder()): карточка
        # варианта ПКМ пока ничем не занята, но кнопка виднее и не
        # требует знать про ПКМ. Открывает файловый менеджер с уже
        # выделенным файлом -- в дополнение к клику по самой карточке
        # (открывает файл В ПРИЛОЖЕНИИ, см. card.mousePressEvent выше), а
        # не вместо него.
        # «Сохранить как» -- рядом с «Показать в Finder», по той же логике
        # (своя кнопка, не пункт меню). В отличие от загрузки/перетаскивания
        # (_install_variant_template(), молча копирует в уже резолвящееся
        # место, без диалога -- см. её докстринг), это единственное место,
        # где оператор ЯВНО решает перенести файл варианта в другую папку
        # (_save_variant_template_as()).
        save_as_btn = QPushButton()
        save_as_btn.setIcon(icons.icon("download", "#8e8e93", 14))
        save_as_btn.setFixedSize(22, 22)
        save_as_btn.setFlat(True)
        save_as_btn.setToolTip("Сохранить как -- перенести файл варианта в другую папку")
        save_as_btn.clicked.connect(functools.partial(self._save_variant_template_as, slot, variant_id))
        layout.addWidget(save_as_btn)

        reveal_btn = QPushButton()
        reveal_btn.setIcon(icons.icon("files", "#8e8e93", 14))
        reveal_btn.setFixedSize(22, 22)
        reveal_btn.setFlat(True)
        reveal_btn.setToolTip("Показать в Finder")
        reveal_btn.clicked.connect(functools.partial(self._reveal_variant_template, slot, variant_id))
        layout.addWidget(reveal_btn)

        remove_btn = QPushButton()
        remove_btn.setIcon(icons.icon("x", "#8e8e93", 14))
        remove_btn.setFixedSize(22, 22)
        remove_btn.setFlat(True)
        remove_btn.setToolTip("Убрать привязанный файл -- вернуться к автосгенерированной заготовке")
        remove_btn.clicked.connect(functools.partial(self._remove_variant_template, slot, variant_id))
        layout.addWidget(remove_btn)

        return card

    def _reveal_variant_template(self, slot: str, variant_id: str):
        """Кнопка-папка на карточке загруженного шаблона -- показывает сам
        .docx-файл варианта в Finder (_reveal_in_finder(), macOS-специфично,
        как и остальные вызовы в приложении), не открывая его ни в
        приложении, ни во внешнем редакторе -- в дополнение к клику по
        карточке (_open_variant_raw_file())."""
        try:
            template_path = self._find_slot_template(slot, variant_id)
        except FileNotFoundError:
            self.show_message(
                "Нет файла шаблона",
                "У этого варианта пока нет сгенерированного .docx-файла.",
                QMessageBox.Icon.Warning,
            )
            return
        self._reveal_in_finder(template_path)

    def _save_variant_template_as(self, slot: str, variant_id: str):
        """«Сохранить как» на карточке загруженного шаблона -- единственное
        место, где оператор ЯВНО переносит .docx-файл варианта в другую
        папку (диалог choose_template_save_path(), src/ui/template_location.py,
        стартует в последней использованной папке). В отличие от загрузки/
        перетаскивания (_install_variant_template()) это не срабатывает
        автоматически -- отдельное намеренное действие. Сохраняет выбранный
        путь в variant.template_path -- дальше find_title_template()/
        find_intro_template() (src/config.py) резолвят файл уже оттуда,
        старая копия на прежнем месте не удаляется (тот же принцип, что и у
        остального кода этой фичи -- см. _delete_variant(), файлы вариантов
        никогда не удаляются с диска автоматически)."""
        try:
            current_path = self._find_slot_template(slot, variant_id)
        except FileNotFoundError:
            self.show_message(
                "Нет файла шаблона",
                "У этого варианта пока нет сгенерированного .docx-файла.",
                QMessageBox.Icon.Warning,
            )
            return

        target_path = choose_template_save_path(self, current_path.name)
        if target_path is None:
            return
        current_path = current_path.resolve()
        if target_path.resolve() == current_path:
            return

        target_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(current_path, target_path)

        store = self._slot_store(slot)
        variants = store.load_variants()
        variant = next((v for v in variants if v.id == variant_id), None)
        if variant is not None:
            variant.template_path = str(target_path)
            store.save_variants(variants)

    def _reveal_preview(self, slot: str, variant_id: str):
        """«Показать в Finder» на поле предпросмотра -- показывает
        последний сгенерированный .docx предпросмотра (лежит рядом с
        файлом шаблона, путь запоминается в self._preview_generated при
        _open_variant_preview()), не открывая его повторно ни в
        приложении, ни во внешнем редакторе -- та же логика, что и у
        _reveal_variant_template() для файла самого варианта."""
        preview_path = self._preview_generated.get((slot, variant_id))
        if preview_path is None or not preview_path.exists():
            self.show_message(
                "Нет файла предпросмотра",
                "Сначала откройте предпросмотр -- файл ещё не сгенерирован.",
                QMessageBox.Icon.Warning,
            )
            return
        self._reveal_in_finder(preview_path)

    def _save_preview_as(self, slot: str, variant_id: str):
        """«Сохранить как» на поле предпросмотра -- копирует последний
        сгенерированный .docx предпросмотра (лежит рядом с самим файлом
        шаблона, см. _open_variant_preview() -- перезаписывается там же
        при следующем клике «Открыть предпросмотр») в постоянное место,
        которое выбирает оператор -- тот же диалог, что и у карточки
        шаблона (choose_template_save_path(), src/ui/template_location.py).
        В отличие от _save_variant_template_as() НЕ трогает
        variant.template_path -- предпросмотр не является шаблоном
        варианта, это отдельный, уже отрендеренный со значениями документ."""
        preview_path = self._preview_generated.get((slot, variant_id))
        if preview_path is None or not preview_path.exists():
            self.show_message(
                "Нет файла предпросмотра",
                "Сначала откройте предпросмотр -- файл ещё не сгенерирован.",
                QMessageBox.Icon.Warning,
            )
            return
        target_path = choose_template_save_path(self, preview_path.name)
        if target_path is None:
            return
        target_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(preview_path, target_path)

    def _remove_variant_template(self, slot: str, variant_id: str):
        """Крестик на карточке загруженного шаблона -- отвязывает файл и
        возвращает панель в состояние "шаблон не загружен" (после
        _render_slot_fields() вместо карточки снова покажется
        _build_template_drop_hint()). Заготовка при этом не просто
        удаляется, а перегенерируется заново (generate_title_fragment()/
        generate_intro_fragment(), тот же вызов, что и при создании нового
        варианта) -- иначе find_title_template()/find_intro_template()
        продолжал бы резолвить СТАРЫЙ файл на диске, и «Открыть
        предпросмотр»/«Посмотреть шаблон» показывали бы отвязанное по
        названию, но реально то же самое содержимое -- вариант никогда не
        должен оставаться "осиротевшим" (см. find_title_template()/
        find_intro_template() в src/config.py). Явный откат к управляемой
        заготовке -- поэтому и путь, выбранный оператором (template_path),
        тоже сбрасывается: заново сгенерированный файл всегда ложится в
        FRAGMENTS_DIR/{slot}_{id}.docx, без диалога выбора места."""
        from ..config import FRAGMENTS_DIR

        store = self._slot_store(slot)
        variants = store.load_variants()
        variant = next((v for v in variants if v.id == variant_id), None)
        if variant is None:
            return

        confirm = QMessageBox.question(
            self, "Убрать шаблон варианта?",
            "Вернуться к автосгенерированной заготовке? Текущий файл, включая любые правки, "
            "сделанные в Word, будет заменён заново сгенерированным.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.Cancel,
        )
        if confirm != QMessageBox.StandardButton.Yes:
            return

        variant.template_filename = ""
        variant.template_path = ""
        store.save_variants(variants)
        self._slot_generate_fragment(slot)(variant, FRAGMENTS_DIR / f"{slot}_{variant_id}.docx")
        self._render_slot_fields(slot, variant_id)

    def _copy_chip(self, widget_name: str, chip: QLabel):
        """Левый клик по плейсхолдеру -- копирует "{{ widget_name }}" в
        буфер обмена (для вставки в открытый в Word файл варианта) и на короткое
        время подсвечивает чип зелёным (QLabel[titleChipCopied="true"], см.
        стили выше) как визуальное подтверждение, что копирование
        произошло. Right-click остался отдельно только для удаления (см.
        _show_chip_context_menu()) -- копирование туда больше не
        привязано."""
        QGuiApplication.clipboard().setText("{{ " + widget_name + " }}")
        chip.setProperty("titleChipCopied", True)
        chip.style().unpolish(chip)
        chip.style().polish(chip)
        QTimer.singleShot(900, lambda w=chip: self._clear_chip_copied(w))

    def _clear_chip_copied(self, chip: QLabel):
        try:
            chip.setProperty("titleChipCopied", False)
            chip.style().unpolish(chip)
            chip.style().polish(chip)
        except RuntimeError:
            pass  # панель могла перестроиться и удалить чип раньше таймера -- не ошибка

    def _build_placeholder_menu(self, slot: str, variant_id: str, field_ids) -> QMenu:
        """Меню «Вставить плейсхолдер» -- по строке на поле общего каталога
        (get_all_field_labels(), см. её докстринг: встроенные
        TITLE_FIELD_LABELS + пользовательские из JSON). Строка кликабельна
        целиком (добавляет плейсхолдер в текущий вариант) и несёт корзину
        справа -- УДАЛЕНИЕ ИЗ КАТАЛОГА теперь доступно для ЛЮБОГО поля, не
        только пользовательского (пользователь явно попросил "для всех" --
        раньше встроенные TITLE_FIELD_LABELS были защищены от удаления по
        аналогии со встроенными вариантами, но здесь это не то же самое:
        "удаление" встроенного поля -- не правка кода, а его id в
        отдельном списке-исключении hidden_builtin_fields, см.
        _delete_catalog_field()/title_variants_store.get_all_field_labels()).
        Удаление действует сразу для всех слотов -- каталог общий.

        Каждая строка -- QWidgetAction с ручной QWidget-разметкой, а не
        обычный QAction (тот не умеет два независимых клика в одной
        строке) -- тот же приём, что уже даёт составные строки в
        includedBlockList/сайдбаре (см. _process_block_drop()): текст --
        обычный QLabel с mousePressEvent, кнопка -- обычный QPushButton, её
        собственный клик не всплывает к обработчику текста. QMenu не
        подсвечивает QWidgetAction-строки синим при наведении сама (в
        отличие от обычных QAction) -- подсвечиваем вручную через
        enter/leaveEvent самого row; раз теперь ВСЕ строки такие, подсветка
        стала единообразной для всего меню (раньше, пока часть строк были
        обычными QAction, а часть -- QWidgetAction, подсвечивались
        синим только первые).

        Сами строки лежат не по одной QWidgetAction на пункт меню
        (было раньше), а ВСЕ вместе в одном QScrollArea, который и есть
        единственная QWidgetAction меню. Каталог полей растёт (каждое «+
        Новое поле» добавляет строку навсегда) и легко перерастает высоту
        экрана -- тогда штатный QMenu сам переходит в режим прокрутки
        (стрелочки вверху/внизу поповера), а этот режим не расчитан на
        пункты с тяжёлыми виджетами (QHBoxLayout + QLabel + QPushButton в
        каждой строке) -- на реальном каталоге в полсотни полей прокрутка
        ощутимо подлагивала. QScrollArea прокручивает свой viewport
        обычным Qt-механизмом (клиппинг + перерисовка только видимой
        части) и с тем же набором строк не лагает вообще.

        Строка поиска -- первым пунктом меню, ДО списка (не рядом с «+
        Новое поле…» внизу -- фильтровать нужно то, что над ней, а не сам
        добавляющий пункт). QLineEdit.textChanged просто скрывает
        (setVisible(False)) строки, чей label не содержит введённый текст
        (регистронезависимо) -- QVBoxLayout сам схлопывает место под
        скрытыми виджетами, отдельного re-layout не требуется. Фокус в
        поле сразу при открытии меню (aboutToShow) -- каталог уже вырос
        до полусотни+ полей (см. выше), пользователю обычно быстрее
        напечатать несколько букв, чем листать список глазами."""
        from ..services.title_variants_store import get_all_field_labels

        # Поля, уже используемые в ДРУГОМ разделе, идут первыми в списке --
        # без этого они вперемешку с ещё не задействованными полями, и
        # пользователь не видит с ходу, что часть каталога уже занята
        # (sorted() -- стабильная сортировка, порядок внутри каждой из двух
        # групп остаётся как в get_all_field_labels()).
        rows = []
        for field_id, label in get_all_field_labels().items():
            already = field_id in field_ids
            # "Уже вставлен" -- строго про ТЕКУЩИЙ вариант этого слота (see
            # docstring), а не про общий каталог -- тот же field_id мог
            # быть отдельно вставлен и в другой слот, это не должно мешать
            # вставить его и сюда. Такое использование в другом слоте --
            # не блокирующая, отдельная пометка (other_slot_labels), а не
            # тот же статус "✓ уже вставлен".
            other_slot_labels = [] if already else self._slots_using_field(slot, field_id)
            rows.append((field_id, label, already, other_slot_labels))
        rows.sort(key=lambda row: 0 if row[3] else 1)

        menu = QMenu(self)

        search_edit = QLineEdit()
        search_edit.setPlaceholderText("Поиск плейсхолдера…")
        search_edit.setClearButtonEnabled(True)
        search_edit.addAction(icons.icon("search", "#8e8e93", 13), QLineEdit.ActionPosition.LeadingPosition)
        search_edit.setStyleSheet(
            "QLineEdit { background: #1c1c1e; border: 0.5px solid #38383a; border-radius: 6px; "
            "color: #e5e5e7; font-size: 12px; padding: 5px 8px; margin: 6px 8px; }"
            "QLineEdit:focus { border-color: #0a84ff; }"
        )
        search_action = QWidgetAction(menu)
        search_action.setDefaultWidget(search_edit)
        menu.addAction(search_action)
        menu.aboutToShow.connect(lambda: QTimer.singleShot(0, search_edit.setFocus))

        list_container = QWidget()
        list_container.setStyleSheet("background: #2c2c2e;")
        list_layout = QVBoxLayout(list_container)
        list_layout.setContentsMargins(0, 0, 0, 0)
        list_layout.setSpacing(0)

        # (виджет строки, label в нижнем регистре) -- используется ниже
        # обработчиком search_edit.textChanged для фильтрации по подстроке.
        searchable_rows = []
        for field_id, label, already, other_slot_labels in rows:
            used_elsewhere = bool(other_slot_labels)
            if already:
                row_text = label + " ✓"
            elif used_elsewhere:
                other_names = "«" + "», «".join(other_slot_labels) + "»"
                row_text = (
                    f'{html.escape(label)} '
                    f'<span style="color:#5a5a5c; font-size:10.5px;">· уже в {html.escape(other_names)}</span>'
                )
            else:
                row_text = label

            row = QWidget()
            row.enterEvent = lambda event, r=row: r.setStyleSheet("background: #0a84ff;")
            row.leaveEvent = lambda event, r=row: r.setStyleSheet("")
            row_layout = QHBoxLayout(row)
            row_layout.setContentsMargins(10, 2, 4, 2)
            row_layout.setSpacing(4)
            text_label = QLabel(row_text)
            if used_elsewhere:
                text_label.setTextFormat(Qt.TextFormat.RichText)
            if already:
                text_label.setStyleSheet("color: #5a5a5c;")
            else:
                text_label.setCursor(Qt.CursorShape.PointingHandCursor)

                def _add_and_close(event, fid=field_id, m=menu):
                    self._add_variant_placeholder(slot, variant_id, fid)
                    m.close()

                text_label.mousePressEvent = _add_and_close
            row_layout.addWidget(text_label, stretch=1)
            delete_btn = QPushButton()
            delete_btn.setIcon(icons.icon("trash", "#8e8e93", 12))
            delete_btn.setFixedSize(22, 22)
            delete_btn.setFlat(True)
            delete_btn.setToolTip("Удалить поле из каталога")
            delete_btn.enterEvent = lambda event, b=delete_btn: b.setIcon(icons.icon("trash", "#ff453a", 12))
            delete_btn.leaveEvent = lambda event, b=delete_btn: b.setIcon(icons.icon("trash", "#8e8e93", 12))

            def _delete_and_reopen(checked=False, fid=field_id, lbl=label, m=menu):
                # m.pos() -- ПОКА меню ещё реально открыто (сам клик по
                # корзинке пришёл из его же виджета), захватываем позицию
                # ДО вызова _delete_catalog_field(): там QMessageBox.question()
                # -- модальный, и уже само его появление молча закрывает
                # открытый QMenu (стандартное поведение Qt-попапов -- любое
                # другое окно, становясь активным, закрывает все открытые
                # попапы). См. _reopen_insert_menu() -- он открывает меню
                # заново на этом месте, что и даёт эффект "список не
                # закрылся".
                pos = m.pos()
                self._delete_catalog_field(slot, variant_id, fid, lbl)
                QTimer.singleShot(
                    0, lambda: QTimer.singleShot(0, functools.partial(self._reopen_insert_menu, slot, pos))
                )

            delete_btn.clicked.connect(_delete_and_reopen)
            row_layout.addWidget(delete_btn)
            list_layout.addWidget(row)
            searchable_rows.append((row, label.lower()))

        # QScrollArea ниже -- setWidgetResizable(True), поэтому list_container
        # растягивается минимум под высоту вьюпорта. Без завершающего
        # stretch лишнее место (когда после фильтра видимых строк мало)
        # распределялось бы QVBoxLayout между самими строками -- каждая
        # растягивалась бы на всю оставшуюся высоту (видно по результату
        # поиска: 2-3 строки вместо компактного списка занимали всю
        # видимую область). Stretch забирает остаток на себя, строки
        # остаются естественной высоты и прижаты к верху.
        list_layout.addStretch(1)

        def _filter_rows(text: str, entries=searchable_rows):
            needle = text.strip().lower()
            for row_widget, label_lower in entries:
                row_widget.setVisible(not needle or needle in label_lower)

        search_edit.textChanged.connect(_filter_rows)

        scroll_area = QScrollArea()
        scroll_area.setWidget(list_container)
        scroll_area.setWidgetResizable(True)
        scroll_area.setFrameShape(QFrame.Shape.NoFrame)
        scroll_area.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        # На macOS по умолчанию (системная настройка "показывать полосы
        # прокрутки: при прокрутке") скроллбар оверлейный -- рисуется
        # ПОВЕРХ содержимого, без выделенного места под него, из-за чего
        # корзинки у правого края строк оказывались частично под ним.
        # Собственный QSS для QScrollBar переключает Qt с нативной
        # NSScrollView-полосы на обычную стилизованную -- та всегда
        # занимает свою полосу пространства во viewport, содержимое
        # подвинется, а не перекроется.
        scroll_area.setStyleSheet(
            "QScrollArea { background: #2c2c2e; border: none; }"
            "QScrollBar:vertical { background: transparent; width: 10px; margin: 2px 1px; }"
            "QScrollBar::handle:vertical { background: #48484a; border-radius: 5px; min-height: 24px; }"
            "QScrollBar::handle:vertical:hover { background: #5a5a5c; }"
            "QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0; }"
            "QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical { background: none; }"
        )
        # QAbstractScrollArea.sizeHint() по умолчанию не связан с реальным
        # содержимым viewport -- без явной ширины QMenu мог бы отвести под
        # список меньше места, чем нужно самой длинной строке (особенно с
        # припиской "· уже в «...»"), и она бы обрезалась/переносилась.
        scroll_area.setMinimumWidth(list_container.sizeHint().width())
        # 420px -- примерно 12 строк, дальше уже сам QScrollArea, не QMenu.
        scroll_area.setMaximumHeight(420)
        list_action = QWidgetAction(menu)
        list_action.setDefaultWidget(scroll_area)
        menu.addAction(list_action)

        menu.addSeparator()
        new_field_action = menu.addAction(icons.icon("plus", "#0a84ff", 13), "Новое поле…")
        new_field_action.triggered.connect(
            lambda checked=False: self._create_and_add_variant_field(slot, variant_id)
        )
        return menu

    def _reopen_insert_menu(self, slot: str, pos):
        """Открывает меню «Вставить плейсхолдер» заново на месте pos --
        вызывается с задержкой после удаления поля через корзинку (см.
        _delete_and_reopen() внутри _build_placeholder_menu()), чтобы
        закрытие меню модальным QMessageBox.question() из
        _delete_catalog_field() не выглядело для оператора так, будто
        список закрылся -- иначе после каждого удаления пришлось бы заново
        кликать «Вставить плейсхолдер», чтобы удалить следующее поле.

        Кнопка ищется по имени (insertPlaceholderBtn_{slot}, см.
        _build_placeholder_catalog()), а не берётся по ссылке, сохранённой
        в замыкании -- если поле реально удалено, _render_slot_fields()
        (тоже отложенная на singleShot(0, ...), см. _delete_catalog_field())
        успевает перестроить fields_layout ДО этого вызова (двойной
        singleShot(0, ...) у места вызова гарантирует порядок) и завести
        СОВСЕМ ДРУГОЙ QToolButton с новым меню; при отмене удаления
        (ответ "Нет") ничего не перестраивалось -- кнопка та же, что и
        была, просто снова показываем то же самое меню."""
        fields_panel = self._slot_widget(slot, "fields_panel")
        btn = fields_panel.findChild(QToolButton, f"insertPlaceholderBtn_{slot}")
        if btn is not None and btn.menu() is not None:
            btn.menu().popup(pos)

    def _delete_catalog_field(self, slot: str, variant_id: str, field_id: str, label: str):
        """Корзина у поля в меню «Вставить плейсхолдер» -- удаляет поле из
        ОБЩЕГО каталога, а не из реквизитов какого-то конкретного варианта.
        Поле пропадёт из списка «Вставить плейсхолдер» для ЛЮБОГО варианта
        (титульных листов, вводной части, приложения 1 -- каталог общий,
        см. _build_placeholder_menu()). Уже вставленный в чей-то вариант
        плейсхолдер с этим id при этом тоже удаляется -- из
        variant.subtitle_fields ВСЕХ вариантов всех слотов (не только
        текущего), чтобы висящий чип с подписью-фоллбэком на голый id
        (см. _render_slot_fields()) не оставался в интерфейсе после того,
        как поле пропало из каталога.

        Работает и для пользовательских полей (заведённых через «+ Новое
        поле…» -- field_catalog в JSON, физически удаляется), и для
        встроенных (TITLE_FIELD_LABELS -- сама константа в коде не
        правится, вместо этого id добавляется в hidden_builtin_fields,
        список-исключение, который get_all_field_labels() вычитает при
        сборке итогового каталога, см. title_variants_store.py).

        Формула этого поля (если была, см. field_formulas/
        _open_formula_editor()) удаляется вместе с ним -- без самого поля
        в каталоге она бессмысленна. Формулы ДРУГИХ полей, ссылавшихся на
        field_id, не правятся -- просто перестанут считаться
        (evaluate_formula() вернёт None, "—"), тот же принцип деградации,
        что и у голого id вместо подписи. Таблица этого поля (field_tables/
        _open_table_editor()) удаляется тем же принципом; чипы этого поля
        ВНУТРИ чужих таблиц (вставленные через «Вставить плейсхолдер» в
        редакторе таблиц) не правятся -- останутся ссылкой на пропавший id,
        тот же принцип деградации.

        _render_slot_fields() (перестраивает fields_layout, включая
        свежий _build_placeholder_menu() без удалённого поля) отложена на
        следующий тик цикла событий -- тот же приём и по той же причине,
        что и в _open_variant_preview(): вызов идёт из delete_btn.clicked,
        а сам delete_btn -- часть виджета текущего fields_layout, который
        _render_slot_fields() тут же уничтожит; синхронный вызов означал
        бы удаление виджета прямо изнутри обработки его же клика."""
        answer = QMessageBox.question(
            self, "Удалить поле?",
            f"Удалить поле «{label}» из каталога плейсхолдеров? Оно пропадёт из списка «Вставить "
            "плейсхолдер» для всех вариантов (титульных листов, вводной части, приложения 1) — "
            "каталог общий.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return

        from ..services.title_variants_store import (
            load_field_catalog, save_field_catalog,
            load_hidden_builtin_fields, save_hidden_builtin_fields,
            load_field_formulas, save_field_formulas,
            load_field_tables, save_field_tables,
        )

        catalog = load_field_catalog()
        if field_id in catalog:
            catalog.pop(field_id, None)
            save_field_catalog(catalog)
        else:
            hidden = load_hidden_builtin_fields()
            if field_id not in hidden:
                hidden.append(field_id)
                save_hidden_builtin_fields(hidden)

        formulas = load_field_formulas()
        if field_id in formulas:
            formulas.pop(field_id, None)
            save_field_formulas(formulas)

        tables = load_field_tables()
        if field_id in tables:
            tables.pop(field_id, None)
            save_field_tables(tables)

        for slot_key in self._CONSTRUCTOR_SLOTS:
            slot_store = self._slot_store(slot_key)
            variants = slot_store.load_variants()
            changed = False
            for variant in variants:
                if field_id in variant.subtitle_fields:
                    variant.subtitle_fields.remove(field_id)
                    changed = True
            if changed:
                slot_store.save_variants(variants)

        QTimer.singleShot(0, functools.partial(self._render_slot_fields, slot, variant_id))

    def _ensure_fragment_exists(self, slot: str, variant):
        """Сеть безопасности -- генерирует .docx-заготовку варианта, только
        если файла ещё нет (создан не был или удалён вручную с диска).
        НЕ вызывается безусловно -- иначе затирала бы ручную правку,
        сделанную пользователем в Word (см. generate_title_fragment()/
        generate_intro_fragment()). Целевой путь -- тот же, что и при
        обычной резолюции (_variant_fragment_path()): если у варианта
        выбрано своё место хранения (template_path), перегенерирует именно
        там, а не в FRAGMENTS_DIR."""
        fragment_path = self._variant_fragment_path(slot, variant)
        if not fragment_path.exists():
            self._slot_generate_fragment(slot)(variant, fragment_path)

    def _add_variant_placeholder(self, slot: str, variant_id: str, field_id: str):
        store = self._slot_store(slot)
        variants = store.load_variants()
        variant = next((v for v in variants if v.id == variant_id), None)
        if variant is None or field_id in variant.subtitle_fields:
            return
        variant.subtitle_fields.append(field_id)
        store.save_variants(variants)
        self._ensure_fragment_exists(slot, variant)
        self._render_slot_fields(slot, variant_id)

    def _remove_variant_placeholder(self, slot: str, variant_id: str, field_id: str):
        store = self._slot_store(slot)
        variants = store.load_variants()
        variant = next((v for v in variants if v.id == variant_id), None)
        if variant is None:
            return
        variant.subtitle_fields = [f for f in variant.subtitle_fields if f != field_id]
        store.save_variants(variants)
        self._render_slot_fields(slot, variant_id)

    def _create_and_add_variant_field(self, slot: str, variant_id: str):
        from ..services.title_variants_store import load_field_catalog, save_field_catalog

        label, ok = QInputDialog.getText(self, "Новое поле", "Название поля:")
        label = label.strip()
        if not ok or not label:
            return
        field_id = "field_" + uuid4().hex[:8]
        catalog = load_field_catalog()
        catalog[field_id] = label
        save_field_catalog(catalog)
        self._add_variant_placeholder(slot, variant_id, field_id)

    def _show_chip_context_menu(self, slot: str, variant_id: str, field_id: str, global_pos):
        """ПКМ по плейсхолдеру -- «Создать формулу»/«Редактировать формулу»,
        «Переименовать» и «Удалить» (копирование по-прежнему по левому
        клику, см. _copy_chip()).

        «Создать формулу» -- ПЕРВЫЙ пункт, всегда доступен (формула --
        свойство самого поля-в-документе, привязанное к field_id в общем
        каталоге, а не пользовательского/встроенного статуса поля, в
        отличие от «Переименовать» ниже, см. _open_formula_editor()).
        Подпись переключается на «Редактировать формулу», если формула для
        этого field_id уже задана.

        «Переименовать» меняет подпись поля в ОБЩЕМ каталоге
        (field_catalog, см. _rename_catalog_field()) -- у чипа нет
        собственной, per-вариантной подписи, label всегда берётся из
        get_all_field_labels(), поэтому переименование не может быть
        чем-то иным, кроме правки каталога. Доступно только для
        пользовательских полей (заведённых через «+ Новое поле…») -- тот
        же принцип, что уже действует для удаления поля из каталога в меню
        «Вставить плейсхолдер» (_build_placeholder_menu()): встроенные
        TITLE_FIELD_LABELS через UI не редактируются.

        «Создать таблицу» -- тем же принципом, что и «Создать формулу»
        выше: всегда доступно, подпись переключается на «Редактировать
        таблицу» для уже заполненных полей (см. _open_table_editor()),
        хранится по field_id в общем каталоге (field_tables). В отличие от
        формулы, таблица не сводится к одному вычисленному значению --
        get_form_data() подставляет вместо неё маркер, который на
        настоящую .docx-таблицу меняет уже ПОСЛЕ рендера
        _splice_table_placeholders() (см. её докстринг).

        «Удалить» -- убирает плейсхолдер из РЕКВИЗИТОВ этого варианта
        (variant.subtitle_fields), а не из каталога -- другая операция,
        см. _remove_variant_placeholder()."""
        from ..services.title_variants_store import load_field_catalog, load_field_formulas, load_field_tables

        menu = QMenu(self)
        formula_label = "Редактировать формулу" if field_id in load_field_formulas() else "Создать формулу"
        formula_action = menu.addAction(icons.icon("formula", "#bf5af2", 13), formula_label)
        table_label = "Редактировать таблицу" if field_id in load_field_tables() else "Создать таблицу"
        table_action = menu.addAction(icons.icon("table", "#64d2ff", 13), table_label)
        rename_action = None
        if field_id in load_field_catalog():
            rename_action = menu.addAction(icons.icon("edit", "#c7c7cc", 13), "Переименовать")
        delete_action = menu.addAction(icons.icon("trash", "#ff453a", 13), "Удалить")
        chosen = menu.exec(global_pos)
        if chosen == formula_action:
            self._open_formula_editor(slot, variant_id, field_id)
        elif chosen == table_action:
            self._open_table_editor(slot, variant_id, field_id)
        elif chosen == delete_action:
            self._remove_variant_placeholder(slot, variant_id, field_id)
        elif rename_action is not None and chosen == rename_action:
            self._rename_catalog_field(slot, variant_id, field_id)

    def _open_formula_editor(self, slot: str, variant_id: str, field_id: str):
        """Открывает FormulaEditorDialog для field_id -- на успешном
        сохранении/удалении правит ОБЩИЙ каталог формул (field_formulas,
        title_variants_store.py), а не что-то в этом конкретном варианте:
        формула, как и подпись поля, действует везде, где встречается этот
        плейсхолдер (см. _show_chip_context_menu()). После любого исхода,
        кроме «Отмена», перестраивает ОБА слота (не только текущий) -- в
        другом слоте могла быть открыта форма с тем же полем."""
        from ..services.title_variants_store import get_all_field_labels, load_field_formulas, save_field_formulas

        formulas = load_field_formulas()
        dialog = FormulaEditorDialog(
            field_id=field_id,
            field_label=get_all_field_labels().get(field_id, field_id),
            all_fields=get_all_field_labels(),
            all_formulas=formulas,
            existing_formula=formulas.get(field_id),
            resolve_placeholder=self._placeholder_numeric_value,
            parent=self,
        )
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return

        if dialog.removed:
            formulas.pop(field_id, None)
        else:
            formulas[field_id] = {"tokens": dialog.tokens, "decimals": dialog.decimals}
        save_field_formulas(formulas)

        for refresh_slot in self._CONSTRUCTOR_SLOTS:
            refresh_variant_id = self._filled_slot_variant(refresh_slot)
            if refresh_variant_id is not None:
                self._render_slot_fields(refresh_slot, refresh_variant_id)

    def _open_table_editor(self, slot: str, variant_id: str, field_id: str):
        """Открывает TableEditorDialog для field_id -- на успешном
        сохранении/удалении правит ОБЩИЙ каталог таблиц (field_tables,
        title_variants_store.py), тем же охватом, что и формула (см.
        _open_formula_editor()): действует везде, где встречается этот
        плейсхолдер, а не только в этом варианте."""
        from ..services.title_variants_store import (
            get_all_field_labels, load_field_formulas, load_field_tables, save_field_tables,
        )

        tables = load_field_tables()
        dialog = TableEditorDialog(
            field_id=field_id,
            field_label=get_all_field_labels().get(field_id, field_id),
            all_fields=get_all_field_labels(),
            all_formulas=load_field_formulas(),
            all_tables=tables,
            existing_table=tables.get(field_id),
            create_field=self._create_catalog_field_only,
            parent=self,
        )
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return

        if dialog.removed:
            tables.pop(field_id, None)
        else:
            tables[field_id] = {"has_header": dialog.has_header, "rows": dialog.rows}
        save_field_tables(tables)

        for refresh_slot in self._CONSTRUCTOR_SLOTS:
            refresh_variant_id = self._filled_slot_variant(refresh_slot)
            if refresh_variant_id is not None:
                self._render_slot_fields(refresh_slot, refresh_variant_id)

    def _create_catalog_field_only(self, label: str) -> str:
        """Заводит новое поле ТОЛЬКО в общем каталоге (field_catalog), не
        вставляя его в реквизиты никакого варианта -- в отличие от
        _create_and_add_variant_field() (кнопка «+ Новое поле…» в
        «Вставить плейсхолдер» реквизитов, которая сразу добавляет поле и в
        текущий вариант). Нужен редактору таблиц: поле, заведённое «на
        лету» прямо в ячейке, должно быть доступно для вставки в ЛЮБУЮ
        другую таблицу/формулу через общий каталог, но не обязано само по
        себе становиться отдельным плейсхолдером в чьих-то реквизитах."""
        from ..services.title_variants_store import load_field_catalog, save_field_catalog

        field_id = "field_" + uuid4().hex[:8]
        catalog = load_field_catalog()
        catalog[field_id] = label
        save_field_catalog(catalog)
        return field_id

    def _rename_catalog_field(self, slot: str, variant_id: str, field_id: str):
        """«Переименовать» из контекстного меню чипа -- правит подпись
        ПОЛЬЗОВАТЕЛЬСКОГО поля в общем каталоге (field_catalog), общем для
        обоих слотов и всех вариантов (см. _show_chip_context_menu()). Не
        трогает сам field_id/плейсхолдер "{{ ... }}" -- меняется только
        человекочитаемое название, которое видит оператор."""
        from ..services.title_variants_store import get_all_field_labels, load_field_catalog, save_field_catalog

        current_label = get_all_field_labels().get(field_id, field_id)
        new_label, ok = QInputDialog.getText(
            self, "Переименовать поле", "Новое название:", text=current_label,
        )
        new_label = new_label.strip()
        if not ok or not new_label or new_label == current_label:
            return

        catalog = load_field_catalog()
        if field_id not in catalog:
            return  # встроенное поле -- сюда попасть не должно (см. _show_chip_context_menu())
        catalog[field_id] = new_label
        save_field_catalog(catalog)
        self._render_slot_fields(slot, variant_id)

    def _upload_variant_template(self, slot: str, variant_id: str):
        """«Загрузить шаблон Word» -- диалог выбора файла, пользователь
        указывает ЛЮБОЙ существующий .docx (например, уже готовый
        реальный отчёт), в который собирается вставлять плейсхолдеры.
        Сама привязка файла (без копирования -- см. _install_variant_template(),
        подтверждение замены, запрос приложения для редактирования) вынесена в
        _install_variant_template() -- та же логика нужна и для
        перетаскивания .docx прямо на панель каталога плейсхолдеров (см.
        setAcceptDrops в _build_placeholder_catalog()), кнопка при этом
        остаётся как равноправный способ, не заменяется перетаскиванием."""
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Выбрать документ Word", "", "Документы Word (*.docx)"
        )
        if not file_path:
            return
        self._install_variant_template(slot, variant_id, Path(file_path))

    def _install_variant_template(self, slot: str, variant_id: str, source_path: Path):
        """Общая часть установки .docx как шаблона варианта -- БЕЗ копии:
        variant.template_path указывает прямо на выбранный файл, в том
        месте и под тем именем, где он реально лежит у пользователя (см.
        обсуждение задачи -- копия в FRAGMENTS_DIR под переименованным
        title_{id}.docx путала пользователя: «Показать в Finder»/открытие
        в Word показывали не тот файл, что он выбрал). Дальше рендер при
        сборке документа (_calculate_constructor()) и «Показать в
        Finder»/«Открыть предпросмотр» резолвят СТРОГО тот же путь
        (find_title_template()/find_intro_template(), см. src/config.py).
        open_with_prompt() затем спрашивает, каким приложением
        редактировать (см. src/ui/open_with.py) -- открывает и правит
        ИМЕННО этот файл, а не его копию: если это уже готовый реальный
        документ пользователя, правки в Word сразу применяются к нему.
        Явная копия на новое место -- отдельное, осознанное действие
        («Сохранить как», _save_variant_template_as())."""
        store = self._slot_store(slot)
        variants = store.load_variants()
        variant = next((v for v in variants if v.id == variant_id), None)
        if variant is None:
            return

        source_path = source_path.resolve()
        current_path = self._variant_fragment_path(slot, variant)
        if variant.template_filename and current_path.exists() and current_path.resolve() != source_path:
            confirm = QMessageBox.question(
                self, "Привязать другой файл варианта?",
                f"У варианта уже привязан файл ({current_path.name}). Привязать вместо него выбранный "
                "файл?\n\nПрежний файл на диске не изменится и не удалится -- просто перестанет "
                "использоваться этим вариантом.",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.Cancel,
            )
            if confirm != QMessageBox.StandardButton.Yes:
                return

        variant.template_filename = source_path.name
        variant.template_path = str(source_path)
        store.save_variants(variants)

        open_with_prompt(source_path, self)
        self._render_slot_fields(slot, variant_id)

    @staticmethod
    def _docx_drop_urls(mime_data):
        """Из перетаскиваемых данных -- только локальные .docx-файлы (папки,
        не-Word файлы, файлы из другого приложения без локального пути и
        т.п. игнорируются -- перетаскивание тогда просто не подсвечивается
        и не срабатывает, без сообщения об ошибке)."""
        if not mime_data.hasUrls():
            return []
        return [
            url for url in mime_data.urls()
            if url.isLocalFile() and url.toLocalFile().lower().endswith(".docx")
        ]

    def _refresh_title_template_drop_style(self, drop_zone: QWidget):
        """unpolish/polish только на drop_zone (container) не хватает --
        QSS-правило подсветки нацелено на ВЛОЖЕННЫЙ виджет
        (titleTemplateDropHint, через compound-селектор
        "QWidget#titleTemplateDropZone[dragActive] QWidget#titleTemplateDropHint"),
        а Qt не инвалидирует закэшированный стиль потомка только из-за
        того, что у родителя поменялось свойство -- перерисовать явно
        нужно и сам этот потомок (если он вообще есть: в состоянии "файл
        уже загружен" рамки нет, там просто текст, обновлять нечего).
        Это и была причина, почему функционально перетаскивание уже
        работало (файл копировался, диалог открывался), а подсветка -- нет
        (пользователь подтвердил на реальном запуске)."""
        drop_zone.style().unpolish(drop_zone)
        drop_zone.style().polish(drop_zone)
        drop_zone.update()
        for hint_box in drop_zone.findChildren(QWidget, "titleTemplateDropHint"):
            hint_box.style().unpolish(hint_box)
            hint_box.style().polish(hint_box)
            hint_box.update()

    def _title_template_drag_enter(self, event, drop_zone: QWidget):
        if self._docx_drop_urls(event.mimeData()):
            event.acceptProposedAction()
            drop_zone.setProperty("dragActive", True)
            self._refresh_title_template_drop_style(drop_zone)
        else:
            event.ignore()

    def _title_template_drag_move(self, event):
        """Без этого обработчика перетаскивание "гаснет" сразу после входа
        в виджет: базовая QWidget.dragMoveEvent() ничего не принимает
        (event.ignore() по умолчанию), и Qt показывает курсор "нельзя
        бросить" на КАЖДОЕ движение мыши внутри области, даже если
        dragEnterEvent секундой раньше согласился -- accept там разовый, не
        держит область "валидной" на всё время перетаскивания. Без этого
        подсветка из dragEnterEvent виднелась только на долю секунды входа
        в область и dropEvent часто вообще не успевал сработать."""
        if self._docx_drop_urls(event.mimeData()):
            event.acceptProposedAction()
        else:
            event.ignore()

    def _title_template_drag_leave(self, event, drop_zone: QWidget):
        drop_zone.setProperty("dragActive", False)
        self._refresh_title_template_drop_style(drop_zone)

    def _title_template_drop(self, event, slot: str, variant_id: str, drop_zone: QWidget):
        drop_zone.setProperty("dragActive", False)
        self._refresh_title_template_drop_style(drop_zone)
        urls = self._docx_drop_urls(event.mimeData())
        if not urls:
            event.ignore()
            return
        event.acceptProposedAction()
        self._install_variant_template(slot, variant_id, Path(urls[0].toLocalFile()))

    def _find_slot_template(self, slot: str, variant_id: str):
        from ..config import find_appendix_template, find_intro_template, find_section_template, find_title_template
        if slot == "title":
            return find_title_template(variant_id)
        if slot == "intro":
            return find_intro_template(variant_id)
        if slot == "appendix1":
            return find_appendix_template(variant_id)
        return find_section_template(slot, variant_id)

    def _open_variant_preview(self, slot: str, variant_id: str):
        """«Открыть предпросмотр» -- рендерит .docx-файл варианта
        (find_title_template()/find_intro_template(), тот же файл, что
        правится в Word через «Загрузить шаблон Word») с уже введёнными в
        форму значениями -- тот же get_form_data()/DocxTemplate.render(),
        что и «Собрать документ» (_calculate_constructor()). В отличие от
        самого файла варианта (там буквально "{{ field }}"), тут
        плейсхолдеры уже подставлены реальными значениями -- посмотреть,
        как будет выглядеть готовый документ, не собирая его окончательно.

        Сохраняется РЯДОМ с самим файлом шаблона (template_path.parent),
        под понятным именем "<имя шаблона>_предпросмотр.docx" -- та же
        причина, что и у отказа от копии в FRAGMENTS_DIR для самого
        шаблона (см. _install_variant_template()): раньше рендерился в
        новый tempfile.mkdtemp() при каждом клике, и «Показать в
        Finder» открывал случайную папку вроде "title_preview_k_c4ewb2" --
        непонятно и негде искать повторно. Один и тот же путь на каждый
        клик означает перезапись поверх предыдущего предпросмотра -- если
        файл всё ещё открыт в Word/Pages, tpl.save() упадёт с ошибкой
        доступа, отловлено ниже отдельно, а не падает необработанным.

        _splice_table_placeholders() -- ОБЯЗАТЕЛЬНО после render(), до
        save(): табличные поля (field_tables) рендерятся в маркер (см.
        get_form_data()), а не сразу в готовую таблицу -- превратить его в
        настоящую .docx-таблицу можно только когда маркер уже есть в
        дереве документа."""
        try:
            template_path = self._find_slot_template(slot, variant_id)
        except FileNotFoundError:
            self.show_message(
                "Нет файла шаблона",
                "У этого варианта пока нет сгенерированного .docx-файла.",
                QMessageBox.Icon.Warning,
            )
            return

        try:
            tpl = DocxTemplate(template_path)
            tpl.render(self.get_form_data())
            self._splice_table_placeholders(tpl.docx)
        except Exception as e:
            self.show_message(
                "Не удалось открыть шаблон",
                f"Файл {template_path.name} повреждён или не является корректным .docx "
                f"(например, в нём есть ссылка на картинку, которой физически нет в архиве) -- "
                f"откройте его в Word и пересохраните.\n\nОшибка: {e}",
                QMessageBox.Icon.Critical,
            )
            return

        preview_path = template_path.with_name(f"{template_path.stem}_предпросмотр{template_path.suffix}")
        try:
            tpl.save(str(preview_path))
        except OSError:
            self.show_message(
                "Не удалось сохранить предпросмотр",
                f"Файл {preview_path.name} сейчас открыт в другой программе -- закройте его и "
                "попробуйте снова.",
                QMessageBox.Icon.Warning,
            )
            return
        open_with_prompt(preview_path, self)

        # Поле «Открыть предпросмотр итогового документа»
        # (_build_preview_field()) переключается из приглушённого "Здесь
        # будет..." в активный вид только после успешного рендера (не на
        # каждый клик -- если find_*_template() выше упал, помечать
        # нечего). variant_id может относиться и к встроенному варианту --
        # у него fields_layout строится тем же _render_slot_fields(),
        # просто без каталога плейсхолдеров, так что вызов безопасен для
        # любого variant_id.
        #
        # QTimer.singleShot(0, ...), а не прямой вызов -- сюда попадают
        # ИЗ field.mousePressEvent() того самого поля, которое
        # _render_slot_fields() тут же удалит через layout.removeRow(0)
        # (перестраивает fields_layout с нуля). Прямой вызов означал бы
        # удаление виджета прямо изнутри его же обработчика события, ещё
        # до того как Qt закончит обработку самого mousePressEvent -- то
        # же класс проблемы, что уже решён похожим приёмом в
        # _on_block_dropped()/_process_block_drop() (там -- по другой
        # причине, но тот же instrument: отложить на следующий тик).
        self._preview_generated[(slot, variant_id)] = preview_path
        QTimer.singleShot(0, functools.partial(self._render_slot_fields, slot, variant_id))

    def _open_variant_raw_file(self, slot: str, variant_id: str):
        """«Посмотреть шаблон» -- открывает сам .docx-файл варианта как
        есть (find_title_template()/find_intro_template()), без рендера
        значений -- ровно в том состоянии, в котором его оставили в Word
        по завершении редактирования (буквально "{{ field }}" там, где
        плейсхолдер не относится к уже сохранённым правкам). В отличие от
        «Открыть предпросмотр» (_open_variant_preview()) ничего не
        рендерит и не копирует во временный файл -- открывает напрямую тот
        же файл, что правится через «Загрузить шаблон Word»."""
        try:
            template_path = self._find_slot_template(slot, variant_id)
        except FileNotFoundError:
            self.show_message(
                "Нет файла шаблона",
                "У этого варианта пока нет сгенерированного .docx-файла.",
                QMessageBox.Icon.Warning,
            )
            return
        open_with_prompt(template_path, self)

    def show_message(self, title, text, icon=QMessageBox.Icon.Information):
        """Универсальный метод показа сообщений"""
        msg = QMessageBox()
        msg.setWindowTitle(title)
        msg.setText(text)
        msg.setIcon(icon)
        msg.setStandardButtons(QMessageBox.StandardButton.Ok)
        msg.exec()

    def _check_prerequisite(self, step: str) -> bool:
        """Проверяет, что предыдущий шаг из STEP_ORDER выполнен.

        Если нет — показывает предупреждение вместо тихого сбоя
        (см. STEP_ORDER) и возвращает False.
        """
        idx = self.STEP_ORDER.index(step)
        if idx == 0:
            return True
        prev = self.STEP_ORDER[idx - 1]
        if prev not in self._completed_steps:
            self.show_message(
                "Нарушен порядок действий",
                f"Сначала выполните шаг {self.STEP_LABELS[prev]}.",
                QMessageBox.Icon.Warning,
            )
            return False
        return True

    # --- Методы импорта/экспорта ---

    def import_csv(self):
        """Импорт баллонов из CSV файла."""
        if self.file_handler:
            self.file_handler.import_csv_balloon_list()

    def export_csv(self):
        """Экспорт баллонов в CSV файл."""
        if self.file_handler:
            self.file_handler.export_csv_balloon_list()

    def save_project(self):
        """Сохранение проекта в JSON файл."""
        if self.file_handler:
            self.file_handler.save_project_json()

    def open_project(self):
        """Загрузка проекта из JSON файла."""
        if self.file_handler:
            self.file_handler.open_project_json()

    def fill_table(self):
        """Заполнение таблицы баллонов. Заполняется 1й столбец!!!"""
        self.text = (self.zav_nums.toPlainText()).split(', ')
        amount = self.amount.value()
        table = self.table_ballons
        if len(self.text) == amount:
            table.setRowCount(amount)
            for row, zav_num in enumerate(self.text):
                item = QTableWidgetItem(zav_num)
                table.setItem(row, 0, item)
            self._completed_steps.add("amount")

        else:
            print(f"Кол-во баллонов {amount} не совпадает с введёными зав. №№ {len(self.text)}")

        self.data.update({"tables": [{"num": str(i + 1)} for i in range(len(self.text))]})

    def s_min_min_calc(self):
        """Вычисляет минимальное значение из второго столбца и выводит в QPlainTextEdit"""
        if not self._check_prerequisite("s_min_min"):
            return

        self.s_min_lst = []
        years_min_max_lst = []
        table = self.table_ballons

        # Проверка наличия данных в self.text
        if not self.text:
            print("Ошибка: список заводских номеров пуст. Нажмите кнопку 'Количество'")
            self.s_min_total.setPlainText("Нет данных")
            return

        # Собираем все числовые значения из второго столбца
        for row in range(table.rowCount()):
            item = table.item(row, 1)
            if item is not None and item.text():
                try:
                    self.s_min_lst.append(parse_ru(item.text()))
                except ValueError:
                    print(f"Пропуск нечислового значения в строке {row}")
                    continue

        # Вычисляем минимум (если есть данные)
        try:
            min_result = find_min_thickness(self.s_min_lst)
        except ValueError:
            self.s_min_total.setPlainText("Нет данных")
            self.zav_s_min.setPlainText("Нет данных")
            print("Ошибка: нет числовых данных для вычисления минимума")
        else:
            self.s_min_total.setPlainText(format_ru(min_result.s_min))
            if min_result.s_min_index < len(self.text):
                self.zav_s_min.setPlainText(self.text[min_result.s_min_index])
            else:
                print(f"Ошибка: индекс {min_result.s_min_index} выходит за пределы списка {len(self.text)}")
                self.zav_s_min.setPlainText("Нет данных")
            self._completed_steps.add("s_min_min")

        # Собираем все года из третьего столбца.
        for row in range(table.rowCount()):
            item = table.item(row, 2)
            if item is not None and item.text():
                try:
                    years_min_max_lst.append(int(item.text()))
                except ValueError:
                    print(f"Пропуск нечислового значения в строке {row}")
                    continue

        # Вычисляем диапазон годов изготовления и добавляем в словарь data на вывод в ворд.
        self.data.update({"min_year": format_year_range(years_min_max_lst)})

    def calc_thick(self):
        """Функция - генератор толщин."""
        if not self._check_prerequisite("thickness"):
            return

        thick_table = self.table_thick
        amount = self.amount.value()
        tolshiny_lst = []

        # Устанавливаем высоту строки (вызовите это один раз при инициализации)
        thick_table.verticalHeader().setDefaultSectionSize(100)

        if len(self.text) != amount:
            # Обработка несоответствия количества элементов
            thick_table.setRowCount(0)
            return

        thick_table.setRowCount(amount)

        for row, zav_num in enumerate(self.text):
            # Устанавливаем заводской номер в первый столбец
            item = QTableWidgetItem(zav_num)
            thick_table.setItem(row, 0, item)

            table = self.table_ballons
            s_min_item = table.item(row, 1)
            g_i_bal_item = table.item(row, 2)
            massa_item = table.item(row, 3)

            # Получаем значения из ячеек
            s_min = s_min_item.text() if s_min_item else ""
            g_i_bal = g_i_bal_item.text() if g_i_bal_item else ""
            massa = massa_item.text() if massa_item else ""

            tolshiny_dict = {
                "zav": zav_num,
                "s_min": s_min,
                "g_i_bal": g_i_bal,
                "massa": massa
            }

            # Проверяем, что есть данные в self.s_min_lst
            if row < len(self.s_min_lst):
                try:
                    nums = float(self.s_min_lst[row])

                    # Генерируем 20 замеров толщины вокруг измеренного минимума
                    res_thick = generate_thickness_measurements(nums)

                    # Добавляем значения в словарь
                    for i, value in enumerate(res_thick, 1):
                        tolshiny_dict[f"s{i}"] = format_ru_fixed(value)

                    # Форматируем в 5 строк по 4 числа
                    res_thick_str = format_thickness_block(res_thick)

                    # Устанавливаем значения во второй столбец
                    item2 = QTableWidgetItem(res_thick_str)
                    thick_table.setItem(row, 1, item2)
                    tolshiny_lst.append(tolshiny_dict)

                except (ValueError, TypeError) as e:
                    print(f"Ошибка обработки данных для строки {row}: {e}")
                    continue

        self.data.update({"ballony": tolshiny_lst})
        self._completed_steps.add("thickness")

    def s_max_lst(self):
        """Собираем все макс толщины в список."""
        s_max_lst = []
        table = self.table_ballons

        for row in range(table.rowCount()):
            item = table.item(row, 2)
            if item is not None and item.text():
                try:
                    s_max_lst.append(float(item.text().replace(',', '.')))
                except ValueError:
                    print(f"Пропуск нечислового значения в строке {row}")
                    continue
        print(s_max_lst)

    def ovalnost_calc(self):
        """Расчёт овальности."""
        if not self._check_prerequisite("ovalness"):
            return

        bal_oval = []

        for zav in self.text:
            bal_oval_dict = {
                "z_n": zav
            }
            for i, m in enumerate(generate_ovalness_measurements(count=3)):
                bal_oval_dict.update({
                    f"d_max_rand{i}": f'{m.d_max}',
                    f"d_min_rand{i}": f'{m.d_min}',
                    f"oval{i}": f'{m.ovalness}'
                })

            bal_oval.append(bal_oval_dict)
        self.data.update({"bal_oval": bal_oval})
        self._completed_steps.add("ovalness")

    def tverdost(self) -> None:
        """Расчёт твёрдости и подготовка данных для Word."""
        if not self._check_prerequisite("hardness"):
            return
        try:
            # 1. Получаем предел прочности (Rm) из интерфейса — то же поле,
            # что уже вводится оператором для расчёта прочности в prochnost().
            rm = parse_ru(self.vrem_sopr_min.toPlainText())

            # 2. Расчёт минимальной и максимальной твёрдости по ГОСТ
            hb_range = calculate_hardness_range(rm)

            # 3. Генерация значений для каждого баллона (если нужно)
            tverdost_data = []

            for zav in self.text:
                tverdost_dict = {
                    "zav": zav
                }
                measurements = generate_hardness_measurements(hb_range.hb_min, hb_range.hb_max)
                for i, hb_random in enumerate(measurements, 1):
                    tverdost_dict.update({f"hb_{i}": f"{hb_random}"})
                tverdost_data.append(tverdost_dict)

            # 4. Формируем словарь для плейсхолдеров Word
            self.data.update({
                "hb_min": format_ru(hb_range.hb_min),
                "hb_max": format_ru(hb_range.hb_max),
                "tverdost_data": tverdost_data  # Список для цикла в Word
            })
            self._completed_steps.add("hardness")

        except ValueError as e:
            print(f"Ошибка ввода данных: {e}")
            self.show_message("Ошибка ввода", str(e), QMessageBox.Icon.Warning)

    def prochnost(self):
        if not self._check_prerequisite("strength"):
            return
        try:
            # Получаем данные из полей
            pred_tek_min = parse_ru(self.pred_tek_min.toPlainText())
            vrem_sopr_min = parse_ru(self.vrem_sopr_min.toPlainText())
            p_rab_MPa = parse_ru(self.p_rab_MPa.toPlainText())
            p_gidro = parse_ru(self.p_gidro.toPlainText())
            d_vnutr = parse_ru(self.d_vnutr.toPlainText())
            s_isp = parse_ru(self.s_isp.toPlainText())
            p_pnevma = parse_ru(self.p_pnevma.toPlainText())
            p_rab = parse_ru(self.p_rab.toPlainText())

            # Расчёт на прочность по ГОСТ 34233.1
            result = calculate_strength(
                pred_tek_min, vrem_sopr_min, p_rab_MPa, p_gidro,
                d_vnutr, s_isp, p_pnevma, p_rab,
            )

            # Давление для этапов пневматического испытания.
            self.data.update({"p_rab_025": f"{result.p_rab_025}"})
            self.data.update({"p_rab_05": f"{result.p_rab_05}"})
            self.data.update({"p_rab_075": f"{result.p_rab_075}"})

            # Вывод в QPlainTextEdit
            self.sigma.setPlainText(format_ru(result.sigma))
            self.sigma_gidro.setPlainText(format_ru(result.sigma_gidro))
            self.s_rasch.setPlainText(format_ru(result.s_rasch))
            self.s_rasch_gidro.setPlainText(format_ru(result.s_rasch_gidro))
            self.s_max_rasch.setPlainText(format_ru(result.s_max_rasch))
            self.p_pnevma_kgs.setPlainText(format_ru(result.p_pnevma_kgs))
            self.p_dop.setPlainText(format_ru(result.p_dop))
            self._completed_steps.add("strength")

        except ValueError as e:
            print(f"Ошибка ввода: {e}")
            self.sigma.setPlainText("Ошибка")
            self.sigma_gidro.setPlainText("Ошибка")
            self.s_max_rasch.setPlainText("Ошибка")

    def ost_res(self):
        if not self._check_prerequisite("residual_life"):
            return
        try:
            # Получаем данные из полей с проверкой на пустые значения
            s_isp = parse_ru(self.s_isp.toPlainText()) if self.s_isp.toPlainText() else 0.0
            c0_plus_dop = parse_ru(self.c0_plus_dop.toPlainText()) if self.c0_plus_dop.toPlainText() else 0.0
            s_min_total = parse_ru(self.s_min_total.toPlainText()) if self.s_min_total.toPlainText() else 0.0
            years_of_operation = parse_ru(self.yearsOfExpluatation.toPlainText()) \
                if self.yearsOfExpluatation.toPlainText() else 0.0

            # Получаем s_max_rasch (если это QPlainTextEdit)
            s_max_rasch = parse_ru(self.s_max_rasch.toPlainText()) if hasattr(self,
                                                                               's_max_rasch') and self.s_max_rasch.toPlainText() else 0.0

            # Расчёт скорости коррозии и остаточного ресурса
            # (raises ValueError, если срок эксплуатации равен нулю)
            result = calculate_residual_life(s_isp, c0_plus_dop, s_min_total, years_of_operation, s_max_rasch)

            # Вывод результатов
            self.a_corr.setPlainText(format_ru(result.corrosion_rate))
            self.tk_years.setPlainText(format_ru(result.remaining_years))
            self.tk_just.setPlainText(result.comment)
            self._completed_steps.add("residual_life")

        except ValueError as e:
            print(f"Ошибка ввода данных: {e}")
            self.a_corr.setPlainText("Ошибка")
            self.tk_years.setPlainText("Ошибка")

    # --- Методы для трубопровода (ГОСТ 32388-2013) ---

    def _install_segment_type_combo(self, table, row, current_text=SEGMENT_TYPES[0]):
        """Устанавливает выпадающий список типа элемента в ячейку (row, 1)
        таблицы участков трассы -- вместо свободного ввода текста."""
        combo = QComboBox()
        combo.addItems(SEGMENT_TYPES)
        if current_text not in SEGMENT_TYPES:
            combo.addItem(current_text)
        combo.setCurrentText(current_text)
        table.setCellWidget(row, 1, combo)

    def _install_ae_class_combo(self, table, row, current_text=AE_CLASS_TYPES[0]):
        """Устанавливает выпадающий список "Класс источника" в ячейку
        (row, 2) table_pnevmo_ae (Приложение 8, Таблица 1) -- тот же приём,
        что и _install_segment_type_combo(), фиксированный список
        AE_CLASS_TYPES."""
        combo = QComboBox()
        combo.addItems(AE_CLASS_TYPES)
        if current_text not in AE_CLASS_TYPES:
            combo.addItem(current_text)
        combo.setCurrentText(current_text)
        table.setCellWidget(row, 2, combo)

    def _install_uzk_segment_combo(self, table, row, current_text=None):
        """Устанавливает выпадающий список "Участок" (номер участка на
        схеме) в ячейку (row, 1) table_uzk (Приложение 5) -- вместо
        свободного ввода текста, только номера от 1 до segments_count
        (Приложение 4, "Количество участков") включительно, тот же приём,
        что и _install_segment_type_combo(). При смене выбора синхронизирует
        колонку "Типоразмер" той же строки (см. _on_uzk_segment_changed)."""
        options = [str(n) for n in range(1, self.segments_count.value() + 1)]
        combo = QComboBox()
        combo.addItems(options)
        if current_text and current_text not in options:
            combo.addItem(current_text)
        if current_text:
            combo.setCurrentText(current_text)
        combo.currentTextChanged.connect(lambda _=None, c=combo: self._on_uzk_segment_changed(c))
        table.setCellWidget(row, 1, combo)

    def _uzk_typorazmer_for_segment(self, segment_number_text):
        """Значение "Типоразмер" из table_segments (Приложение 4, "Участки
        трассы (ввод)") для строки с номером участка segment_number_text --
        поиск по значению колонки "№" (не по индексу строки, строки могли
        быть пересозданы через fill_segments_table())."""
        table = self.table_segments
        for row in range(table.rowCount()):
            number_item = table.item(row, 0)
            if number_item is not None and number_item.text().strip() == segment_number_text:
                size_item = table.item(row, 2)
                return size_item.text() if size_item else ""
        return ""

    def _sync_uzk_row_size(self, table, row):
        """Проставляет в (row, 2) table_uzk "Типоразмер" выбранного в (row,
        1) участка -- то же значение, что в одноимённой колонке
        table_segments для этого номера участка."""
        combo = table.cellWidget(row, 1)
        segment_number = combo.currentText() if combo else ""
        table.setItem(row, 2, QTableWidgetItem(self._uzk_typorazmer_for_segment(segment_number)))

    def _on_uzk_segment_changed(self, combo):
        """Реакция на смену номера участка в table_uzk -- находит строку
        комбобокса по факту (не по захваченному при подключении сигнала
        индексу: он "протухает" при удалении строк выше) и пересинхронизирует
        "Типоразмер" этой строки."""
        table = self.table_uzk
        for row in range(table.rowCount()):
            if table.cellWidget(row, 1) is combo:
                self._sync_uzk_row_size(table, row)
                return

    def _renumber_uzk_rows(self):
        """Колонка "№" table_uzk -- сквозная нумерация по позиции строки,
        та же логика, что и в fill_segments_table() (Приложение 4)."""
        table = self.table_uzk
        for row in range(table.rowCount()):
            table.setItem(row, 0, QTableWidgetItem(str(row + 1)))

    def _add_uzk_row(self):
        """Добавляет одну пустую строку в table_uzk (Приложение 5) с
        выпадающим списком номера участка в колонке 1, авто-номером в
        колонке 0 и подставленным "Типоразмер" в колонке 2."""
        table = self.table_uzk
        row = table.rowCount()
        table.insertRow(row)
        self._install_uzk_segment_combo(table, row)
        self._renumber_uzk_rows()
        self._sync_uzk_row_size(table, row)

    def _remove_uzk_row(self):
        """Удаляет выбранную строку table_uzk и перенумеровывает колонку
        "№" оставшихся строк (см. _renumber_uzk_rows)."""
        self._remove_table_row(self.table_uzk)
        self._renumber_uzk_rows()

    def _add_pnevmo_ae_row(self):
        """Добавляет одну пустую строку в table_pnevmo_ae (Приложение 8,
        Таблица 1) с выпадающим списком "Класс источника" в колонке 2 --
        для точечной правки поверх _fill_pnevmo_ae_table()."""
        table = self.table_pnevmo_ae
        row = table.rowCount()
        table.insertRow(row)
        self._install_ae_class_combo(table, row)

    def _fill_pnevmo_ae_table(self):
        """Заполняет table_pnevmo_ae (Приложение 8, Таблица 1) по числу
        датчиков (pnevmo_sensors_count, п.7): группами по 5 строк на
        каждый ПАЭ со стандартными ступенями нагрузки. Полностью
        перезаписывает таблицу -- как fill_segments_table(), кнопка не
        добавляет, а пересобирает. Дальше строки редактируются вручную
        (класс источника, нагрузка) или добавляются/удаляются по одной."""
        stages = ["0,3 Рразр", "0,6 Рразр", "Рразр", "Рисп.", "Рразр"]
        table = self.table_pnevmo_ae
        table.setRowCount(0)
        for paje_num in range(1, self.pnevmo_sensors_count.value() + 1):
            for stage in stages:
                row = table.rowCount()
                table.insertRow(row)
                table.setItem(row, 0, QTableWidgetItem(str(paje_num)))
                table.setItem(row, 1, QTableWidgetItem(stage))
                self._install_ae_class_combo(table, row)

    def _update_pnevmo_mirrors(self):
        """Зеркалит в read-only поля Приложения 8 (пункт 3) значения,
        введённые оператором ранее в других местах формы -- обязательное
        требование задачи "заполняются автоматически в read-only". Сами
        поля-зеркала (pnevmo_*_display) не входят в widget_names_pipeline.py
        и в шаблон .docx не попадают -- шаблон использует исходные ключи
        (obj_naznach, reg_number и т.п.), см. _update_calc_temp_display().
        pnevmo_obj_naznach/ae_zakl_obj_control сюда не входят -- они теперь
        самостоятельные редактируемые поля, см.
        _update_pnevmo_obj_naznach_display()/_update_ae_zakl_obj_control_display().
        ae_zakl_date_display и ae_zakl_report_date_display тоже сюда не
        входят -- обновляются сразу по dateChanged (pnevmo_date и
        report_date соответственно), см. _update_ae_zakl_date_display()/
        _update_ae_zakl_report_date_display() (раньше ждали переключения
        вкладки, из-за чего показывали старую дату)."""
        self.pnevmo_reg_number_display.setPlainText(self.reg_number.toPlainText())
        self.pnevmo_year_start_display.setPlainText(self.year_start.toPlainText())
        self.pnevmo_p_rab_display.setPlainText(self.p_rab_kgs.toPlainText())
        self.pnevmo_work_medium_display.setPlainText(self.work_medium.currentText())
        self.pnevmo_steel_grade_display.setPlainText(self._first_pipe_material_value(4))
        self.pnevmo_pipe_size_display.setPlainText(self._first_pipe_material_value(3))

        # Приложение 9 -- те же исходные значения, тот же приём зеркал.
        self.ae_zakl_reg_number_display.setPlainText(self.reg_number.toPlainText())
        self.ae_zakl_location_display.setPlainText(self.obj_location.toPlainText())

    def _update_ae_zakl_date_display(self):
        """Зеркалит pnevmo_date (Приложение 8, "Дата проведения") в
        read-only ae_zakl_date_display (Приложение 9, "Дата проведения
        контроля") сразу при изменении даты -- раньше обновлялось только
        при переключении вкладки (_update_pnevmo_mirrors()), из-за чего
        после правки даты в Приложении 8 тут держалась старая дата (по
        умолчанию 01.01.2000) до следующего переключения вкладки."""
        self.ae_zakl_date_display.setPlainText(self.pnevmo_date.date().toString("dd.MM.yyyy"))

    def _update_ae_zakl_report_date_display(self):
        """Зеркалит report_date (титульный лист, "Дата отчёта") в
        read-only ae_zakl_report_date_display (Приложение 9, "Дата
        отчёта (к «УТВЕРЖДАЮ»)") сразу при изменении даты -- тот же
        баг и тот же приём, что и в _update_ae_zakl_date_display(): без
        прямой подписки на dateChanged поле держало дату по умолчанию
        (01.01.2000) до переключения вкладки, из-за чего дата отчёта в
        Приложении 9 расходилась с титульным листом, хотя оба места
        рендерятся из одного и того же report_date в шаблоне .docx."""
        self.ae_zakl_report_date_display.setPlainText(self.report_date.date().toString("dd.MM.yyyy"))

    def _update_pnevmo_obj_naznach_display(self):
        """Подсказка pnevmo_obj_naznach (Приложение 8, п.3 "Трубопровод
        (наименование)") значением obj_naznach (раздел 2) -- поле
        редактируемое, см. _sync_mirror_field()."""
        self._sync_mirror_field(
            self.pnevmo_obj_naznach, self.obj_naznach.toPlainText(),
            "_pnevmo_obj_naznach_auto_value",
        )

    def _update_ae_zakl_obj_control_display(self):
        """Подсказка ae_zakl_obj_control (Приложение 9, "Объект контроля")
        значением pnevmo_obj_naznach (Приложение 8, "Трубопровод
        (наименование)") -- общий по смыслу плейсхолдер для обоих полей, см.
        _sync_mirror_field()."""
        self._sync_mirror_field(
            self.ae_zakl_obj_control, self.pnevmo_obj_naznach.toPlainText(),
            "_ae_zakl_obj_control_auto_value",
        )

    def _first_pipe_material_value(self, col):
        """Значение колонки col первой строки table_pipe_materials (Таблица
        6 -- Сведения о трубах) или "" если таблица пуста. Типоразмер и
        марка стали больше не отдельные поля формы -- вводятся один раз в
        Таблице 6, остальные места документа (участки трассы, расчёт на
        прочность) берут значение оттуда, из первого (представительного)
        элемента."""
        table = self.table_pipe_materials
        if table.rowCount() == 0:
            return ""
        item = table.item(0, col)
        if item is not None:
            return item.text().strip()
        cell_widget = table.cellWidget(0, col)
        # NB: "if cell_widget" была бы неверна -- пустой QComboBox (без
        # добавленных пунктов, только с введённым текстом) в Python ложный
        # (__len__() == 0), хотя currentText() при этом валиден.
        if isinstance(cell_widget, QComboBox):
            return cell_widget.currentText().strip()
        return ""

    def fill_segments_table(self):
        """Заполнение таблицы участков трассы. STEP_ORDER: 'segments'.

        Идемпотентно: повторное нажатие (например, чтобы заново отметить
        шаг выполненным после открытия сохранённого проекта -- см.
        FileHandler) не затирает уже введённые тип элемента и типоразмер
        существующих строк значениями по умолчанию, только достраивает
        новые строки при увеличении segments_count."""
        count = self.segments_count.value()
        default_size = self._first_pipe_material_value(3) or "-"
        table = self.table_segments
        existing_rows = table.rowCount()
        table.setRowCount(count)
        for row in range(count):
            table.setItem(row, 0, QTableWidgetItem(str(row + 1)))
            if row < existing_rows:
                type_combo = table.cellWidget(row, 1)
                current_type = type_combo.currentText() if type_combo else SEGMENT_TYPES[0]
                size_item = table.item(row, 2)
                current_size = size_item.text() if size_item and size_item.text() else default_size
            else:
                current_type = SEGMENT_TYPES[0]
                current_size = default_size
            self._install_segment_type_combo(table, row, current_type)
            table.setItem(row, 2, QTableWidgetItem(current_size))
        self._completed_steps.add("segments")
        self._update_toc_progress()

    def _read_segments(self):
        """Читает участки трассы из table_segments в список SegmentSpec."""
        table = self.table_segments
        segments = []
        for row in range(table.rowCount()):
            number_item = table.item(row, 0)
            size_item = table.item(row, 2)
            if number_item is None or not number_item.text():
                continue
            type_combo = table.cellWidget(row, 1)
            element_type = type_combo.currentText() if type_combo else SEGMENT_TYPES[0]
            segments.append(SegmentSpec(
                number=int(number_item.text()),
                element_type=element_type,
                size=size_item.text() if size_item and size_item.text() else "",
            ))
        return segments

    def calc_pipeline_thickness(self):
        """Синтетическая генерация замеров толщины по участкам. STEP_ORDER: 'thickness'."""
        if not self._check_prerequisite("thickness"):
            return
        try:
            s_min = parse_ru(self.thick_seed_min.toPlainText())
            segments = self._read_segments()
            if not segments:
                raise ValueError("Сначала заполните участки трассы")

            measurements = generate_pipeline_thickness_measurements(segments, s_min)

            table = self.table_thick_pipeline
            table.setRowCount(len(measurements))
            for row, m in enumerate(measurements):
                table.setItem(row, 0, QTableWidgetItem(str(m.number)))
                table.setItem(row, 1, QTableWidgetItem(m.element_type))
                table.setItem(row, 2, QTableWidgetItem(m.size))
                table.setItem(row, 3, QTableWidgetItem(format_ru_fixed(m.thickness, 2)))

            self.data.update({"segments": [
                {
                    "number": m.number, "element_type": m.element_type,
                    "size": m.size, "thickness": format_ru_fixed(m.thickness, 2),
                }
                for m in measurements
            ]})

            # Фактическая минимальная толщина -- минимум из сгенерированных
            # замеров, автоматически подставляется как вход для расчёта на
            # прочность (calc_sf), как Sф в контрольном примере из отчёта.
            s_fact_min = min(m.thickness for m in measurements)
            self.calc_sf.setPlainText(format_ru_fixed(s_fact_min, 2))

            self._completed_steps.add("thickness")
            self._update_toc_progress()

        except ValueError as e:
            print(f"Ошибка ввода данных: {e}")
            self.show_message("Ошибка ввода", str(e), QMessageBox.Icon.Warning)

    def calc_pipeline_strength_ui(self):
        """Расчёт на прочность по ГОСТ 32388-2013. STEP_ORDER: 'strength'."""
        if not self._check_prerequisite("strength"):
            return
        try:
            p_working = parse_ru(self.p_rab_mpa.toPlainText())
            d_outer = parse_ru(self.calc_da.toPlainText())
            temp = parse_ru(self.calc_temp.toPlainText())
            phi = parse_ru(self.calc_phi.toPlainText())
            c2 = parse_ru(self.calc_c2.toPlainText())
            s_actual = parse_ru(self.calc_sf.toPlainText())
            steel_grade = self._first_pipe_material_value(4)
            if not steel_grade:
                raise ValueError(
                    "Сначала заполните марку стали в Таблице 6 (Сведения о трубах)"
                )
            self.calc_steel_grade.setPlainText(steel_grade)

            try:
                allowable_stress = get_allowable_stress(steel_grade, temp)
                self.calc_sigma_allow.setPlainText(format_ru(allowable_stress))
            except (KeyError, ValueError) as e:
                manual_value = self.calc_sigma_allow.toPlainText().strip()
                if not manual_value:
                    raise ValueError(
                        f"{e} Введите [σ] вручную в поле «Допускаемое "
                        "напряжение [σ], МПа»."
                    )
                allowable_stress = parse_ru(manual_value)

            result = calculate_pipeline_strength(
                p_working=p_working, d_outer=d_outer, allowable_stress=allowable_stress,
                s_actual=s_actual, c2=c2, phi=phi,
            )

            self.calc_sr.setPlainText(format_ru(result.s_calc))
            self.calc_s_reject.setPlainText(format_ru(result.s_reject))
            self.calc_p_allow.setPlainText(format_ru(result.p_allow))
            self.calc_strength_conclusion.setPlainText(
                "Условие прочности выполняется" if result.strength_ok
                else "Условие прочности не выполняется"
            )
            self._completed_steps.add("strength")
            self._update_toc_progress()

        except (ValueError, KeyError) as e:
            print(f"Ошибка ввода данных: {e}")
            self.show_message("Ошибка ввода", str(e), QMessageBox.Icon.Warning)

    def calc_pipeline_residual_life_ui(self):
        """Остаточный ресурс по скорости коррозии. STEP_ORDER: 'residual_life'."""
        if not self._check_prerequisite("residual_life"):
            return
        try:
            s_nominal = parse_ru(self.calc_sn.toPlainText())
            s_actual = parse_ru(self.calc_sf.toPlainText())
            s_reject = parse_ru(self.calc_s_reject.toPlainText())
            years = parse_ru(self.calc_years_operation.toPlainText())
            k = parse_ru(self.calc_k.toPlainText()) if self.calc_k.toPlainText().strip() else 1.0

            # calc_corrosion_rate редактируется вручную: если поле уже
            # заполнено (расчётом или инженером), его текущее значение
            # берётся как есть, а не перезаписывается формулой -- иначе
            # правка терялась бы при каждом повторном нажатии кнопки.
            corrosion_text = self.calc_corrosion_rate.toPlainText().strip()
            corrosion_override = parse_ru(corrosion_text) if corrosion_text else None

            result = calculate_pipeline_residual_life(
                s_nominal=s_nominal, s_actual=s_actual, s_reject=s_reject,
                years_of_operation=years, k=k, corrosion_rate_override=corrosion_override,
            )

            self.calc_corrosion_rate.setPlainText(format_ru(result.corrosion_rate))
            self.calc_remaining_years.setPlainText(format_ru(result.remaining_years))
            # calc_residual_comment -- вывод формулируется инженером текстом
            # вручную (как thick_conclusion/uzk_conclusion), кнопка его не
            # перезаписывает -- иначе правки в UI терялись бы при каждом
            # повторном расчёте.

            # final_years_allowed (раздел 8) -- подсказка расчётным остаточным
            # ресурсом, только если поле ещё пустое: инженер вправе утвердить
            # меньший регламентный срок, повторный расчёт его правку не сотрёт.
            # setPlainText ниже сам вызовет _update_final_deadline_date() через
            # textChanged (см. подключение сигнала в __init__).
            if not self.final_years_allowed.toPlainText().strip():
                self.final_years_allowed.setPlainText(str(int(result.remaining_years)))

            self._completed_steps.add("residual_life")
            self._update_toc_progress()

        except ValueError as e:
            print(f"Ошибка ввода данных: {e}")
            self.show_message("Ошибка ввода", str(e), QMessageBox.Icon.Warning)

    def _add_table_row(self, table):
        """Добавляет пустую строку в конец таблицы -- рассмотренные
        документы, элементы ВИК. Свободный ввод, не расчётный шаг -- не
        входит в STEP_ORDER. Специалисты используют отдельный
        _add_specialist_row() -- их ячейки не свободный текст, а
        редактируемые выпадающие списки."""
        table.insertRow(table.rowCount())

    def _remove_table_row(self, table):
        """Удаляет выбранную строку таблицы (если строка выбрана). Работает
        одинаково для ячеек-текста и ячеек-виджетов (комбобоксов)."""
        row = table.currentRow()
        if row >= 0:
            table.removeRow(row)

    def _remove_combo_current_item(self, combo):
        """Удаляет из выпадающего списка текущий пункт (толщиномер,
        дефектоскоп, аппаратура АЭ и т.п. -- редактируемые комбобоксы, куда
        новые варианты добавляются вводом текста)."""
        index = combo.currentIndex()
        if index >= 0:
            combo.removeItem(index)
        else:
            combo.clearEditText()

    def _install_growable_combo(self, table, row, col, current_text=""):
        """Устанавливает редактируемый выпадающий список в ячейку (row, col)
        произвольной таблицы -- вместо свободного ввода текста. В отличие от
        _install_segment_type_combo() (фиксированный SEGMENT_TYPES), список
        здесь растёт вводом оператора (как thick_device и т.п.) и изначально
        собирается из уже введённых значений этой же колонки в других
        строках -- чтобы повторно использовать ранее введённые значения
        (ФИО/должности/удостоверения в специалистах, марки стали/ГОСТы в
        трубах и т.п.) через выпадающий список."""
        combo = QComboBox()
        combo.setEditable(True)
        seen = []
        for r in range(table.rowCount()):
            if r == row:
                continue
            existing = table.cellWidget(r, col)
            if isinstance(existing, QComboBox):
                text = existing.currentText()
                if text and text not in seen:
                    seen.append(text)
        combo.addItems(seen)
        if current_text and current_text not in seen:
            combo.addItem(current_text)
        combo.setCurrentText(current_text)
        table.setCellWidget(row, col, combo)

    def _add_specialist_row(self):
        """Добавляет строку в table_specialists -- специалиста выбирают из
        справочника «Сотрудники» (EmployeesTabController), а не вводят
        текстом: так подтвердил пользователь, весь состав специалистов
        отчёта должен идти через справочник (иначе для строки не с кем
        связать клише -- см. _specialist_kleishe_image()). Если нужного
        человека нет в справочнике -- его сначала заводят на вкладке
        «Сотрудники».

        4 ячейки по-прежнему QComboBox (редактируемый, как и раньше --
        _cell_text()/_table_to_dicts() рассчитаны именно на этот тип), но
        предзаполненные из Employee, а не пустые."""
        employees = load_employees()
        if not employees:
            self.show_message(
                "Справочник пуст",
                "Сначала добавьте сотрудников на вкладке «Сотрудники» -- "
                "специалисты отчёта выбираются оттуда.",
                QMessageBox.Icon.Warning,
            )
            return

        labels = [f"{e.position} — {e.full_name}" for e in employees]
        label, ok = QInputDialog.getItem(
            self, "Выбор специалиста", "Сотрудник:", labels, editable=False
        )
        if not ok:
            return
        employee = employees[labels.index(label)]

        cert_full = ""
        cert_short = ""
        if employee.certificates:
            cert_full = employee.certificates[0]
            if len(employee.certificates) > 1:
                cert_full, ok = QInputDialog.getItem(
                    self, "Выбор удостоверения", "Удостоверение:",
                    employee.certificates, editable=False,
                )
                if not ok:
                    cert_full = employee.certificates[0]
            match = re.search(r"№\s*\S+", cert_full)
            cert_short = match.group(0) if match else cert_full

        table = self.table_specialists
        row = table.rowCount()
        table.insertRow(row)
        for col, text in enumerate((employee.position, employee.full_name, cert_full, cert_short)):
            self._install_growable_combo(table, row, col, text)
        self._specialist_employee_ids.append(employee.id)
        self._refresh_program_specialist_combo()

    def _remove_specialist_row(self):
        """Удаляет выбранную строку table_specialists, синхронно убирает её
        employee_id из _specialist_employee_ids (см. _add_specialist_row()) и
        обновляет комбобоксы выбора специалиста (program_specialist и
        т.п.) -- раньше эти комбобоксы обновлялись только при переключении
        вкладки, теперь вкладок нет, обновление дергается прямо из точек,
        где меняется table_specialists, см. _refresh_program_specialist_combo()."""
        row = self.table_specialists.currentRow()
        if row >= 0 and row < len(self._specialist_employee_ids):
            self._specialist_employee_ids.pop(row)
        self._remove_table_row(self.table_specialists)
        self._refresh_program_specialist_combo()

    def _specialist_kleishe_image(self, doc, employees, employee_id, used_paths):
        """InlineImage клише сотрудника для подстановки в form_data (см.
        calculate(), KLEISHE_ROLE_PREFIXES) -- тот же приём, что и
        nk_scheme_image/pnevmo_graph_image: пустая строка, если клише нет
        (сотрудник не выбран, не привязан к справочнику или не загрузил
        клише), без ошибки рендера. Путь к файлу резолвит
        resolve_kleishe_path() (src/services/employees_store.py).

        used_paths -- set, куда добавляется путь картинки, если она
        подставлена -- после doc.render() по нему находим rId вставленных
        картинок для _float_kleishe_drawings_behind_text() (картинка клише
        должна плавать "за текстом", а не как схема НК/график нагружения,
        см. calculate()).

        width/height не заданы намеренно -- InlineImage тогда берёт
        реальный размер картинки (пиксели + DPI из самого файла, см.
        docx.image.image.Image.width/height), а не произвольно
        подогнанный -- раньше здесь стоял фиксированный width=Mm(30),
        из-за чего клише в документе не совпадало по размеру с исходным
        файлом."""
        path = resolve_kleishe_path(employees, employee_id)
        if not path:
            return ""
        used_paths.add(path)
        return InlineImage(doc, str(path))

    def _float_kleishe_drawings_behind_text(self, doc, used_paths):
        """После doc.render() переводит уже вставленные картинки клише из
        обычного инлайн-положения в плавающее "за текстом" (Word:
        Обтекание текстом -> За текстом) -- без изменения размера. Только
        клише -- схема НК/график нагружения (тоже InlineImage) не
        затрагиваются, т.к. их rId в used_paths не попадают.

        rId картинок определяем через get_or_add_image() -- она дедуплицирует
        по содержимому так же, как это уже сделал docxtpl при рендере (см.
        docx.parts.story.StoryPart.get_or_add_image()), поэтому для уже
        вставленной картинки метод просто возвращает тот же rId, без
        побочных эффектов."""
        if not used_paths:
            return
        target_rids = {doc.part.get_or_add_image(str(path))[0] for path in used_paths}
        float_drawings_behind_text(doc.docx, target_rids)

    def _add_pipe_material_row(self):
        """Добавляет строку в table_pipe_materials (Таблица 6 -- Сведения о
        трубах): № проставляется автоматически, колонки "Наименование
        элемента", "Марка стали, ГОСТ или ТУ" и "Трубы, ГОСТ или ТУ" --
        растущие выпадающие списки (см. _install_growable_combo()),
        "Количество" и "Типоразмер" остаются свободным текстом."""
        table = self.table_pipe_materials
        row = table.rowCount()
        table.insertRow(row)
        table.setItem(row, 0, QTableWidgetItem(f"{row + 1}."))
        self._install_growable_combo(table, row, 1)
        self._install_growable_combo(table, row, 4)
        self._install_growable_combo(table, row, 5)

    def _seed_program_table_defaults(self):
        """Предзаполняет table_program (Приложение 1 -- Программа) стандартным
        составом работ из PROGRAM_DEFAULT_ITEMS -- только если таблица ещё
        пустая (не перетирает восстановленный из проекта или уже
        отредактированный оператором список). Дальше строки полностью
        редактируются/удаляются/добавляются через UI, как обычные строки."""
        if self.table_program.rowCount() > 0:
            return
        table = self.table_program
        for level, text in PROGRAM_DEFAULT_ITEMS:
            row = table.rowCount()
            table.insertRow(row)
            number_item = QTableWidgetItem("")
            number_item.setData(Qt.ItemDataRole.UserRole, level)
            table.setItem(row, 0, number_item)
            table.setItem(row, 1, QTableWidgetItem(text))
        self._renumber_program_table()

    def _add_program_item_row(self):
        """Добавляет пункт верхнего уровня в table_program -- № вида "N.",
        см. _renumber_program_table()."""
        self._insert_program_row(level=0)

    def _add_program_subitem_row(self):
        """Добавляет подпункт в table_program -- № вида "N.M." под текущим
        (последним) пунктом верхнего уровня, см. _renumber_program_table()."""
        self._insert_program_row(level=1)

    def _insert_program_row(self, level):
        """Общая часть _add_program_item_row()/_add_program_subitem_row():
        строка всегда добавляется в конец таблицы (как и все остальные
        таблицы в проекте -- ни у одной сейчас нет reorder/insert-в-середину),
        уровень хранится в Qt.ItemDataRole.UserRole на item(row, 0)."""
        table = self.table_program
        row = table.rowCount()
        table.insertRow(row)
        number_item = QTableWidgetItem("")
        number_item.setData(Qt.ItemDataRole.UserRole, level)
        table.setItem(row, 0, number_item)
        table.setItem(row, 1, QTableWidgetItem(""))
        self._renumber_program_table()

    def _remove_program_row(self):
        """Удаляет выбранную строку table_program и пересчитывает номера --
        обычный _remove_table_row() номера не трогает."""
        self._remove_table_row(self.table_program)
        self._renumber_program_table()

    def _renumber_program_table(self):
        """Пересчитывает колонку "№ п/п" в table_program по уровням: level 0
        -> "N.", level 1 -> "N.M." под текущим top-level пунктом. Уровень
        читается из Qt.ItemDataRole.UserRole на item(row, 0); если он не
        выставлен (например, после восстановления проекта из JSON -- Project
        хранит только текст ячеек, см. src/ui/file_handler.py), определяется
        эвристикой по уже отображённому номеру ("1.1." -> подпункт, иначе --
        пункт верхнего уровня), чтобы повторное открытие сохранённого
        проекта не расплющивало уже сохранённую иерархию при следующем
        добавлении строки. Подпункт раньше первого пункта верхнего уровня
        трактуется как пункт верхнего уровня (без "0.1.")."""
        table = self.table_program
        top = 0
        sub = 0
        for row in range(table.rowCount()):
            item = table.item(row, 0)
            if item is None:
                item = QTableWidgetItem("")
                table.setItem(row, 0, item)
            level = item.data(Qt.ItemDataRole.UserRole)
            if level is None:
                level = 1 if item.text().strip(".").count(".") >= 1 else 0
            if level != 0 and top == 0:
                level = 0
            item.setData(Qt.ItemDataRole.UserRole, level)
            if level == 0:
                top += 1
                sub = 0
                item.setText(f"{top}.")
            else:
                sub += 1
                item.setText(f"{top}.{sub}.")

    def _refresh_program_specialist_combo(self, saved_indices=None):
        """Обновляет списки в program_specialist (поле "Программу составил",
        Приложение 1), act2_specialist (поле "Анализ документации провёл",
        Приложение 2), vik_specialist (поле "Контроль провёл", Приложение 3),
        thick_specialist (поле "Измерение провёл", Приложение 4),
        uzk_specialist (поле "Измерение провёл", Приложение 5),
        calc_specialist (поле "Расчёт выполнил", Приложение 6),
        pnevmo_specialist (поле "Контроль выполнил", Приложение 8) и
        ae_zakl_specialist (поле "Заключение составил", Приложение 9).
        Источник вариантов для всех один и тот же: table_specialists (1.3
        Сведения о специалистах, Таблица 2) -- вызывается прямо из точек,
        где меняется эта таблица (_add_specialist_row/_remove_specialist_row),
        плюс после загрузки проекта (см. file_handler.py) -- раньше
        обновление держалось на переключении вкладки "Приложения"/"Расчёты",
        вкладок больше нет. Ни один из комбобоксов не входит ни в один
        список widget_names_pipeline.py (как pnevmo_pressure_hint) -- каждый
        даёт индекс строки специалиста, а не текст для .docx напрямую,
        итоговые плейсхолдеры собирает calculate(). Заодно обновляет
        read-only зеркала пункта 3 Приложения 8, см. _update_pnevmo_mirrors().

        saved_indices -- необязательный dict {имя_комбобокса: индекс},
        восстановленный из report_data при загрузке проекта (см.
        FileHandler.open_project_json()/MainWindow._open_document()).
        Перекрывает обычную логику "сохранить текущий выбор" в
        _refresh_specialist_combo() -- для только что открытого документа
        combo.currentData() всегда None (комбобоксы ещё не заполнены),
        без этого выбор специалиста откатывался на первую строку."""
        saved_indices = saved_indices or {}
        for name in self.SPECIALIST_COMBO_NAMES:
            self._refresh_specialist_combo(getattr(self, name), saved_indices.get(name))
        self._update_pnevmo_mirrors()

    def _refresh_specialist_combo(self, combo, desired=None):
        """Перезаполняет один комбобокс-выбор специалиста вариантами из
        table_specialists, сохраняя выбор (по индексу строки), если он
        всё ещё существует -- общая логика для program_specialist и
        act2_specialist, см. _refresh_program_specialist_combo().

        desired -- индекс строки, который нужно выставить явно (при
        загрузке проекта, когда живого текущего выбора в ещё не
        заполненном комбобоксе нет); по умолчанию берётся
        combo.currentData() -- обычный случай live-редактирования
        таблицы специалистов, когда выбор уже стоит в комбобоксе."""
        previous = desired if desired is not None else combo.currentData()
        combo.blockSignals(True)
        combo.clear()
        for row in range(self.table_specialists.rowCount()):
            position = self._cell_text(self.table_specialists, row, 0)
            name = self._cell_text(self.table_specialists, row, 1)
            label = " — ".join(part for part in (position, name) if part) or f"Специалист {row + 1}"
            combo.addItem(label, row)
        if previous is not None:
            index = combo.findData(previous)
            if index >= 0:
                combo.setCurrentIndex(index)
        combo.blockSignals(False)

    def _switch_view(self, view):
        """Переключает viewStack между документом и общими справочниками
        (Сотрудники/Приборы/Документы).

        Имя страницы резолвится через getattr() лениво -- ТОЛЬКО для
        запрошенного view, а не собирает сразу все четыре в dict-литерал:
        у конструктора документов (equipment_type.id == "constructor")
        свой mainSplitter/viewStack (см. constructor_window.ui), но нет
        вкладок tab_employees/tab_instruments/tab_orgdocs -- сборка
        словаря со всеми четырьмя сразу упала бы AttributeError, даже
        если реально нужна только "document"."""
        page_names = {
            "document": "tab_document",
            "employees": "tab_employees",
            "instruments": "tab_instruments",
            "orgdocs": "tab_orgdocs",
        }
        self.viewStack.setCurrentWidget(getattr(self, page_names[view]))
        self._current_view = view
        self._update_report_buttons_visibility()

    _ACTIVITY_VIEW_TITLES = {"search": "Поиск", "employees": "Сотрудники"}

    def _switch_activity_view(self, view):
        """Переключатель активити-бара конструктора документов (docs/design/
        вводная_часть.html, switchView()) -- аналог _switch_view() выше, но
        для конструктора: переключает ЦЕЛИКОМ страницу viewStack, а не
        только основную область, поскольку у "editor" и "database" разный
        сайдбар (палитра блоков vs. дерево объектов), а не общий, как у
        трубопровода. "search"/"employees" -- заглушка (page_stub), как и
        в самом мокапе (activityViewTitles там же) -- раздел не
        реализован, страница просто показывает название.

        Повторный клик по уже активной иконке «Редактор документов» вместо
        обычного переключения сворачивает/разворачивает constructorSidebar
        (см. _toggle_constructor_sidebar()) -- тот же паттерн, что в
        VSCode (клик по активной иконке активити-бара прячет её панель), и
        в мокапе (toggleEditorSidebar()). Только для "editor" -- у
        "database"/заглушек своего сайдбара для сворачивания либо нет
        (page_stub), либо он уже сворачивается отдельной кнопкой
        (sidebarBtn_collapse/_toggle_sidebar(), другой, не связанный
        виджет -- см. обсуждение задачи)."""
        if view == "editor" and getattr(self, "_current_view", None) == "editor":
            self._toggle_constructor_sidebar()
            return
        page_names = {"editor": "tab_document", "database": "page_database"}
        page = getattr(self, page_names.get(view, "page_stub"))
        self.viewStack.setCurrentWidget(page)
        for name in ("search", "database", "editor", "employees"):
            getattr(self, f"activityBtn_{name}").setChecked(name == view)
        if page is self.page_stub:
            self.activityStubTitle.setText(self._ACTIVITY_VIEW_TITLES.get(view, ""))
        elif view == "database":
            self._refresh_objects_tree()
        self._current_view = view

    def _toggle_constructor_sidebar(self):
        """Сворачивает/разворачивает constructorSidebar (палитру разделов) --
        сам виджет просто прячется (setVisible()), соседняя канва
        (mainArea, тот же QHBoxLayout) сама растягивается на освободившееся
        место. Отдельный от _toggle_sidebar()/sidebarBtn_collapse -- тот
        относится к НЕСВЯЗАННОМУ виджету sidebar (дерево «Объекты», общее с
        трубопроводом, см. обсуждение задачи), constructorSidebar с ним не
        путать."""
        self.constructorSidebar.setVisible(not self.constructorSidebar.isVisible())

    def _reset_form(self):
        """Очищает форму под новый/другой документ -- обратная операция к
        get_form_data()/_fill_ui_from_project(), проходит по тем же
        спискам виджетов (widget_names_pipeline.py и т.п.), только очищая
        вместо чтения. Раньше такой возможности не было вообще -- "начать
        заново" означало перезапустить приложение; нужна для переключения
        между документами дерева объектов без перезапуска.

        Комбобоксы намеренно НЕ .clear() -- это стёрло бы предзаполненные
        варианты (у work_medium это единственный источник выбора: азот/
        воздух/кислород/гелий/аргон, не растится вводом). setCurrentIndex(0)
        воспроизводит тот же вид, что при самом первом запуске окна (ни у
        одного из комбобоксов currentIndex в .ui не выставлен явно, Qt по
        умолчанию показывает первый пункт).

        _current_document_path обнуляется ДО очистки виджетов, а не после
        -- сама очистка (setPlainText(""), setCurrentIndex(0) и т.п.)
        дёргает те же сигналы, что и правки оператора, и без этого
        _mark_dirty() успел бы ложно пометить грязным только что
        сохранённый документ, который мы покидаем (см. Фаза 5.4)."""
        self._current_document_path = None
        for name in self.PLAIN_TEXT_EDIT_NAMES:
            getattr(self, name).setPlainText("")
        for name in self.COMBO_BOX_NAMES:
            getattr(self, name).setCurrentIndex(0)
        for name in self.DATE_EDIT_NAMES:
            getattr(self, name).setDate(QDate.currentDate())
        for name in self.SPIN_BOX_NAMES:
            widget = getattr(self, name)
            widget.setValue(widget.minimum())
        for name in self.TABLE_WIDGET:
            getattr(self, name).setRowCount(0)

        self.data = {}
        self._completed_steps = set()

        if self.equipment_type.id == "pipeline":
            self._seed_program_table_defaults()
            self._specialist_employee_ids = []
            self._current_toc_items = []
            self._current_toc_active_item = None
            self._update_breadcrumb()

    def _refresh_objects_tree(self):
        """Перестраивает objectsTree с нуля из файловой системы (см.
        src/services/workspace.py) -- папки произвольной вложенности
        (docs/design/вводная_часть.html: "Папки -- произвольная
        вложенность"), документы -- листья внутри них, см.
        _build_objects_tree_level(). Вызывается при старте и сразу после
        создания/переименования/удаления папки или документа --
        построение дешёвое (десятки папок/файлов, не тысячи), отдельный
        кэш не заводится.

        Каждая строка несёт словарь в Qt.ItemDataRole.UserRole --
        {"kind": "folder", "path": ...} или {"kind": "document", "path":
        ...} (TOC-строки ниже используют тот же приём с "groupbox"
        вместо "path", см. _populate_document_toc()). Тип строки
        определяется ЭТИМ полем, а не глубиной вложенности в дереве --
        раньше (до произвольной вложенности папок) глубина однозначно
        отличала объект/документ/раздел TOC, сейчас документ и TOC-раздел
        могут оказаться на любом уровне.

        Иконки -- folder (золотистая) на папке, file-text (синяя) на
        документе: визуально сразу отличимо, что папка, а что файл
        (тот же принцип для обоих окон, использующих это дерево --
        трубопровод и конструктор документов)."""
        self.objectsTree.clear()
        self._build_objects_tree_level(None, workspace.OUTPUT_DIR)

        # objectsTree.clear() выше уничтожает и дочерние строки
        # оглавления под строкой текущего документа (см.
        # _populate_document_toc()), а _current_toc_items/
        # _current_toc_active_item эти QTreeWidgetItem не сбрасывает --
        # без восстановления это висячие ссылки на удалённые C++
        # объекты (падение "wrapped C/C++ object of type QTreeWidgetItem
        # has been deleted" при следующем _update_breadcrumb()/
        # _on_document_scrolled()). Отстраиваем TOC текущего документа
        # заново, если он есть в новом дереве -- только у трубопровода
        # (см. equipment_type.id == "pipeline"): у конструктора
        # документов TOC-виджетов (_toc_groupboxes/tab_document_scrollContent)
        # нет вовсе, вызов _populate_document_toc() упал бы AttributeError.
        if self.equipment_type.id == "pipeline" and self._current_document_path is not None:
            doc_item = self._find_document_tree_item(self._current_document_path)
            if doc_item is not None:
                self._populate_document_toc(doc_item)

    def _build_objects_tree_level(self, parent_item, folder_dir):
        """Рекурсивно наполняет objectsTree содержимым folder_dir --
        сперва подпапки (алфавит), затем документы (алфавит), на любую
        глубину. parent_item=None -- верхний уровень дерева
        (workspace.OUTPUT_DIR). workspace.list_objects()/list_documents()
        уже были path-based и не завязаны на глубину сами по себе --
        ограничение было только в построении дерева, не в файловой
        модели (см. src/services/workspace.py)."""
        for name in workspace.list_objects(folder_dir):
            child_dir = folder_dir / name
            folder_item = QTreeWidgetItem([name])
            folder_item.setIcon(0, icons.icon("folder", "#c9a227", 14))
            folder_item.setData(0, Qt.ItemDataRole.UserRole, {"kind": "folder", "path": child_dir})
            if parent_item is None:
                self.objectsTree.addTopLevelItem(folder_item)
            else:
                parent_item.addChild(folder_item)
            self._build_objects_tree_level(folder_item, child_dir)
            folder_item.setExpanded(True)
        for path, label in workspace.list_documents(folder_dir, self.equipment_type.id):
            doc_item = QTreeWidgetItem([label])
            doc_item.setIcon(0, icons.icon("file-text", "#0a84ff", 14))
            doc_item.setData(0, Qt.ItemDataRole.UserRole, {"kind": "document", "path": path})
            if parent_item is None:
                self.objectsTree.addTopLevelItem(doc_item)
            else:
                parent_item.addChild(doc_item)

    def _filter_objects_tree(self, text):
        """Фильтр по вводу в sidebarSearchBox -- рекурсивный (произвольная
        вложенность папок): строка видна, если её текст совпал с
        введённым, ИЛИ у неё есть видимый (по тому же правилу) потомок --
        иначе поиск по имени папки прятал бы всё содержимое внутри, а
        совпавший документ в глубокой подпапке остался бы скрыт вместе со
        свёрнутыми предками. TOC-строки (см. _populate_document_toc()) в
        поиске не участвуют -- как и раньше, когда фильтр вообще не
        спускался до их уровня."""
        text = text.strip().lower()

        def apply_filter(item):
            self_match = text in item.text(0).lower()
            any_child_match = False
            for i in range(item.childCount()):
                child = item.child(i)
                data = child.data(0, Qt.ItemDataRole.UserRole)
                if isinstance(data, dict) and data.get("kind") == "toc":
                    continue
                if apply_filter(child):
                    any_child_match = True
            visible = not text or self_match or any_child_match
            item.setHidden(not visible)
            return self_match or any_child_match

        for i in range(self.objectsTree.topLevelItemCount()):
            apply_filter(self.objectsTree.topLevelItem(i))

    def _toggle_sidebar(self):
        """Сворачивает/разворачивает сайдбар (Фаза 7, как в референсе) --
        весь sidebar прячется целиком, вместо него показывается узкая
        полоса revealStrip (22px) с одной кнопкой разворота. Оба виджета
        -- постоянные дети mainSplitter (Фаза 4.3), переключается только
        видимость -- QSplitter сам схлопывает скрытого ребёнка до 0 при
        пересчёте раскладки, отдельно двигать min/maxWidth не нужно.
        childrenCollapsible=false защищает только от случайного
        схлопывания перетаскиванием мышью, программному setVisible() не
        мешает."""
        self._sidebar_collapsed = not self._sidebar_collapsed
        self.sidebar.setVisible(not self._sidebar_collapsed)
        self.revealStrip.setVisible(self._sidebar_collapsed)
        if self._sidebar_collapsed:
            self._sidebar_expanded_sizes = self.mainSplitter.sizes()
        else:
            self.mainSplitter.setSizes(self._sidebar_expanded_sizes)
            # У конструктора documents-дерево живёт на отдельной странице
            # viewStack (page_database, см. _switch_activity_view()) --
            # разворачивание/сворачивание sidebar внутри неё не должно
            # менять текущую страницу вообще, поэтому здесь только
            # трубопровод (у него единственная страница "document" делит
            # mainSplitter с этим же sidebar) пересобирает TOC-биндинги.
            if self.equipment_type.id == "pipeline":
                self._switch_view(self._current_view)

    def _on_objects_tree_item_clicked(self, item, column):
        """Клик по строке objectsTree -- три вида строк, различаются
        полем "kind" в Qt.ItemDataRole.UserRole (см.
        _build_objects_tree_level()/_populate_document_toc()), а НЕ
        глубиной вложенности -- с произвольной вложенностью папок
        документ и раздел TOC могут оказаться на любом уровне. Папка
        (kind="folder") -- ничего не делаем, QTreeWidget сам
        разворачивает/сворачивает. Раздел оглавления (kind="toc", только
        дочерние строки документа, только у трубопровода) -- скроллим к
        разделу, см. _scroll_to_toc_item().

        Документ (kind="document") -- у трубопровода клик сразу открывает
        его (см. _open_document()), как и раньше. У конструктора
        документов -- «База документов» теперь отдельная страница со
        своим превью (см. _show_document_preview()), клик по документу
        только показывает его содержимое справа, не переключая на
        «Редактор документов» -- явный переход туда теперь через ПКМ ->
        «Открыть в редакторе документов» (см. _show_document_context_menu()),
        чтобы не терять/не подменять несохранённые правки открытого
        сейчас документа одним неосторожным кликом по дереву."""
        data = item.data(0, Qt.ItemDataRole.UserRole)
        if not isinstance(data, dict):
            return
        if data.get("kind") == "document":
            if self.equipment_type.id == "constructor":
                self._show_document_preview(data["path"])
            else:
                self._open_document(data["path"])
        elif data.get("kind") == "toc":
            self._scroll_to_toc_item(item)

    def _show_objects_tree_context_menu(self, pos):
        """ПКМ по objectsTree -- вид меню зависит от "kind" строки под
        курсором (см. _on_objects_tree_item_clicked()): папка ->
        _show_folder_context_menu(), документ ->
        _show_document_context_menu(), раздел TOC -- меню нет, там
        нечем управлять."""
        item = self.objectsTree.itemAt(pos)
        if item is None:
            return
        data = item.data(0, Qt.ItemDataRole.UserRole)
        if not isinstance(data, dict):
            return
        if data.get("kind") == "folder":
            self._show_folder_context_menu(item, pos)
        elif data.get("kind") == "document":
            self._show_document_context_menu(item, pos)

    def _show_document_context_menu(self, item, pos):
        """ПКМ по строке документа. «Переименовать» из референсного
        мокапа (docs/design/pipeline_sidebar_mockup.html) отдельным
        пунктом меню не реализовано -- ярлык документа это имя файла
        (path.stem, см. workspace.list_documents()), а переименование
        уже доступно через «Сохранить проект»: диалог всегда просит имя
        файла (FileHandler._prompt_document_path()), и при вводе
        другого имени старый файл удаляется, а не остаётся сиротой
        (см. FileHandler.save_project_json()).

        «Открыть в редакторе документов» -- только для конструктора
        (первым пунктом): у него клик по строке в «Базе документов»
        теперь только показывает превью справа (см.
        _on_objects_tree_item_clicked()/_show_document_preview()),
        поэтому явный переход к редактированию нужен отдельным
        действием. У трубопровода клик уже открывает документ сразу --
        дублировать это тем же пунктом меню не нужно."""
        path = item.data(0, Qt.ItemDataRole.UserRole)["path"]
        menu = QMenu(self)
        if self.equipment_type.id == "constructor":
            menu.addAction("Открыть в редакторе документов", lambda: self._open_document(path))
            menu.addSeparator()
        menu.addAction("Дублировать", lambda: self._duplicate_document(path))
        menu.addAction("Показать в Finder", lambda: self._reveal_in_finder(path))
        menu.addSeparator()
        menu.addAction("Удалить", lambda: self._delete_document(path))
        menu.exec(self.objectsTree.mapToGlobal(pos))

    def _show_folder_context_menu(self, item, pos):
        """ПКМ по строке папки -- object_dir берётся из данных строки
        (UserRole), а не восстанавливается из имени относительно
        OUTPUT_DIR: при произвольной вложенности папка может лежать на
        любой глубине, а не только прямо в OUTPUT_DIR."""
        object_dir = item.data(0, Qt.ItemDataRole.UserRole)["path"]
        menu = QMenu(self)
        menu.addAction("Новая подпапка", lambda: self._create_subfolder_dialog(object_dir))
        menu.addAction("Создать документ здесь", lambda: self._create_document_in_object(object_dir))
        menu.addSeparator()
        menu.addAction("Переименовать папку", lambda: self._rename_object_dialog(object_dir))
        menu.addAction("Показать в Finder", lambda: self._reveal_in_finder(object_dir))
        menu.addSeparator()
        menu.addAction("Удалить папку", lambda: self._delete_object_dialog(object_dir))
        menu.exec(self.objectsTree.mapToGlobal(pos))

    def _reveal_in_finder(self, path):
        """Открывает Finder с выделенным файлом/папкой -- macOS-
        специфично (`open -R`), как и сам пункт меню в референсе
        ("Показать в Finder"): приложение не претендует на
        кроссплатформенность."""
        subprocess.run(["open", "-R", str(path)])

    def _duplicate_document(self, path):
        """«Дублировать» -- копирует .json документа в той же папке
        под новым именем (см. workspace.duplicate_document())."""
        workspace.duplicate_document(path)
        self._refresh_objects_tree()

    def _delete_document(self, path):
        """«Удалить» документ -- необратимо, с подтверждением. Если
        удаляется текущий открытый документ, форма сбрасывается
        (как при старте, документов больше нет для показа)."""
        reply = QMessageBox.question(
            self, "Удалить документ",
            f"Удалить документ «{path.stem}»? Это необратимо.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return
        if path == self._current_document_path:
            if self.equipment_type.id == "constructor":
                self._reset_constructor_form()
            else:
                self._reset_form()
                self._switch_view("document")
        if getattr(self, "_database_preview_path", None) == path:
            # Удалили именно тот документ, что сейчас показан справа в
            # «Базе документов» -- иначе там осталось бы превью уже
            # несуществующего файла (см. _show_document_preview()).
            self._database_preview_path = None
            self._clear_database_preview()
            self.databaseHintLabel.setText("Выберите документ слева, чтобы посмотреть его содержимое.")
            self.databaseHintLabel.setVisible(True)
            self.databaseBreadcrumbLabel.setText("База документов")
        path.unlink(missing_ok=True)
        self._refresh_objects_tree()
        self._update_breadcrumb()

    def _create_document_in_object(self, folder_dir):
        """«Создать документ здесь» из меню папки / из выпадающего меню
        кнопки «Создать документ» -- folder_dir уже существует (папка
        выбрана по ПКМ либо из списка существующих объектов), в отличие
        от _create_document_in_new_object() ниже, где папку сначала
        нужно создать."""
        doc_path = workspace.create_document(folder_dir, self.equipment_type.id)
        self._refresh_objects_tree()
        self._open_document(doc_path)

    def _create_subfolder_dialog(self, parent_dir):
        """«Новая подпапка» из ПКМ-меню папки -- дерево «Объекты»
        поддерживает произвольную вложенность (docs/design/
        вводная_часть.html: "Папки -- произвольная вложенность"),
        workspace.create_object() уже принимает любой base_dir --
        ограничение было только в построении дерева (см.
        _build_objects_tree_level()), не в файловой модели."""
        name, ok = QInputDialog.getText(self, "Новая подпапка", "Название папки:")
        if not ok or not name.strip():
            return
        workspace.create_object(name, base_dir=parent_dir)
        self._refresh_objects_tree()

    def _rename_object_dialog(self, object_dir):
        """«Переименовать папку» -- в отличие от документа, у папки
        реальное имя ровно совпадает с именем папки на диске, так что
        переименование осмысленно и видно в дереве сразу. Открытый
        документ может лежать не только прямо внутри object_dir, но и в
        любой её вложенной подпапке (произвольная глубина) -- проверяем
        через parents, а не точное равенство .parent, и пересчитываем
        путь заменой только переименованного префикса."""
        new_name, ok = QInputDialog.getText(
            self, "Переименовать папку", "Новое название:", text=object_dir.name,
        )
        if not ok or not new_name.strip():
            return
        new_dir = workspace.rename_object(object_dir, new_name)
        if self._current_document_path is not None and (
            self._current_document_path == object_dir
            or object_dir in self._current_document_path.parents
        ):
            rel = self._current_document_path.relative_to(object_dir)
            self._current_document_path = new_dir / rel
        self._refresh_objects_tree()
        self._update_breadcrumb()

    def _delete_object_dialog(self, object_dir):
        """«Удалить папку» -- необратимо, удаляет папку целиком со всеми
        документами и вложенными подпапками, число документов (на любой
        глубине, rglob -- не только прямо внутри object_dir)
        показывается в подтверждении, чтобы не удалить что-то по
        ошибке."""
        doc_count = len(list(object_dir.rglob("*.json")))
        reply = QMessageBox.question(
            self, "Удалить папку",
            f"Удалить папку «{object_dir.name}» и все документы внутри ({doc_count})? Это необратимо.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return
        if self._current_document_path is not None and (
            self._current_document_path == object_dir
            or object_dir in self._current_document_path.parents
        ):
            if self.equipment_type.id == "constructor":
                self._reset_constructor_form()
            else:
                self._reset_form()
                self._switch_view("document")
        workspace.delete_object(object_dir)
        self._refresh_objects_tree()
        self._update_breadcrumb()

    def _open_document(self, path):
        """Открывает документ дерева. Если это уже открытый документ --
        просто переключает вид, не трогая форму (иначе несохранённые
        правки терялись бы). Иначе: тихо сохраняет текущий документ (если
        он был), сбрасывает форму и наполняет данными выбранного —
        порядок как в _reset_form()/_fill_ui_from_project(), см. план
        Фазы 2.

        Конструктор документов делит с трубопроводом objectsTree/дерево
        объектов (см. equipment_type.id == "constructor" в __init__), но
        не TOC/специалистов/индикатор несохранённых правок ниже -- их у
        него просто нет, см. _open_constructor_document()."""
        if self.equipment_type.id == "constructor":
            self._open_constructor_document(path)
            return
        if path == self._current_document_path:
            self._switch_view("document")
            return

        if self._current_document_path is not None:
            try:
                self.file_handler._save_current_project()
            except Exception as e:
                print(f"Не удалось автосохранить текущий документ: {e}")
            # Оглавление уходящего документа -- дочерние строки под его
            # строкой в дереве -- снимается: статус разделов считается
            # по живым виджетам формы, а они сейчас переиспользуются под
            # другой документ (см. _reset_form()).
            old_item = self._find_document_tree_item(self._current_document_path)
            if old_item is not None:
                old_item.takeChildren()

        self._reset_form()
        project = Project.load_from_file(path)
        self.file_handler._fill_ui_from_project(project)
        self._current_document_path = path
        self._document_dirty = False
        self._update_dirty_indicator()
        saved_indices = {
            name: project.report_data.get(name) for name in self.SPECIALIST_COMBO_NAMES
        }
        self._refresh_program_specialist_combo(saved_indices)
        self._switch_view("document")
        doc_item = self._find_document_tree_item(path)
        if doc_item is not None:
            self._populate_document_toc(doc_item)
        self._update_breadcrumb()

    def _reset_constructor_form(self):
        """Сбрасывает конструктор документов под другой/новый документ
        дерева -- аналог _reset_form() для остальных типов, но по всем
        слотам (см. _CONSTRUCTOR_SLOTS): у конструктора нет
        фиксированного списка виджетов реквизитов, они создаются заново
        под набор полей конкретного варианта (_render_slot_fields()).

        _clear_included_block(slot) сам по себе не убирает уже
        отрисованные динамические поля прошлого варианта (только прячет
        fields_panel и возвращает included_list к пустому плейсхолдеру,
        см. его докстринг) -- без явной очистки dynamic_field_names/
        PLAIN_TEXT_EDIT_NAMES/строк fields_layout здесь виджеты прошлого
        документа остались бы висеть (пустой текст, но всё ещё
        зарегистрированные) и просочились бы в get_form_data() следующего
        «Сохранить проект». Та же очистка, что в начале
        _render_slot_fields(), просто без немедленной перерисовки под
        новый вариант -- открывающий вызывающий код (_open_constructor_document())
        сам наполнит слоты заново через _fill_ui_from_project()."""
        for slot in self._CONSTRUCTOR_SLOTS:
            dynamic_names = self._dynamic_field_names[slot]
            for name in dynamic_names:
                if name in self.PLAIN_TEXT_EDIT_NAMES:
                    self.PLAIN_TEXT_EDIT_NAMES.remove(name)
            self._dynamic_field_names[slot] = []
            layout = self._slot_widget(slot, "fields_layout")
            while layout.rowCount():
                layout.removeRow(0)
            self._clear_included_block(slot)
        self.data = {}
        self._completed_steps = set()
        self._current_document_path = None

    def _open_constructor_document(self, path):
        """Открывает документ дерева для конструктора -- упрощённый
        аналог _open_document() для трубопровода: без TOC, специалистов и
        индикатора несохранённых правок (их у конструктора нет), но с тем
        же порядком действий -- тихое автосохранение текущего документа,
        сброс формы, наполнение из выбранного. В конце переключает
        активити-бар на «Редактор документов» (см.
        _switch_activity_view()) -- открытие документа из «Базы
        документов» должно сразу показать его для редактирования, а не
        оставлять на странице дерева."""
        if path == self._current_document_path:
            self._switch_activity_view("editor")
            return
        if self._current_document_path is not None:
            try:
                self.file_handler._save_current_project()
            except Exception as e:
                print(f"Не удалось автосохранить текущий документ: {e}")
        self._reset_constructor_form()
        project = Project.load_from_file(path)
        self.file_handler._fill_ui_from_project(project)
        self._current_document_path = path
        self._switch_activity_view("editor")

    def _show_document_preview(self, path):
        """Клик по документу в «Базе документов» (только конструктор) --
        показывает его содержимое справа (databasePreviewContainer), как
        в правой части «Редактора документов», но ИЗ ОТДЕЛЬНЫХ статичных
        QLabel, читая Project прямо с диска -- не трогает self.data/
        PLAIN_TEXT_EDIT_NAMES/_current_document_path текущего
        редактируемого документа. Открытие для правки -- отдельное
        явное действие (ПКМ -> «Открыть в редакторе документов», см.
        _show_document_context_menu()), просмотр не должен незаметно
        подменять то, что сейчас редактируется."""
        from ..models.project import Project
        from ..services.title_variants_store import get_all_field_labels

        self._database_preview_path = path
        try:
            project = Project.load_from_file(path)
        except (OSError, ValueError) as e:
            self._clear_database_preview()
            self.databaseHintLabel.setText(f"Не удалось прочитать документ: {e}")
            self.databaseHintLabel.setVisible(True)
            return

        self.databaseBreadcrumbLabel.setText(f"База документов / {path.stem}")
        self._clear_database_preview()

        labels = get_all_field_labels()
        rendered_any = False
        for slot in self._CONSTRUCTOR_SLOTS:
            variant_id = project.report_data.get(f"_included_{slot}_variant_id")
            if not variant_id:
                continue
            get_all_variants = self._slot_store(slot).get_all_variants
            variant = get_all_variants().get(variant_id)
            if variant is None:
                continue
            rendered_any = True
            self.databasePreviewLayout.addWidget(
                self._build_database_preview_block(slot, variant, project.report_data, labels)
            )

        self.databaseHintLabel.setVisible(not rendered_any)
        if not rendered_any:
            self.databaseHintLabel.setText("В этом документе нет заполненных блоков.")

    def _clear_database_preview(self):
        """Убирает предыдущее превью из databasePreviewLayout (полная
        пересборка на каждый клик, список полей короткий -- десятки
        строк, не тысячи, отдельный diff/переиспользование виджетов не
        нужны)."""
        layout = self.databasePreviewLayout
        while layout.count():
            item = layout.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()

    def _build_database_preview_block(self, slot, variant, report_data, labels):
        """Один блок превью (заголовок варианта + его реквизиты) -- та же
        пара «включённый блок / реквизиты», что рисует _process_block_drop()/
        _render_slot_fields() в самом редакторе, только read-only и без
        привязки к self.data (значения читаются напрямую из report_data
        уже сохранённого документа, а не с живых виджетов формы).

        Клик по заголовку сворачивает/разворачивает реквизиты -- то же
        поведение и тот же шеврон, что у _toggle_fields_panel() в самом
        редакторе (см. docs/design/вводная_часть.html, buildDbStructureBlock())."""
        container = QWidget()
        layout = QVBoxLayout(container)
        layout.setContentsMargins(0, 0, 0, 18)
        layout.setSpacing(8)

        header = QWidget()
        header_layout = QHBoxLayout(header)
        header_layout.setContentsMargins(0, 0, 0, 0)
        header_layout.setSpacing(6)
        header.setCursor(Qt.CursorShape.PointingHandCursor)
        icon_label = QLabel()
        icon_label.setPixmap(icons.render("file-text", "#0a84ff", 15))
        icon_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        header_layout.addWidget(icon_label)
        title_label = QLabel(variant.document_title)
        title_label.setObjectName("databasePreviewTitle")
        title_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        header_layout.addWidget(title_label)
        header_layout.addStretch()
        chevron_label = QLabel()
        chevron_label.setPixmap(icons.render("chevron-down", "#8e8e93", 13))
        chevron_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        header_layout.addWidget(chevron_label)
        layout.addWidget(header)

        form = QFormLayout()
        form.setLabelAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignTop)
        for field_id in variant.subtitle_fields:
            widget_name = self._slot_placeholder_name(slot, field_id)
            value = report_data.get(widget_name) or "—"
            label = QLabel(labels.get(field_id, field_id))
            label.setObjectName("databasePreviewFieldLabel")
            value_label = QLabel(str(value))
            value_label.setObjectName("databasePreviewFieldValue")
            value_label.setWordWrap(True)
            form.addRow(label, value_label)
        # Обёртка QFormLayout в собственный QWidget (а не layout.addLayout(form))
        # -- когда несколько таких блоков подряд лежат в одном QVBoxLayout
        # (databasePreviewLayout, один блок на слот title/intro), при прямой
        # вложенности layout-в-layout Qt считает heightForWidth многострочных
        # QLabel неверно: блоку достаётся высота меньше его собственного
        # sizeHint(), и последняя многострочная строка визуально наезжает на
        # следующий блок снизу (воспроизведено отдельным скриптом на паре
        # смежных блоков с многострочным ФИО). Отдельный QWidget с setLayout()
        # даёт слою один авторитетный heightForWidth вместо рассогласованной
        # совместной развёртки вложенных анонимных layout'ов.
        form_widget = QWidget()
        form_widget.setLayout(form)
        layout.addWidget(form_widget)

        header.mousePressEvent = functools.partial(
            self._toggle_database_preview_block, form_widget, chevron_label
        )
        return container

    def _toggle_database_preview_block(self, form_widget, chevron_label, event):
        """Левый клик по заголовку блока превью в «Базе документов» --
        сворачивает/разворачивает его реквизиты. См. _toggle_fields_panel()
        -- тот же приём (чевron меняется на chevron-right/chevron-down)."""
        if event.button() != Qt.MouseButton.LeftButton:
            return
        visible = not form_widget.isVisible()
        form_widget.setVisible(visible)
        chevron_label.setPixmap(
            icons.render("chevron-down" if visible else "chevron-right", "#8e8e93", 13)
        )

    def _find_tree_item_by_path(self, kind, path):
        """Ищет строку objectsTree заданного "kind" ("folder" или
        "document"), чей путь в Qt.ItemDataRole.UserRole совпадает с
        path -- рекурсивно, папки теперь произвольной вложенности (см.
        _build_objects_tree_level()), плоский обход двух уровней уже не
        покрывает всё дерево. Дерево небольшое (десятки строк), полный
        обход на каждый вызов дешевле, чем держать отдельный кэш
        path->item в синхронизации с _refresh_objects_tree()."""
        def search(item):
            data = item.data(0, Qt.ItemDataRole.UserRole)
            if isinstance(data, dict) and data.get("kind") == kind and data.get("path") == path:
                return item
            for i in range(item.childCount()):
                found = search(item.child(i))
                if found is not None:
                    return found
            return None

        for i in range(self.objectsTree.topLevelItemCount()):
            found = search(self.objectsTree.topLevelItem(i))
            if found is not None:
                return found
        return None

    def _find_document_tree_item(self, path):
        """Документ -- самый частый случай поиска (_open_document(),
        _update_breadcrumb() и т.п.), отдельная обёртка ради краткости
        вызовов."""
        return self._find_tree_item_by_path("document", path)

    def _update_breadcrumb(self):
        """Обновляет breadcrumbLabel над лентой документа: объект /
        документ / текущий раздел оглавления. Ярлык документа берётся
        готовым из строки дерева (там уже имя файла из
        workspace.list_documents()), а не пересчитывается заново.

        Объект и документ -- кликабельные ссылки (Фаза 11, как revealFolder()/
        revealDoc() в референсе, docs/design/pipeline_sidebar_mockup.html)
        через встроенную поддержку rich-text гиперссылок в QLabel
        (setOpenExternalLinks(False) + сигнал linkActivated, см.
        _on_breadcrumb_link_activated()) -- три отдельных виджета под
        три сегмента заводить не пришлось. Раздел -- не ссылка, как и в
        референсе (там у него нет своего reveal-обработчика).

        Только для трубопровода -- у него есть свой breadcrumbLabel над
        лентой формы. Конструктор документов делит с трубопроводом общий
        objectsTree/_prompt_document_path()/_rename_object_dialog() и т.п.
        (см. equipment_type.id == "constructor" в __init__), но своего
        breadcrumbLabel не имеет (у него свой crumbLabel -- название первого
        заполненного раздела, не путь дерева, см. _update_crumb()) --
        без этого выхода вызов падал бы AttributeError при любом
        сохранении/переименовании через общие методы."""
        if self.equipment_type.id != "pipeline":
            return
        if self._current_document_path is None:
            self.breadcrumbLabel.setText("")
            self.breadcrumbLabel.setVisible(False)
            return
        object_name = Path(self._current_document_path).parent.name
        doc_item = self._find_document_tree_item(self._current_document_path)
        doc_label = doc_item.text(0) if doc_item is not None else ""
        section = self._current_toc_active_item.text(0) if self._current_toc_active_item is not None else ""
        parts = []
        if object_name:
            parts.append(f'<a href="object" style="color:inherit; text-decoration:none;">{html.escape(object_name)}</a>')
        if doc_label:
            parts.append(f'<a href="document" style="color:inherit; text-decoration:none;">{html.escape(doc_label)}</a>')
        if section:
            parts.append(html.escape(section))
        self.breadcrumbLabel.setText(" / ".join(parts))
        self.breadcrumbLabel.setVisible(True)

    def _on_breadcrumb_link_activated(self, href):
        """Клик по сегменту breadcrumb -- разворачивает сайдбар (если
        свёрнут), раскрывает нужную строку дерева и на мгновение
        подсвечивает её (см. revealFolder()/revealDoc() в референсе)."""
        if self._sidebar_collapsed:
            self._toggle_sidebar()
        if href == "document":
            item = self._find_document_tree_item(self._current_document_path)
        else:
            # Прямой родитель документа -- не обязательно верхний
            # уровень дерева при произвольной вложенности папок, ищем по
            # пути, а не по имени среди top-level строк.
            item = self._find_tree_item_by_path("folder", self._current_document_path.parent)
        if item is None:
            return
        # Разворачивает ВСЕХ предков, а не только непосредственного
        # родителя -- при произвольной вложенности папок строка может
        # быть скрыта на любой глубине, если пользователь свернул
        # промежуточную папку вручную.
        ancestor = item.parent()
        while ancestor is not None:
            ancestor.setExpanded(True)
            ancestor = ancestor.parent()
        item.setExpanded(True)
        self.objectsTree.scrollToItem(item)
        self._flash_tree_item(item)

    def _flash_tree_item(self, item):
        """Кратковременная подсветка строки дерева (700ms, как
        folder-flash/doc-flash в референсе) -- в отличие от
        _set_toc_active_item(), это одноразовая вспышка (взгляду
        помочь найти строку после клика по breadcrumb), а не постоянная
        подсветка "текущий раздел"."""
        flash_color = QColor("#0a84ff")
        flash_color.setAlpha(38)
        item.setBackground(0, flash_color)
        QTimer.singleShot(700, lambda: item.setData(0, Qt.ItemDataRole.BackgroundRole, None))

    def _make_dot_icon(self, color):
        """Рисует маленький закрашенный кружок как QIcon -- в проекте нет
        инфраструктуры иконок-ресурсов (.qrc), проще нарисовать пиксмап
        на лету (см. также _make_check_icon()/_make_circle_icon()/
        _make_lock_icon() -- тот же приём для иконок TOC, Фаза 6)."""
        pixmap = QPixmap(10, 10)
        pixmap.fill(Qt.GlobalColor.transparent)
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setBrush(QColor(color))
        painter.setPen(Qt.PenStyle.NoPen)
        painter.drawEllipse(1, 1, 8, 8)
        painter.end()
        return QIcon(pixmap)

    def _connect_dirty_tracking(self):
        """Подключает _mark_dirty к сигналу изменения на каждом виджете
        формы -- по тем же спискам, что _reset_form()/get_form_data(),
        так что набор отслеживаемых полей не может разойтись с реальным
        составом документа."""
        for name in self.PLAIN_TEXT_EDIT_NAMES:
            getattr(self, name).textChanged.connect(self._mark_dirty)
        for name in self.COMBO_BOX_NAMES:
            getattr(self, name).currentIndexChanged.connect(self._mark_dirty)
        for name in self.DATE_EDIT_NAMES:
            getattr(self, name).dateChanged.connect(self._mark_dirty)
        for name in self.SPIN_BOX_NAMES:
            getattr(self, name).valueChanged.connect(self._mark_dirty)
        for name in self.TABLE_WIDGET:
            getattr(self, name).itemChanged.connect(self._mark_dirty)

    def _mark_dirty(self, *args):
        """Слот на любое изменение поля формы. Срабатывает и во время
        программного заполнения формы при открытии документа
        (_fill_ui_from_project() дёргает те же сигналы) -- в этот момент
        _current_document_path ещё None (сбрасывается в _reset_form() до
        заполнения), поэтому ложных срабатываний нет.

        Помимо самого факта "документ изменён" (once-флаг, дальше не
        трогается до сохранения/переоткрытия), пересчитывает иконки
        ✓/○/🔒 у TOC текущего документа на КАЖДОЕ изменение, не только
        первое -- для 18 из 20 разделов (не завязанных на кнопку шага
        STEP_ORDER) статус "готово" определяется именно заполненностью
        полей (см. _groupbox_done()), а значит должен обновляться
        живьём по мере ввода, а не только по клику кнопки расчёта."""
        if self._current_document_path is None:
            return
        if not self._document_dirty:
            self._document_dirty = True
            self._update_dirty_indicator()
        self._update_toc_progress()

    def _update_dirty_indicator(self):
        """Ставит/снимает точку-иконку у строки текущего документа в
        objectsTree."""
        item = self._find_document_tree_item(self._current_document_path)
        if item is None:
            return
        item.setIcon(0, self._dirty_icon if self._document_dirty else QIcon())

    def _create_object_dialog(self):
        """«Создать объект»: спрашивает название, создаёт папку, сразу
        обновляет дерево, чтобы новый объект стало видно."""
        name, ok = QInputDialog.getText(self, "Новый объект", "Название объекта:")
        if not ok or not name.strip():
            return
        workspace.create_object(name)
        self._refresh_objects_tree()

    def _create_document_dialog(self):
        """«Создать документ» -- всплывающее меню от кнопки: список
        существующих объектов + «Новый объект…» снизу (соответствует
        шагу 1 референсного мокапа, docs/design/pipeline_sidebar_mockup.html,
        #stepObject). Шаг 2 референса ("Использовать шаблон"/"Создать
        новый шаблон") сознательно не реализован -- реальной
        инфраструктуры нескольких шаблонов на equipment_type нет
        (find_template() захардкожен на один файл в config.py), а
        generate_template() -- разовая операция подготовки заготовки
        под ручную правку в Word (см. template_generator.py), не
        предназначенная запускаться при каждом создании документа.
        Фиктивный шаг выбора шаблона, который ничего не переключает,
        хуже, чем его отсутствие."""
        menu = QMenu(self)
        for object_name in workspace.list_objects():
            object_dir = workspace.OUTPUT_DIR / object_name
            menu.addAction(object_name, lambda d=object_dir: self._create_document_in_object(d))
        if menu.actions():
            menu.addSeparator()
        menu.addAction("Новый объект…", self._create_document_in_new_object)
        button = self.sidebarBtn_createDocument
        menu.exec(button.mapToGlobal(button.rect().bottomLeft()))

    def _create_document_in_new_object(self):
        """«Новый объект…» в меню «Создать документ» -- спрашивает имя,
        создаёт объект (папку верхнего уровня) и сразу документ внутри
        него за один шаг (в отличие от отдельной кнопки «Создать
        объект», после которой пришлось бы ещё раз открывать это же
        меню)."""
        name, ok = QInputDialog.getText(self, "Новый объект", "Название объекта:")
        if not ok or not name.strip():
            return
        object_dir = workspace.create_object(name)
        self._create_document_in_object(object_dir)

    def _show_templates_menu(self):
        """«Мои шаблоны» -- показывает реально существующий шаблон
        (find_template() поддерживает ровно один файл на equipment_type,
        см. TEMPLATE_PATHS_BY_TYPE в config.py -- инфраструктуры выбора
        между несколькими шаблонами в проекте нет). Пункт
        информационный, недоступен для клика -- переключать нечего,
        пока шаблон один; помечать его "по умолч." было бы обманчиво,
        раз альтернативы не существует."""
        from ..config import find_template
        menu = QMenu(self)
        try:
            template_path = find_template(self.equipment_type.id)
            action = menu.addAction(template_path.stem)
        except FileNotFoundError:
            action = menu.addAction("Шаблон не найден")
        action.setEnabled(False)
        button = self.sidebarBtn_templates
        menu.exec(button.mapToGlobal(button.rect().bottomLeft()))

    def _collect_toc_groupboxes(self):
        """Собирает top-level QGroupBox'ы ленты tab_document в порядке
        компоновки -- статический список, форма всегда одна и та же
        независимо от того, какой документ открыт (в отличие от
        _current_toc_items, дочерних строк дерева, которые заново
        строятся под каждым открываемым документом, см.
        _populate_document_toc())."""
        layout = self.tab_document_scrollContent.layout()
        groupboxes = []
        for i in range(layout.count()):
            widget = layout.itemAt(i).widget()
            if isinstance(widget, QGroupBox):
                groupboxes.append(widget)
        return groupboxes

    def _populate_document_toc(self, doc_item):
        """Строит оглавление как дочерние строки под строкой документа в
        objectsTree (Фаза 6) -- по одной на каждый groupbox из
        _toc_groupboxes, в том же порядке. Только для ТЕКУЩЕГО открытого
        документа: статус разделов считается по живым виджетам формы
        (см. _groupbox_done()), а данные другого документа в них не
        загружены. Вызывается из _open_document() после заполнения
        формы; TOC уходящего документа снимается там же через
        item.takeChildren()."""
        self._current_toc_items = []
        for groupbox in self._toc_groupboxes:
            item = QTreeWidgetItem([groupbox.title()])
            item.setData(0, Qt.ItemDataRole.UserRole, {"kind": "toc", "groupbox": groupbox})
            doc_item.addChild(item)
            self._current_toc_items.append(item)
        doc_item.setExpanded(True)
        self._current_toc_active_item = None
        self._update_toc_progress()

    def _scroll_to_toc_item(self, item):
        """Клик по строке оглавления в дереве -- скроллит
        tab_document_scroll так, чтобы верх выбранного раздела оказался
        у верха видимой области (не ensureWidgetVisible(): для разделов
        выше высоты вьюпорта он подтянул бы минимальным движением, а не
        встык к началу раздела). Заблокированные строки (см.
        _toc_lock_reason()) не скроллят -- у них уже снят
        Qt.ItemFlag.ItemIsEnabled в _update_toc_item_style(), но клик
        может дойти и до отключённого item'а, проверка дублируется."""
        groupbox = item.data(0, Qt.ItemDataRole.UserRole)["groupbox"]
        if self._toc_lock_reason(groupbox) is not None:
            return
        self._suppress_toc_spy = True
        self.tab_document_scroll.verticalScrollBar().setValue(groupbox.y())
        self._set_toc_active_item(item)
        self._suppress_toc_spy = False
        self._update_breadcrumb()

    def _on_document_scrolled(self, value):
        """Скролл-спай: при ручной прокрутке ленты подсвечивает строку
        последнего раздела, чей верх уже проскроллен (groupbox.y() <=
        value), в оглавлении текущего открытого документа. Отключается
        на время программного скролла по клику (_suppress_toc_spy),
        иначе клик спорил бы сам с собой."""
        if self._suppress_toc_spy or not self._current_toc_items:
            return
        current = 0
        for i, groupbox in enumerate(self._toc_groupboxes):
            if groupbox.y() <= value:
                current = i
        self._set_toc_active_item(self._current_toc_items[current])
        self._update_breadcrumb()

    def _set_toc_active_item(self, item):
        """Полная заливка активной строки TOC синим, как в референсе
        (docs/design/pipeline_sidebar_mockup.html, .toc-active) -- через
        setBackground()/setForeground() напрямую на item, а не через
        нативное выделение QTreeWidget: нативная подсветка платформенно
        по-разному ведёт себя в фокусе/без фокуса окна, а тут нужна
        подсветка "текущий раздел", не зависящая от фокуса. Сброс роли в
        None (не просто прозрачный цвет) возвращает предыдущему item'у
        подлинный вид по умолчанию."""
        if self._current_toc_active_item is not None:
            self._current_toc_active_item.setData(0, Qt.ItemDataRole.BackgroundRole, None)
            self._current_toc_active_item.setData(0, Qt.ItemDataRole.ForegroundRole, None)
        item.setBackground(0, QColor("#0a84ff"))
        item.setForeground(0, QColor("#ffffff"))
        self._current_toc_active_item = item

    # Пункты оглавления, у которых есть осмысленное понятие "выполнено" из
    # реальной кнопки-расчёта -- только эти два groupbox'а физически
    # содержат кнопки шагов из equipment_types.py STEP_ORDER
    # (segments/thickness -> thick_group, strength/residual_life ->
    # calc_appendix_group). У остальных 18 разделов такой кнопки нет --
    # для них статус считает _groupbox_done() по заполненности полей.
    TOC_PROGRESS_GROUPS = {
        "thick_group": {"segments", "thickness"},
        "calc_appendix_group": {"strength", "residual_life"},
    }

    def _groupbox_done(self, groupbox):
        """"Готово" для раздела оглавления. Для thick_group/
        calc_appendix_group -- точная семантика STEP_ORDER (см.
        TOC_PROGRESS_GROUPS): у них есть настоящая кнопка-расчёт, это
        строго более надёжный сигнал, чем заполненность полей, трогать
        не нужно. Для остальных 18 разделов, где такой кнопки нет --
        обобщённая проверка: все виджеты формы внутри groupbox'а (те же
        типы, что в _reset_form()/get_form_data()) заполнены.

        QDateEdit и QComboBox намеренно не проверяются. У даты всегда
        есть значение (по умолчанию сегодняшнее) -- "пустой" даты, в
        отличие от текстового поля, не бывает. У комбобоксов currentIndex()
        == 0 НЕ значит "не заполнено": report_title -- редактируемый
        комбобокс ровно с одним пунктом (реальный дефолтный текст, не
        плейсхолдер-заглушка), work_medium -- 5 реальных веществ без
        пустого варианта; для обоих индекс 0 -- уже осмысленный выбор,
        который оператор имеет полное право оставить как есть (проверено
        headless-тестом: report_title.count() == 1 в свежесозданном
        документе -- проверка "index == 0 -> пусто" там попросту неверна)."""
        steps = self.TOC_PROGRESS_GROUPS.get(groupbox.objectName())
        if steps is not None:
            return steps <= self._completed_steps
        for edit in groupbox.findChildren(QPlainTextEdit):
            if not edit.toPlainText().strip():
                return False
        for spin in groupbox.findChildren(QSpinBox):
            if spin.value() == spin.minimum():
                return False
        for table in groupbox.findChildren(QTableWidget):
            if table.rowCount() == 0:
                return False
        return True

    def _toc_lock_reason(self, groupbox):
        """Возвращает заголовок блокирующего раздела, если groupbox
        заблокирован (🔒), иначе None. Единственная реальная зависимость
        в данных этого приложения: calc_appendix_group (шаги
        strength/residual_life) требует thick_group (шаги
        segments/thickness) выполненным целиком -- тот же порядок, что
        и в _check_prerequisite(). Другие 18 разделов никогда не
        блокируются -- в реальной модели STEP_ORDER больше зависимостей
        нет; выдумывать их ради сходства с иллюстрацией в референсе не
        нужно (референсный "Приложение 8 заблокировано Приложением 6" --
        условный пример для мокапа, не основанный на реальных данных
        этого приложения)."""
        if groupbox.objectName() == "calc_appendix_group" and not self._groupbox_done(self.thick_group):
            return self.thick_group.title()
        return None

    def _update_toc_item_style(self, item, groupbox):
        """Ставит иконку ✓/○/🔒 и тултип у одной строки TOC."""
        lock_reason = self._toc_lock_reason(groupbox)
        if lock_reason:
            item.setIcon(0, self._toc_lock_icon)
            item.setToolTip(0, f"Сначала завершите «{lock_reason}»")
            item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEnabled)
        else:
            item.setFlags(item.flags() | Qt.ItemFlag.ItemIsEnabled)
            item.setToolTip(0, "")
            item.setIcon(0, self._toc_check_icon if self._groupbox_done(groupbox) else self._toc_circle_icon)

    def _update_toc_progress(self):
        """Обновляет иконки ✓/○/🔒 у всех строк TOC текущего открытого
        документа. Единая точка входа после любого изменения, влияющего
        на статус разделов -- клика по кнопке шага STEP_ORDER (см. 4
        вызова ниже) или правки любого поля формы (см. _mark_dirty()).
        Безопасно вызывать и когда документ не открыт -- _current_toc_items
        тогда пуст, zip() ничего не делает."""
        for item, groupbox in zip(self._current_toc_items, self._toc_groupboxes):
            self._update_toc_item_style(item, groupbox)

    def _make_check_icon(self):
        """✓ готово -- зелёный кружок с белой галочкой (аналог
        ti-circle-check из референса). Рисуется на лету тем же приёмом,
        что и _make_dot_icon()."""
        pixmap = QPixmap(14, 14)
        pixmap.fill(Qt.GlobalColor.transparent)
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setBrush(QColor("#30d158"))
        painter.setPen(Qt.PenStyle.NoPen)
        painter.drawEllipse(1, 1, 12, 12)
        pen = QPen(QColor("#ffffff"))
        pen.setWidthF(1.6)
        pen.setCapStyle(Qt.PenCapStyle.RoundCap)
        pen.setJoinStyle(Qt.PenJoinStyle.RoundJoin)
        painter.setPen(pen)
        painter.drawLine(QPointF(4, 7.2), QPointF(6.2, 9.5))
        painter.drawLine(QPointF(6.2, 9.5), QPointF(10, 4.7))
        painter.end()
        return QIcon(pixmap)

    def _make_circle_icon(self):
        """○ не готово -- серый контур кружка (аналог ti-circle)."""
        pixmap = QPixmap(14, 14)
        pixmap.fill(Qt.GlobalColor.transparent)
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        pen = QPen(QColor("#8e8e93"))
        pen.setWidthF(1.2)
        painter.setPen(pen)
        painter.setBrush(Qt.BrushStyle.NoBrush)
        painter.drawEllipse(2, 2, 10, 10)
        painter.end()
        return QIcon(pixmap)

    def _make_lock_icon(self):
        """🔒 заблокировано -- серый замок, тело + дужка (аналог
        ti-lock)."""
        pixmap = QPixmap(14, 14)
        pixmap.fill(Qt.GlobalColor.transparent)
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        color = QColor("#8e8e93")
        pen = QPen(color)
        pen.setWidthF(1.6)
        painter.setPen(pen)
        painter.setBrush(Qt.BrushStyle.NoBrush)
        painter.drawArc(3, 1, 8, 8, 0, 180 * 16)
        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(color)
        painter.drawRoundedRect(2, 6, 10, 7, 2, 2)
        painter.end()
        return QIcon(pixmap)

    def _update_report_buttons_visibility(self):
        """Скрывает кнопки "Выгрузить в Word"/"Сохранить проект"/"Открыть
        проект", пока открыты "Сотрудники"/"Приборы"/"Документы" -- это
        общие справочники компании, не часть текущего отчёта (см.
        EmployeesTabController, InstrumentsTabController,
        OrgDocsTabController), эти действия к ним не относятся."""
        is_directory_view = self._current_view in ("employees", "instruments", "orgdocs")
        self.pushButt_generateWord.setVisible(not is_directory_view)
        self.pushButton_saveProject.setVisible(not is_directory_view)
        self.pushButton_openProject.setVisible(not is_directory_view)

    def _cell_text(self, table, row, col):
        """Текст ячейки (row, col) независимо от того, обычный это
        QTableWidgetItem или виджет (QComboBox, см. _install_growable_combo())."""
        item = table.item(row, col)
        if item is not None:
            return item.text()
        cell_widget = table.cellWidget(row, col)
        if isinstance(cell_widget, QComboBox):
            return cell_widget.currentText()
        return ""

    def _update_p_rab_kgs(self):
        """Автоматически пересчитывает "Давление, кгс/см2" из "Давление,
        МПа" (1 МПа = 1/0,0980665 кгс/см2, точный коэффициент перевода) --
        p_rab_kgs доступно только для чтения (см. .ui), оператор вводит
        давление один раз в МПа."""
        try:
            mpa = parse_ru(self.p_rab_mpa.toPlainText())
        except ValueError:
            self.p_rab_kgs.setPlainText("")
            return
        self.p_rab_kgs.setPlainText(format_ru_fixed(mpa / 0.0980665, 1))

    def _sync_mirror_field(self, target, source_text, state_attr):
        """Копирует source_text в target, пока оператор не ввёл в target
        своё значение. "Тронуто" определяется не через "непусто" (иначе при
        обычном посимвольном наборе в source-поле target подхватил бы
        только самый первый введённый символ и на этом застрял бы -- после
        него target уже "непусто", хотя это ещё чисто автоподстановка) -- а
        через сравнение текущего текста target с последним значением,
        которое сюда же поставила именно эта автоподстановка (state_attr,
        хранится на self). Если оператор изменил target -- текст разойдётся
        с state_attr, и подстановка остановится."""
        current = target.toPlainText()
        if current and current != getattr(self, state_attr, None):
            return
        target.setPlainText(source_text)
        setattr(self, state_attr, source_text)

    # Пары (target, source, state_attr) для _sync_mirror_field() -- вынесены
    # сюда одним списком, чтобы завести его один раз и переиспользовать и в
    # сигналах __init__ (через свои _update_*_display()), и в
    # _seed_mirror_states_after_load() ниже.
    MIRROR_FIELD_PAIRS = (
        ("calc_temp", "work_temp", "_calc_temp_auto_value"),
        ("pnevmo_obj_naznach", "obj_naznach", "_pnevmo_obj_naznach_auto_value"),
        ("ae_zakl_obj_control", "pnevmo_obj_naznach", "_ae_zakl_obj_control_auto_value"),
        ("calc_years_operation", "years_of_operation", "_calc_years_operation_auto_value"),
    )

    def _seed_mirror_states_after_load(self):
        """Восстанавливает state_attr у "зеркал, пока не тронутых"
        (_sync_mirror_field(): calc_temp/pnevmo_obj_naznach/
        ae_zakl_obj_control/calc_years_operation) сразу после загрузки
        проекта (FileHandler._fill_ui_from_project()) -- единственный
        признак "не тронуто" там -- current-текст target совпадает с
        state_attr, а на свежем окне state_attr всегда None. Без этого
        восстановленное из JSON непустое значение target (даже если оно
        никогда не редактировалось вручную и просто совпадает с source)
        навсегда воспринималось бы как "оператор его трогал" -- дальше
        правки в source-поле (например years_of_operation, раздел 6)
        переставали подхватываться в target (calc_years_operation,
        Приложение 6), хотя раньше, до сохранения/открытия проекта,
        подхватывались нормально.

        Сеем состояние только когда target совпадает с source -- если они
        разошлись (оператор реально переопределил target перед
        сохранением), это и есть корректное "тронуто", трогать не нужно."""
        for target_name, source_name, state_attr in self.MIRROR_FIELD_PAIRS:
            target = getattr(self, target_name)
            source = getattr(self, source_name)
            if target.toPlainText() == source.toPlainText():
                setattr(self, state_attr, target.toPlainText())

        # years_of_operation -- отдельный случай: источник не одно
        # текстовое поле, а вычисление report_year - year_start, см.
        # _update_years_of_operation_display(). Та же логика "seed только
        # если совпадает с вычисленным значением".
        try:
            computed = str(
                int(parse_ru(self.report_year.toPlainText()))
                - int(parse_ru(self.year_start.toPlainText()))
            )
        except ValueError:
            computed = None
        if computed is not None and self.years_of_operation.toPlainText() == computed:
            self._years_of_operation_auto_value = computed

    def _update_years_of_operation_display(self):
        """Подсказка years_of_operation (раздел 6, "Срок эксплуатации,
        лет") -- по умолчанию год составления отчёта (report_year,
        раздел 1) минус год ввода в эксплуатацию (year_start, раздел 6),
        но поле редактируемое: инженер может исправить вручную (например,
        если объект фактически простаивал часть срока), см.
        _sync_mirror_field()."""
        try:
            report_year = int(parse_ru(self.report_year.toPlainText()))
            year_start = int(parse_ru(self.year_start.toPlainText()))
        except ValueError:
            return
        self._sync_mirror_field(
            self.years_of_operation, str(report_year - year_start),
            "_years_of_operation_auto_value",
        )

    def _update_calc_temp_display(self):
        """Подсказка calc_temp (Приложение 6) значением work_temp (1. Общие
        данные) -- по умолчанию температура стенки принимается равной
        рабочей температуре среды, но поле редактируемое: дальше значение
        полностью в руках оператора, см. _sync_mirror_field()."""
        self._sync_mirror_field(self.calc_temp, self.work_temp.toPlainText(), "_calc_temp_auto_value")

    def _update_calc_years_operation_display(self):
        """Подсказка calc_years_operation (Приложение 6, "Остаточный
        ресурс -- исходные данные") значением years_of_operation (раздел 6)
        -- по умолчанию срок эксплуатации для расчёта остаточного ресурса
        равен сроку эксплуатации объекта, но поле редактируемое, см.
        _sync_mirror_field()."""
        self._sync_mirror_field(
            self.calc_years_operation, self.years_of_operation.toPlainText(),
            "_calc_years_operation_auto_value",
        )

    def _update_calc_p_rab_display(self):
        """Зеркалит p_rab_mpa (1. Общие данные) в read-only
        calc_p_rab_display (Приложение 6) -- само это поле не резолвится
        через widget_names_pipeline.py и в шаблон .docx не попадает, это
        чисто UI-подсказка (тот же приём, что и pnevmo_pressure_hint)."""
        self.calc_p_rab_display.setPlainText(self.p_rab_mpa.toPlainText())

    def _update_pnevmo_pressure(self):
        """Автоматически пересчитывает "Испытательное давление"
        (pnevmo_pressure, Приложения 8-9) из "Давление, кгс/см2"
        (p_rab_kgs, 1. Общие данные) по формуле p_rab_kgs * 1,25 --
        срабатывает при каждом изменении p_rab_kgs, перезаписывая
        прежнее значение поля (тот же принцип, что и у
        _update_final_deadline_date)."""
        try:
            p_rab_kgs = parse_ru(self.p_rab_kgs.toPlainText())
        except ValueError:
            self.pnevmo_pressure.setPlainText("")
            return
        self.pnevmo_pressure.setPlainText(format_ru_fixed(p_rab_kgs * 1.25, 1))

    def _fill_pnevmo_stages_table(self):
        """Автозаполнение table_pnevmo_stages (Приложение 8, "Этапы
        пневматического испытания трубопровода") -- 5 строк, ступени
        давления от p_rab_kgs (0,3 / 0,6 / 1,0 / 1,25 и само значение),
        время выдержки по методике (10 мин на первых трёх ступенях, 15 на
        четвёртой, для пятой -- прочерк). Срабатывает при каждом изменении
        p_rab_kgs, полностью перезаписывая таблицу (тот же принцип, что и у
        _update_pnevmo_pressure())."""
        table = self.table_pnevmo_stages
        try:
            p_rab_kgs = parse_ru(self.p_rab_kgs.toPlainText())
        except ValueError:
            table.setRowCount(0)
            return
        stages = [
            (format_ru_fixed(p_rab_kgs * 0.3, 1), "10"),
            (format_ru_fixed(p_rab_kgs * 0.6, 1), "10"),
            (format_ru_fixed(p_rab_kgs * 1.0, 1), "10"),
            (format_ru_fixed(p_rab_kgs * 1.25, 1), "15"),
            (self.p_rab_kgs.toPlainText(), "-"),
        ]
        table.setRowCount(len(stages))
        for row, (pressure, hold_time) in enumerate(stages):
            table.setItem(row, 0, QTableWidgetItem(str(row + 1)))
            table.setItem(row, 1, QTableWidgetItem(pressure))
            table.setItem(row, 2, QTableWidgetItem(hold_time))

    def _update_pnevmo_pressure_hint(self):
        """Зеркалит "Испытательное давление" (pnevmo_pressure, Приложения
        8-9) в read-only pnevmo_pressure_hint рядом с result_76 (раздел
        7.6). Плейсхолдеры не вкладываются друг в друга, поэтому текст 7.6
        оставляет пробел "P=__________ кгс/см2" под ручное заполнение --
        это поле только подсказывает уже введённое оператором значение,
        само оно в шаблон .docx не попадает."""
        self.pnevmo_pressure_hint.setPlainText(self.pnevmo_pressure.toPlainText())

    def _update_result_76_pressure(self):
        """Подставляет введённое испытательное давление вместо заглушки
        "P=__________" в свободном тексте result_76 (раздел 7.6) -- работает,
        только пока заглушка ещё не тронута оператором, иначе правку текста
        не перезаписывает (тот же принцип, что и у final_years_allowed/
        final_deadline_date)."""
        text = self.result_76.toPlainText()
        if "P=__________" in text:
            self.result_76.setPlainText(
                text.replace("P=__________", f"P={self.pnevmo_pressure.toPlainText()}")
            )

    def _update_final_deadline_date(self):
        """Дата из report_year + final_years_allowed (раздел 8.4.1) -- год
        отчёта плюс разрешённый срок эксплуатации в годах, день/месяц
        берутся из report_date (при его отсутствии -- текущая дата).
        Пересчитывается при каждом изменении final_years_allowed или
        report_year, перезаписывая прежнее значение поля."""
        try:
            years = int(parse_ru(self.final_years_allowed.toPlainText()))
            report_year = int(parse_ru(self.report_year.toPlainText()))
        except ValueError:
            return
        day_month_source = self.report_date.date()
        if day_month_source == QDate(2000, 1, 1):
            day_month_source = QDate.currentDate()
        base = QDate(report_year, day_month_source.month(), day_month_source.day())
        if not base.isValid():
            base = QDate(report_year, day_month_source.month(), 28)
        self.final_deadline_date.setDate(base.addYears(years))

    def _choose_nk_scheme(self):
        """Загрузка изображения схемы НК (Приложение 7). Тот же приём, что
        и клише сотрудника (src/ui/employees_tab.py) -- копия файла в
        общую папку под uuid-именем, само имя лежит в self.data и едет в
        Project.report_data при сохранении/открытии проекта."""
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Выбрать схему НК", "", "Изображения (*.png *.jpg *.jpeg)"
        )
        if not file_path:
            return
        filename = store_kleishe_image(Path(file_path), dest_dir=NK_SCHEME_DIR)
        self.data["nk_scheme_filename"] = filename
        self._set_nk_scheme_preview(filename)

    def _clear_nk_scheme(self):
        self.data["nk_scheme_filename"] = None
        self._set_nk_scheme_preview(None)

    def _set_nk_scheme_preview(self, filename):
        label = self.nk_scheme_preview
        if not filename:
            label.setPixmap(QPixmap())
            label.setText("нет изображения")
            return

        pixmap = QPixmap(str(NK_SCHEME_DIR / filename))
        if pixmap.isNull():
            label.setPixmap(QPixmap())
            label.setText("не удалось загрузить")
            return

        label.setText("")
        label.setPixmap(pixmap.scaled(
            label.width(), label.height(), Qt.AspectRatioMode.KeepAspectRatio,
        ))

    def _choose_pnevmo_graph(self):
        """Загрузка изображения графика нагружения (Приложение 8, Рисунок
        1) -- полная копия _choose_nk_scheme() (Приложение 7) под новую
        папку PNEVMO_GRAPH_DIR."""
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Выбрать график нагружения", "", "Изображения (*.png *.jpg *.jpeg)"
        )
        if not file_path:
            return
        filename = store_kleishe_image(Path(file_path), dest_dir=PNEVMO_GRAPH_DIR)
        self.data["pnevmo_graph_filename"] = filename
        self._set_pnevmo_graph_preview(filename)

    def _clear_pnevmo_graph(self):
        self.data["pnevmo_graph_filename"] = None
        self._set_pnevmo_graph_preview(None)

    def _set_pnevmo_graph_preview(self, filename):
        label = self.pnevmo_graph_preview
        if not filename:
            label.setPixmap(QPixmap())
            label.setText("нет изображения")
            return

        pixmap = QPixmap(str(PNEVMO_GRAPH_DIR / filename))
        if pixmap.isNull():
            label.setPixmap(QPixmap())
            label.setText("не удалось загрузить")
            return

        label.setText("")
        label.setPixmap(pixmap.scaled(
            label.width(), label.height(), Qt.AspectRatioMode.KeepAspectRatio,
        ))

    def _table_to_dicts(self, table, keys):
        """Конвертирует QTableWidget (по столбцам, слева направо) в
        list[dict] по переданным keys -- для полей, которые в self.data
        должны попасть списком словарей под Jinja-цикл {% for %}, а не
        списком списков "как есть" (см. get_form_data(), TABLE_WIDGET).
        Ячейка может быть обычным QTableWidgetItem или виджетом (см.
        table_specialists -- редактируемые комбобоксы), поэтому при
        отсутствии item проверяем cellWidget."""
        rows = []
        for row in range(table.rowCount()):
            item_dict = {}
            for col, key in enumerate(keys):
                item = table.item(row, col)
                if item is not None:
                    item_dict[key] = item.text()
                else:
                    cell_widget = table.cellWidget(row, col)
                    if isinstance(cell_widget, QComboBox):
                        item_dict[key] = cell_widget.currentText()
                    else:
                        item_dict[key] = ""
            rows.append(item_dict)
        return rows

    # Поля, чей плейсхолдер в шаблоне трубопровода использует докстпл-синтаксис
    # {{r field }} (сырой XML, не экранированный текст) -- см. _apply_line_breaks().
    # Список должен буквально совпадать с тем, что реально помечено {{r %}
    # в templates/Шаблон_трубопровод.docx -- {{r %} требует RichText-совместимое
    # значение ВСЕГДА (даже однострочное: RichText("текст") без \n -- валидный
    # <w:r><w:t>текст</w:t></w:r>), иначе докстпл вставляет голый текст мимо
    # <w:t>, и Word/python-docx его теряют -- проверено эмпирически на
    # org_activity_scope (однострочное значение, {{r %} без RichText → пропало).
    RICH_TEXT_FIELDS = frozenset({
        "intro_text", "result_71", "result_72", "result_73", "result_74",
        "result_75", "result_76", "result_77", "act2_intro_text",
        "vik_guidance_docs", "thick_guidance_docs", "uzk_guidance_docs",
        "uzk_evaluation_text", "calc_gost_basis", "calc_residual_methodology_note",
        "calc_residual_formula_text", "calc_corrosion_formula_text",
        "calc_worked_example_note", "ae_zakl_sources_text",
        "ae_zakl_evaluation_text", "org_activity_scope",
        "pnevmo_sensor_placement_note", "org_license_issuer",
    })

    # Поля из RICH_TEXT_FIELDS, чей плейсхолдер-run в шаблоне помечен italic
    # -- RichText не наследует форматирование run'а, который заменяет (см.
    # _apply_line_breaks), поэтому курсив нужно проставлять явно, иначе он
    # тихо теряется при рендере (эмпирически найдено на org_activity_scope).
    RICH_TEXT_ITALIC_FIELDS = frozenset({"org_activity_scope"})

    @classmethod
    def _apply_line_breaks(cls, form_data: dict) -> None:
        """Заменяет в form_data (на месте) значения полей из RICH_TEXT_FIELDS
        на docxtpl.RichText с настоящими переносами строк (<w:br/> между
        строками, без новых <w:p> -- так же устроены реальные переносы в
        эталонном документе). Применяется КО ВСЕМ полям из списка, а не
        только к содержащим '\\n' -- их плейсхолдер в шаблоне уже помечен
        {{r %} (сырой XML), и без RichText-обёртки даже однострочное
        значение теряется при рендере (см. RICH_TEXT_FIELDS)."""
        for key in cls.RICH_TEXT_FIELDS:
            value = form_data.get(key)
            if not isinstance(value, str):
                continue
            italic = key in cls.RICH_TEXT_ITALIC_FIELDS
            rt = RichText()
            lines = value.split("\n")
            for i, line in enumerate(lines):
                if i > 0:
                    rt.xml += "<w:r><w:br/></w:r>"
                rt.add(line, italic=italic)
            form_data[key] = rt


if __name__ == "__main__":
    app = QApplication([])
    window = MainWindow()
    window.show()
    app.exec()
