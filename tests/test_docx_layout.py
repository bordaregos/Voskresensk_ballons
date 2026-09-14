import io
import struct
import zlib

from docx import Document
from docx.oxml.ns import qn

from src.services.docx_layout import float_drawings_behind_text


def _minimal_png_bytes() -> bytes:
    """1x1 RGB PNG без внешних зависимостей (PIL не тянем в requirements-dev)."""
    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    signature = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw_scanline = b"\x00" + b"\xff\x00\x00"
    idat = zlib.compress(raw_scanline)
    return signature + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


def _add_inline_picture(document, width=100000):
    run = document.add_paragraph().add_run()
    run.add_picture(io.BytesIO(_minimal_png_bytes()), width=width)
    drawing = run._element.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    blip = inline.find(f".//{qn('a:blip')}")
    rid = blip.get(qn("r:embed"))
    return drawing, rid


def test_no_target_rids_changes_nothing():
    document = Document()
    _add_inline_picture(document)

    changed = float_drawings_behind_text(document, set())

    assert changed == 0


def test_unrelated_rid_is_not_touched():
    document = Document()
    drawing, rid = _add_inline_picture(document)

    changed = float_drawings_behind_text(document, {"rId999"})

    assert changed == 0
    assert drawing.find(qn("wp:inline")) is not None
    assert drawing.find(qn("wp:anchor")) is None


def test_matching_rid_converts_inline_to_anchor_behind_text():
    document = Document()
    drawing, rid = _add_inline_picture(document)
    inline_extent = drawing.find(qn("wp:inline")).find(qn("wp:extent"))
    original_cx, original_cy = inline_extent.get("cx"), inline_extent.get("cy")

    changed = float_drawings_behind_text(document, {rid})

    assert changed == 1
    assert drawing.find(qn("wp:inline")) is None
    anchor = drawing.find(qn("wp:anchor"))
    assert anchor is not None
    assert anchor.get("behindDoc") == "1"
    assert anchor.get("layoutInCell") == "1"

    # Размер не изменился -- extent перенесён, а не создан заново.
    extent = anchor.find(qn("wp:extent"))
    assert extent.get("cx") == original_cx
    assert extent.get("cy") == original_cy

    # Картинка (a:graphic) и её blip/rId тоже перенесены как есть.
    blip = anchor.find(f".//{qn('a:blip')}")
    assert blip.get(qn("r:embed")) == rid


def test_multiple_matching_drawings_get_distinct_relative_heights():
    document = Document()
    _, rid1 = _add_inline_picture(document)
    _, rid2 = _add_inline_picture(document)

    changed = float_drawings_behind_text(document, {rid1, rid2})

    assert changed == 2
    heights = set()
    for drawing in document.element.body.iter(qn("w:drawing")):
        anchor = drawing.find(qn("wp:anchor"))
        assert anchor is not None
        heights.add(anchor.get("relativeHeight"))
    assert len(heights) == 2
